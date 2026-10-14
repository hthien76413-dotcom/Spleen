#!/usr/bin/env python3
"""Tests for tools/compare_extractions.py. All data are invented.

Run from the repository root:   python3 tools/test_compare_extractions.py
(not with -I: the optional cross-check needs the `formulas` package, which may live in the user site-packages;
without it that one test is skipped.)

What is checked
  1. kappa arithmetic against a textbook example (a=20, b=5, c=10, d=15 -> po 0.70, pe 0.50, kappa 0.40) and the
     undefined case (no variation).
  2. every column of the template (tools/build_extraction_workbook.py, read as source, not executed) is classified in
     compare_extractions.py exactly once, so a change to the template cannot silently go uncompared.
  3. the Python re-computation of the derived columns equals the workbook's own formulas (evaluated with `formulas`).
  4. an end-to-end run on two invented extractions with known disagreements; counts and kappas are hand-calculated.
"""
import ast
import os
import subprocess
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)

import compare_extractions as ce  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

BUILDER = os.path.join(HERE, "build_extraction_workbook.py")
ROB_BUILDER = os.path.join(HERE, "build_rob_workbook.py")


def template_codes(builder=BUILDER, names=("PATIENT_COLS", "REPORT_COLS")):
    """Field codes of the named column lists, read from a generator's source with ast (the generator is not run)."""
    with open(builder, encoding="utf-8") as fh:
        tree = ast.parse(fh.read())
    out = {}
    for node in tree.body:
        if isinstance(node, ast.Assign) and isinstance(node.targets[0], ast.Name):
            name = node.targets[0].id
            if name in names:
                out[name] = [row[0] for row in ast.literal_eval(node.value)]
    return out


def build_template(path, builder=BUILDER):
    subprocess.run([sys.executable, "-I", builder, path], check=True, capture_output=True)


def col_index(ws):
    return {str(ws.cell(row=1, column=c).value): c for c in range(1, ws.max_column + 1) if ws.cell(row=1, column=c).value}


class KappaTests(unittest.TestCase):
    def test_textbook_example(self):
        pairs = [("Y", "Y")] * 20 + [("Y", "N")] * 5 + [("N", "Y")] * 10 + [("N", "N")] * 15
        k = ce.cohen_kappa(pairs)
        self.assertEqual(k["n"], 50)
        self.assertAlmostEqual(k["po"], 0.70)
        self.assertAlmostEqual(k["pe"], 0.50)
        self.assertAlmostEqual(k["kappa"], 0.40)

    def test_undefined_without_variation(self):
        k = ce.cohen_kappa([("V", "V")] * 4)
        self.assertIsNone(k["kappa"])
        self.assertAlmostEqual(k["po"], 1.0)

    def test_empty(self):
        self.assertEqual(ce.cohen_kappa([])["n"], 0)

    def test_perfect_and_chance(self):
        self.assertAlmostEqual(ce.cohen_kappa([("A", "A"), ("B", "B")])["kappa"], 1.0)
        self.assertAlmostEqual(ce.cohen_kappa([("A", "B"), ("B", "A")])["kappa"], -1.0)


class TemplateCoverageTests(unittest.TestCase):
    def test_every_template_column_is_classified_once(self):
        codes = template_codes()
        for name, fields in (("PATIENT_COLS", ce.PATIENT_FIELDS), ("REPORT_COLS", ce.REPORT_FIELDS)):
            classified = [c for group in fields.values() for c in group]
            self.assertEqual(len(classified), len(set(classified)), f"{name}: a column is in two groups")
            self.assertEqual(sorted(classified), sorted(codes[name]),
                             f"{name}: template and compare_extractions.py disagree on the column list")

    def test_rob_template_columns_are_classified_once(self):
        codes = template_codes(ROB_BUILDER, ("ROB_COLS",))["ROB_COLS"]
        classified = [c for group in ce.ROB_FIELDS.values() for c in group]
        self.assertEqual(len(classified), len(set(classified)))
        self.assertEqual(sorted(classified), sorted(codes))


class DerivedTests(unittest.TestCase):
    def test_rules_by_hand(self):
        d = ce.derive_category
        self.assertEqual(d(2, None), "<=2 d")
        self.assertEqual(d(2.5, None), "3-7 d")     # boundary rule: whole-day classes, 2.5 falls in the next one
        self.assertEqual(d(7, None), "3-7 d")
        self.assertEqual(d(7.5, None), "8-30 d")
        self.assertEqual(d(30, None), "8-30 d")
        self.assertEqual(d(31, None), ">30 d")
        self.assertEqual(d("NR", "3-7 d"), "3-7 d")  # rule 5: imprecise but all readings in one class
        self.assertEqual(d("NR", None), "NR")
        self.assertIsNone(d(None, None))
        o = ce.derive_outcome
        self.assertEqual(o("Non-viable", "Total splenectomy"), "Non-viable")
        self.assertEqual(o("Viable", "Total splenectomy"), "Viable")   # removed but viable still counts as viable
        self.assertEqual(o("Partially non-viable", "Partial splenectomy +/- splenopexy"), "Non-viable")
        self.assertEqual(o("Partially non-viable", "Splenopexy"), "Review")
        self.assertEqual(o("NR", "Total splenectomy"), "NR - removed, check reason")
        self.assertEqual(o("NR", "Splenopexy"), "NR")
        self.assertEqual(o("Not applicable", "No operation"), "Not applicable")
        self.assertIsNone(o(None, "Splenopexy"))

    @unittest.skipUnless(__import__("importlib").util.find_spec("formulas"), "package `formulas` not installed")
    def test_python_derivation_equals_workbook_formulas(self):
        import formulas  # noqa: WPS433
        with tempfile.TemporaryDirectory() as tmp:
            tpl = os.path.join(tmp, "tpl.xlsx")
            build_template(tpl)
            wb = load_workbook(tpl)
            ws = wb["Patients"]
            idx = col_index(ws)
            letters = {c: ws.cell(row=1, column=i).column_letter for c, i in idx.items()}
            # keep the formulas in a few rows only, so evaluation is fast
            cases = [
                # (idx_dur_days, idx_dur_cat_imprecise, viability, procedure, platelets)
                (0.5, None, "Non-viable", "Total splenectomy", 451),
                (2, None, "Viable", "Splenopexy", 450),
                (2.5, None, "Partially non-viable", "Partial splenectomy +/- splenopexy", "NR"),
                (7, None, "Partially non-viable", "Splenopexy", 100),
                (7.5, None, "NR", "Total splenectomy", None),
                (30, None, "NR", "Detorsion only", 300),
                (31, None, "Not applicable", "No operation", 300),
                ("NR", "3-7 d", "Viable", "Total splenectomy", "NR"),
                ("NR", None, "Viable", "Autotransplantation", 500),
                (None, None, None, None, None),
            ]
            first = ce.FIRST_DATA_ROW
            n_keep = len(cases)
            for r in range(first + n_keep, ws.max_row + 1):
                for code in ("own_patient", "idx_dur_category", "thrombocytosis", "primary_outcome"):
                    ws.cell(row=r, column=idx[code]).value = None
            for i, (days, imp, via, proc, plt) in enumerate(cases):
                r = first + i
                for code, val in (("report_id", "WS0001"), ("idx_dur_days", days), ("idx_dur_cat_imprecise", imp),
                                  ("viability", via), ("procedure", proc), ("platelets", plt)):
                    ws.cell(row=r, column=idx[code]).value = val
            for name in wb.sheetnames:  # evaluating the other sheets is not needed
                if name != "Patients":
                    del wb[name]
            # data validations reference the deleted Lists sheet; drop them for the evaluation copy
            ws.data_validations.dataValidation = []
            small = os.path.join(tmp, "small.xlsx")
            wb.save(small)
            model = formulas.ExcelModel().loads(small).finish()
            sol = model.calculate()
            fname = os.path.basename(small).upper()

            def cell(col, row):
                key = f"'[{os.path.basename(small)}]PATIENTS'!{letters[col]}{row}"
                for k in sol:
                    if k.upper() == key.upper():
                        v = sol[k].value[0, 0]
                        return None if v in ("", None) or str(v) == "empty" else v
                self.fail(f"cell {key} not found in the evaluated model")

            for i, (days, imp, via, proc, plt) in enumerate(cases):
                r = first + i
                self.assertEqual(cell("idx_dur_category", r), ce.derive_category(ce.clean(days), imp), f"category, row {r}")
                self.assertEqual(cell("primary_outcome", r), ce.derive_outcome(via, proc), f"outcome, row {r}")
                self.assertEqual(cell("thrombocytosis", r), ce.derive_thrombocytosis(ce.clean(plt)), f"platelets, row {r}")


class EndToEndTests(unittest.TestCase):
    """Two invented extractions. Hand calculation is in the comments."""

    ROWS_E1 = {
        # pid: (report_id, idx_dur_days, idx_dur_cat_imprecise, viability, procedure, sex)
        "WS0001_P1": ("WS0001", 1, None, "Non-viable", "Total splenectomy", "M"),
        "WS0002_P1": ("WS0002", 3, None, "Viable", "Splenopexy", "F"),
        "WS0003_P1": ("WS0003", 7, None, "Viable", "Splenopexy", "F"),
        "WS0004_P1": ("WS0004", "NR", "3-7 d", "Viable", "Splenopexy", "M"),
        "WS0005_P1": ("WS0005", 40, None, "Non-viable", "Total splenectomy", "F"),
        "WS0006_P1": ("WS0006", 0.5, None, "Partially non-viable", "Partial splenectomy +/- splenopexy", "M"),
        "WS0007_P1": ("WS0007", 2, None, "Viable", "Splenopexy", "F"),
        "WS0008_P1": ("WS0008", 5, None, "Viable", "Splenopexy", "M"),          # only in E1
        "WS0010_P1": ("WS0010", 10, None, "Viable", "Splenopexy", "M"),
        "WS0011_P1": ("WS0011", 100, None, "viable", "Splenopexy", "F"),         # lower case: must still agree
    }
    ROWS_E2 = {
        "WS0001_P1": ("WS0001", 1.2, None, "Non-viable", "Total splenectomy", "M"),   # 1 vs 1.2: agree within tolerance
        "WS0002_P1": ("WS0099", 3, None, "Viable", "Splenopexy", "F"),                # report_id differs: identity flag
        "WS0003_P1": ("WS0003", 8, None, "Viable", "Splenopexy", "F"),                # 7 vs 8: within tolerance, category differs
        "WS0004_P1": ("WS0004", "NR", None, "Viable", "Splenopexy", "M"),             # imprecise category blank in E2
        "WS0005_P1": ("WS0005", 25, None, "Non-viable", "Total splenectomy", "F"),    # 40 vs 25: disagreement
        "WS0006_P1": ("WS0006", 0.5, None, "Partially non-viable", "Partial splenectomy +/- splenopexy", "M"),
        "WS0007_P1": ("WS0007", 2, None, "Non-viable", "Total splenectomy", "F"),     # outcome disagreement
        "WS0009_P1": ("WS0009", 4, None, "Viable", "Splenopexy", "M"),                # only in E2
        "WS0010_P1": ("WS0010", 10, None, "Viable", "Splenopexy", "F"),               # sex disagreement
        "WS0011_P1": ("WS0011", 100, None, "Viable", "Splenopexy", "F"),
    }

    def fill(self, tpl, out, rows, extra_dup=None):
        wb = load_workbook(tpl)
        ws = wb["Patients"]
        idx = col_index(ws)
        r = ce.FIRST_DATA_ROW
        items = list(rows.items()) + ([extra_dup] if extra_dup else [])
        for pid, (rep, days, imp, via, proc, sex) in items:
            for code, val in (("patient_id", pid), ("report_id", rep), ("idx_dur_days", days),
                              ("idx_dur_cat_imprecise", imp), ("viability", via), ("procedure", proc), ("sex", sex)):
                ws.cell(row=r, column=idx[code]).value = val
            r += 1
        wb.save(out)

    def test_end_to_end(self):
        with tempfile.TemporaryDirectory() as tmp:
            tpl = os.path.join(tmp, "tpl.xlsx")
            build_template(tpl)
            f1, f2, rep = (os.path.join(tmp, n) for n in ("e1.xlsx", "e2.xlsx", "cmp.xlsx"))
            self.fill(tpl, f1, self.ROWS_E1, extra_dup=("WS0001_P1", ("WS0001", 9, None, "Viable", "Splenopexy", "M")))
            self.fill(tpl, f2, self.ROWS_E2)
            self.assertEqual(ce.main([f1, f2, "-o", rep]), 0)

            wb = load_workbook(rep)
            dis = [tuple(c.value for c in row) for row in wb["Disagreements"].iter_rows(min_row=2)]
            got = sorted((d[1], d[2], d[5]) for d in dis)
            self.assertEqual(got, sorted([
                ("WS0002_P1", "report_id", "identity"),
                ("WS0005_P1", "idx_dur_days", "numeric"),
                ("WS0007_P1", "viability", "categorical"),
                ("WS0007_P1", "procedure", "categorical"),
                ("WS0010_P1", "sex", "categorical"),
            ]))
            inc = [tuple(c.value for c in row) for row in wb["Incomplete"].iter_rows(min_row=2)]
            self.assertEqual([(i[1], i[2]) for i in inc], [("WS0004_P1", "idx_dur_cat_imprecise")])
            unm = sorted((u[1], u[2]) for u in (tuple(c.value for c in row) for row in wb["Unmatched"].iter_rows(min_row=2)))
            self.assertEqual(unm, [("E1", "WS0001_P1"), ("E1 only", "WS0008_P1"), ("E2 only", "WS0009_P1")])

            kap = {r[1]: r for r in (tuple(c.value for c in row) for row in wb["Kappa"].iter_rows(min_row=2))}
            # Primary exposure category, 9 matched patients with a category from both:
            #   E1: <=2 d x3, 3-7 d x3, 8-30 d x1, >30 d x2;  E2: <=2 d x3, 3-7 d x1, 8-30 d x3, >30 d x1, NR x1
            #   agreements 6 -> po = 6/9; pe = (3*3 + 3*1 + 1*3 + 2*1 + 0*1)/81 = 17/81
            exp = next(v for k, v in kap.items() if k.startswith("主要暴露"))
            self.assertEqual(exp[2], 9)
            self.assertAlmostEqual(exp[3], 6 / 9)
            self.assertAlmostEqual(exp[4], (6 / 9 - 17 / 81) / (1 - 17 / 81))
            # Primary outcome (binary Non-viable vs other), 9 matched: E1 NV x3 (P1, P5, P6), E2 NV x4 (+P7)
            #   agreements 8 -> po = 8/9; pe = (3*4 + 6*5)/81 = 42/81
            out = next(v for k, v in kap.items() if k.startswith("主要结局：Non-viable vs"))
            self.assertEqual(out[2], 9)
            self.assertAlmostEqual(out[3], 8 / 9)
            self.assertAlmostEqual(out[4], (8 / 9 - 42 / 81) / (1 - 42 / 81))

            by = {(r[0], r[1]): r for r in (tuple(c.value for c in row) for row in wb["By_field"].iter_rows(min_row=2))}
            days = by[("Patients", "idx_dur_days")]
            # both filled with a value or NR: 9 matched patients; disagreements: only WS0005 (40 vs 25).
            # exact agreement misses 1 vs 1.2 and 7 vs 8 as well -> 6 of 9 exact, 8 of 9 within tolerance
            self.assertEqual((days[2], days[3], days[5]), (9, 8, 6))


class RoBTests(unittest.TestCase):
    """Risk-of-bias workbook: derived adequacy column and the two-reviewer comparison."""

    COMBOS = [  # (q_exposure, q_outcome, q_reporting)
        ("Yes", "Yes", "Yes"), ("Yes", "Yes", "No"), ("Yes", "No", "Yes"), ("No", "Yes", "Yes"),
        ("Yes", "Unclear", "Yes"), ("Yes", "Yes", "Unclear"), ("Unclear", "Unclear", "Unclear"),
        ("Yes", "Yes", None), (None, None, None), ("yes", "YES", "Yes"),
    ]

    def test_adequate_rule_by_hand(self):
        self.assertEqual(ce.derive_adequate("Yes", "Yes", "Yes"), "Yes")
        self.assertEqual(ce.derive_adequate("yes", "YES", "Yes"), "Yes")     # case-insensitive, as in Excel
        self.assertEqual(ce.derive_adequate("Yes", "Unclear", "Yes"), "No")  # Unclear is not "satisfied"
        self.assertEqual(ce.derive_adequate("Yes", "Yes", "No"), "No")
        self.assertIsNone(ce.derive_adequate("Yes", "Yes", None))            # a blank item leaves it open
        self.assertIsNone(ce.derive_adequate(None, None, None))

    @unittest.skipUnless(__import__("importlib").util.find_spec("formulas"), "package `formulas` not installed")
    def test_python_adequate_equals_workbook_formula(self):
        import formulas  # noqa: WPS433
        with tempfile.TemporaryDirectory() as tmp:
            tpl = os.path.join(tmp, "rob.xlsx")
            build_template(tpl, ROB_BUILDER)
            wb = load_workbook(tpl)
            ws = wb["RoB"]
            idx = col_index(ws)
            letter = {c: ws.cell(row=1, column=i).column_letter for c, i in idx.items()}
            first, n = ce.FIRST_DATA_ROW, len(self.COMBOS)
            for r in range(first + n, ws.max_row + 1):
                ws.cell(row=r, column=idx["adequate_primary"]).value = None
            for i, (e, o, p) in enumerate(self.COMBOS):
                for code, val in (("report_id", f"WS{i:04d}"), ("q_exposure", e), ("q_outcome", o), ("q_reporting", p)):
                    ws.cell(row=first + i, column=idx[code]).value = val
            for name in wb.sheetnames:
                if name != "RoB":
                    del wb[name]
            ws.data_validations.dataValidation = []
            small = os.path.join(tmp, "small.xlsx")
            wb.save(small)
            sol = formulas.ExcelModel().loads(small).finish().calculate()
            for i, (e, o, p) in enumerate(self.COMBOS):
                key = f"'[{os.path.basename(small)}]ROB'!{letter['adequate_primary']}{first + i}".upper()
                got = next(sol[k].value[0, 0] for k in sol if k.upper() == key)
                got = None if got in ("", None) or str(got) == "empty" else got
                self.assertEqual(got, ce.derive_adequate(e, o, p), f"combo {i}: {(e, o, p)}")

    def test_end_to_end_two_reviewers(self):
        r1 = {"WS0001": ("Yes", "Yes", "Yes", "Yes", "Yes", "Yes"), "WS0002": ("Yes", "Yes", "Yes", "Yes", "Yes", "Yes"),
              "WS0003": ("No", "Yes", "Yes", "Yes", "Yes", "Yes"), "WS0004": ("Yes", "Unclear", "Yes", "Yes", "Yes", "Yes"),
              "WS0005": ("Yes", "Yes", "Unclear", "Yes", "Yes", "Yes"), "WS0006": ("Yes", "Yes", "Yes", "Yes", "Yes", "Yes")}
        r2 = dict(r1)
        r2["WS0002"] = ("Yes", "No", "Yes", "Yes", "Yes", "Yes")        # outcome item differs
        r2["WS0005"] = ("Yes", "Yes", "Yes", "Yes", "Yes", "Yes")       # reporting item differs
        # tuple order: q_exposure, q_outcome, q_reporting, q_selection, q_alternative, q_followup
        codes = ("q_exposure", "q_outcome", "q_reporting", "q_selection", "q_alternative", "q_followup")
        with tempfile.TemporaryDirectory() as tmp:
            tpl = os.path.join(tmp, "rob.xlsx")
            build_template(tpl, ROB_BUILDER)
            paths = []
            for name, data in (("r1.xlsx", r1), ("r2.xlsx", r2)):
                wb = load_workbook(tpl)
                ws = wb["RoB"]
                idx = col_index(ws)
                for i, (rid, vals) in enumerate(data.items()):
                    ws.cell(row=ce.FIRST_DATA_ROW + i, column=idx["report_id"]).value = rid
                    for code, v in zip(codes, vals):
                        ws.cell(row=ce.FIRST_DATA_ROW + i, column=idx[code]).value = v
                p = os.path.join(tmp, name)
                wb.save(p)
                paths.append(p)
            rep = os.path.join(tmp, "cmp.xlsx")
            self.assertEqual(ce.main([paths[0], paths[1], "-o", rep, "--sheets", "RoB"]), 0)
            wb = load_workbook(rep)
            dis = sorted((tuple(c.value for c in row)[1:3]) for row in wb["Disagreements"].iter_rows(min_row=2))
            self.assertEqual(dis, [("WS0002", "q_outcome"), ("WS0005", "q_reporting")])
            kap = {r[1]: r for r in (tuple(c.value for c in row) for row in wb["Kappa"].iter_rows(min_row=2))}
            ade = next(v for k, v in kap.items() if "adequate_primary" in k)
            # adequate: pairs (Y,Y) (Y,N) (N,N) (N,N) (N,Y) (Y,Y) -> po 4/6; both raters Yes x3, No x3 -> pe 0.5 -> kappa 1/3
            self.assertEqual(ade[2], 6)
            self.assertAlmostEqual(ade[3], 4 / 6)
            self.assertAlmostEqual(ade[4], (4 / 6 - 0.5) / 0.5)


def _have_calc():
    import shutil
    return shutil.which("soffice") is not None


def _lo_sheet_csv(xlsx, sheet_no, outdir):
    """Open a workbook in LibreOffice Calc (a spreadsheet engine independent of both Python and `formulas`) and export
    one sheet (1-based number) as CSV; formulas are evaluated by Calc on load."""
    import csv
    os.makedirs(outdir, exist_ok=True)
    opt = f"csv:Text - txt - csv (StarCalc):44,34,76,1,,0,false,true,false,false,false,{sheet_no}"
    subprocess.run(["soffice", "--headless", "--convert-to", opt, "--outdir", outdir, xlsx],
                   capture_output=True, text=True, timeout=240)
    files = [f for f in os.listdir(outdir) if f.endswith(".csv")]
    assert files, "LibreOffice produced no CSV"
    with open(os.path.join(outdir, files[0]), encoding="utf-8", newline="") as fh:
        return list(csv.reader(fh))


@unittest.skipUnless(_have_calc(), "LibreOffice (soffice) not installed")
class LibreOfficeCalcTests(unittest.TestCase):
    """The workbooks open in a real spreadsheet engine, its results equal the Python rules, drop-downs survive.
    This is LibreOffice Calc, not Excel or WPS: it shows the files load and the formulas mean what they should,
    not that Excel's rendering is fine."""

    CASES = [  # (idx_dur_days, idx_dur_cat_imprecise, viability, procedure, platelets, report_id)
        (0.5, None, "Non-viable", "Total splenectomy", 451, "WS0001"), (2, None, "Viable", "Splenopexy", 450, "OWN"),
        (2.5, None, "Partially non-viable", "Partial splenectomy +/- splenopexy", "NR", "WS0003"),
        (7, None, "Partially non-viable", "Splenopexy", 100, "WS0004"), (7.5, None, "NR", "Total splenectomy", None, "WS0005"),
        (30, None, "NR", "Detorsion only", 300, "WS0006"), (31, None, "Not applicable", "No operation", 300, "WS0007"),
        ("NR", "3-7 d", "Viable", "Total splenectomy", "NR", "WS0008"), ("NR", None, "Viable", "Autotransplantation", 500, "WS0009"),
        (None, None, None, None, None, "WS0010"),
    ]

    def test_extraction_workbook_formulas_in_calc(self):
        with tempfile.TemporaryDirectory() as tmp:
            tpl = os.path.join(tmp, "ext.xlsx")
            build_template(tpl)
            wb = load_workbook(tpl)
            ws = wb["Patients"]
            idx = col_index(ws)
            for i, (d, imp, v, p, plt, rid) in enumerate(self.CASES):
                for code, val in (("patient_id", f"{rid}_P{i}"), ("report_id", rid), ("idx_dur_days", d),
                                  ("idx_dur_cat_imprecise", imp), ("viability", v), ("procedure", p), ("platelets", plt)):
                    ws.cell(row=ce.FIRST_DATA_ROW + i, column=idx[code]).value = val
            filled = os.path.join(tmp, "ext_filled.xlsx")
            wb.save(filled)
            rows = _lo_sheet_csv(filled, wb.sheetnames.index("Patients") + 1, os.path.join(tmp, "out"))
            hdr = rows[0]
            self.assertEqual(len(hdr), len(idx))
            for i, (d, imp, v, p, plt, rid) in enumerate(self.CASES):
                r = rows[ce.FIRST_DATA_ROW - 1 + i]
                got = {c: (r[hdr.index(c)] or None) for c in ("idx_dur_category", "primary_outcome", "thrombocytosis", "own_patient")}
                exp = {"idx_dur_category": ce.derive_category(ce.clean(d), imp), "primary_outcome": ce.derive_outcome(v, p),
                       "thrombocytosis": ce.derive_thrombocytosis(ce.clean(plt)), "own_patient": "Yes" if rid == "OWN" else "No"}
                self.assertEqual(got, exp, f"case {i}")
            blank = rows[ce.FIRST_DATA_ROW - 1 + len(self.CASES) + 2]   # a row nobody filled stays empty
            self.assertEqual([blank[hdr.index(c)] for c in ("own_patient", "idx_dur_category", "thrombocytosis", "primary_outcome")],
                             ["", "", "", ""])

    def test_rob_workbook_formula_in_calc(self):
        with tempfile.TemporaryDirectory() as tmp:
            tpl = os.path.join(tmp, "rob.xlsx")
            build_template(tpl, ROB_BUILDER)
            wb = load_workbook(tpl)
            ws = wb["RoB"]
            idx = col_index(ws)
            for i, (e, o, p) in enumerate(RoBTests.COMBOS):
                for code, val in (("report_id", f"WS{i:04d}"), ("q_exposure", e), ("q_outcome", o), ("q_reporting", p)):
                    ws.cell(row=ce.FIRST_DATA_ROW + i, column=idx[code]).value = val
            filled = os.path.join(tmp, "rob_filled.xlsx")
            wb.save(filled)
            rows = _lo_sheet_csv(filled, wb.sheetnames.index("RoB") + 1, os.path.join(tmp, "out"))
            col = rows[0].index("adequate_primary")
            for i, combo in enumerate(RoBTests.COMBOS):
                self.assertEqual(rows[ce.FIRST_DATA_ROW - 1 + i][col] or None, ce.derive_adequate(*combo), f"combo {combo}")

    def test_dropdown_validations_survive_a_calc_round_trip(self):
        with tempfile.TemporaryDirectory() as tmp:
            for name, builder in (("ext", BUILDER), ("rob", ROB_BUILDER)):
                src = os.path.join(tmp, f"{name}.xlsx")
                build_template(src, builder)
                outdir = os.path.join(tmp, f"{name}_out")
                os.makedirs(outdir)
                subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", outdir, src],
                               capture_output=True, text=True, timeout=240)
                a, b = load_workbook(src), load_workbook(os.path.join(outdir, f"{name}.xlsx"))
                for ws in a.worksheets:
                    n_a = sum(1 for d in ws.data_validations.dataValidation if d.type == "list")
                    n_b = sum(1 for d in b[ws.title].data_validations.dataValidation if d.type == "list")
                    self.assertEqual(n_a, n_b, f"{name}/{ws.title}: list validations lost in Calc")


if __name__ == "__main__":
    unittest.main(verbosity=2)
