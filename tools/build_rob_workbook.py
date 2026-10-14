#!/usr/bin/env python3
"""Build the risk-of-bias workbook template for the wandering-spleen systematic review.

Source of the items and of the adequacy rule: protocol v1.4 (tools/build_protocol.js) Section 15 and Section 18,
sensitivity analysis 5. The tool is that of Murad et al. for case reports and case series (protocol ref. 2); the
protocol names six of its items and omits challenge-rechallenge and dose-response. The decision rules on the Rules
sheet (when to answer Yes, No or Unclear) are Claude's operationalisation, awaiting the three reviewers' approval and a
calibration; they are not in the protocol.

Usage:  python3 -I tools/build_rob_workbook.py "<output>.xlsx"
Requires openpyxl.
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

VERSION = "v2026-10-14a"
OUT = sys.argv[1] if len(sys.argv) > 1 else "Risk of bias workbook - template (draft).xlsx"
N_ROWS = 800    # one row per report; each reviewer (and the consensus) has a separate file
FIRST = 4       # first data row; rows 1-3 = header block (field code / label / hint), as in the extraction workbook

LISTS = {
    "rob": ["Yes", "No", "Unclear"],
    "rob_fu": ["Yes", "No", "Unclear", "Not applicable"],
    "reviewer": ["R1", "R2", "Consensus"],
}
LIST_RANGE = {}

thin = Side(style="thin", color="BBBBBB")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
HDR_FONT = Font(bold=True, color="FFFFFF")
GREY = PatternFill("solid", fgColor="E7E6E6")
FILL = {"id": "44546A", "item": "2F75B5", "derived": "7F7F7F", "source": "7F7F7F"}

# (code, label, hint, kind, width, section) -- kind: text | date | ("list", name) | ("formula", key)
ROB_COLS = [
    ("report_id", "Report ID", "WS 编号（与提取表 Reports 表一致）；本机构 5 例用 OWN", "text", 12, "id"),
    ("reviewer", "Reviewer", "R1 = 舒俊，R2 = 彭飞，Consensus = 共识（讨论后的最终值）", ("list", "reviewer"), 12, "id"),
    ("date_assessed", "Date assessed", "yyyy-mm-dd", "date", 12, "id"),
    ("q_selection", "Selection", "病例是否代表研究者的全部经验？Yes / No / Unclear（规则见 Rules）", ("list", "rob"), 12, "item"),
    ("q_exposure", "Ascertainment of exposure", "症状开始和手术的时间是否写得精确？Yes / No / Unclear", ("list", "rob"), 13, "item"),
    ("q_outcome", "Ascertainment of outcome", "是否写明脾脏有无活力及其依据？Yes / No / Unclear", ("list", "rob"), 13, "item"),
    ("q_alternative", "Alternative explanations", "是否考虑并排除了其他解释？Yes / No / Unclear", ("list", "rob"), 13, "item"),
    ("q_followup", "Adequacy of follow-up", "随访是否足够？Yes / No / Unclear / Not applicable", ("list", "rob_fu"), 13, "item"),
    ("q_reporting", "Sufficiency of reporting", "报告是否详细到可以提取？Yes / No / Unclear", ("list", "rob"), 13, "item"),
    ("adequate_primary", "Adequate for primary question (derived)",
     "公式，勿改：暴露、结局、报告三项全为 Yes 才是 Yes，否则 No；三项有空白则留空", ("formula", "adequate"), 16, "derived"),
    ("support_selection", "Support: selection", "依据（页/表/原文一句话）", "text", 22, "source"),
    ("support_exposure", "Support: exposure", "依据（页/表/原文一句话）；必填", "text", 22, "source"),
    ("support_outcome", "Support: outcome", "依据（页/表/原文一句话）；必填", "text", 22, "source"),
    ("support_other", "Support: other items", "依据", "text", 22, "source"),
    ("notes", "Notes", "需要共识讨论的问题", "text", 30, "source"),
]


def f_adequate(r, L):
    e, o, p = (f'{L[c]}{r}' for c in ("q_exposure", "q_outcome", "q_reporting"))
    return (f'=IF(OR({e}="",{o}="",{p}=""),"",IF(AND({e}="Yes",{o}="Yes",{p}="Yes"),"Yes","No"))')


FORMULAS = {"adequate": f_adequate}

README_LINES = [
    ("title", f"偏倚风险评估工作簿（模板，草稿）{VERSION} — 儿童游走脾系统综述（PROSPERO CRD420261520329）"),
    ("p", "依据：协议 v1.4 第 15 节（Murad 等的病例报告和病例系列工具，协议参考文献 2）和第 18 节敏感性分析 5。协议列了六项：选择、暴露的确定、结局的确定、其他解释、随访是否足够、报告是否充分；'激发再激发'和'剂量反应'两项不适用，已省略。协议不计算总分；'暴露、结局、报告三项都满足'的报告才算对主要问题'足够'（adequate）。"),
    ("h", "1. 怎么用"),
    ("b", "每篇纳入的报告（含本机构 OWN）由两位评审员各自独立评一行：reviewer 填 R1 或 R2；两人各用自己的文件（建议 rob_R1.xlsx、rob_R2.xlsx），完成之前不看对方。分歧经讨论后，在共识文件（rob_consensus.xlsx）里填 reviewer = Consensus 的行。"),
    ("b", "评的是**报告**，不是患者：一篇报告一行。同一篇报告里各患者的情况不同时（如只有一例有精确时间），按报告整体判断，在 Notes 里写明；这一点协议没有规定，是 Claude 的建议。"),
    ("b", "每一项只能选 Yes / No / Unclear（随访一项可选 Not applicable）。判断规则见 Rules 表，**那是 Claude 的操作化，不是协议原文，三位确认后才生效**；建议在正式评估前，两位评审员先各评 10 篇，讨论并修订规则（与提取表的校准用同一批 10 篇即可）。"),
    ("b", "暴露、结局两项的依据（页、表或原文一句话）必填，以便共识时核对。"),
    ("b", "灰色列 adequate_primary 是公式，不要改：暴露、结局、报告三项全为 Yes → Yes，其他情况 → No，三项有空白 → 留空。**Unclear 算'不满足'**——这是 Claude 对协议'all satisfied'的读法，需要确认。"),
    ("b", "敏感性分析 5（只用被判为足够的报告）用共识文件的 adequate_primary；分析脚本 tools/analysis_ipd.R 读这个文件。"),
    ("b", "协议没有规定偏倚风险评估中分歧的裁决人（筛选里是 Jun Yang）。建议沿用 Jun Yang，需要确认。"),
    ("b", "不要在文件里写姓名、住院号或具体日期。"),
    ("h", "2. 没有验证的事"),
    ("b", "灰色列的公式已用三种方式核对：Python 的 formulas 库、独立的 Python 规则、LibreOffice Calc（10 组测试值，结果一致）；LibreOffice Calc 能正常打开本文件，所有下拉列表都被识别。**但没有在 Excel 或 WPS 里打开过，版面也没有人眼检查过。** 第一次用时请随便填几行试一下，有问题告诉 Claude。"),
    ("b", "Murad 工具原文的问题措辞这里没有逐字引用（协议只给了六项的概括）；使用前请对照原文（BMJ Evid Based Med 2018;23:60–3）核对每一项的含义。"),
]

RULES = [
    ("Selection / 选择",
     "协议：selection (whether the patients represent the whole experience of the investigators)",
     "Yes：连续系列，或明确写了纳入某机构某时期的全部患者。\nNo：明确或明显是挑选的（如从更大的一组中只选手术病例、只选切脾病例、'典型病例'）。\nUnclear：没有说明；单个病例报告默认 Unclear（无法判断类似患者是否未被报告），除非作者明确说明。"),
    ("Ascertainment of exposure / 暴露的确定",
     "协议：ascertainment of exposure (whether the timing of symptom onset and of operation is stated precisely)",
     "Yes：按提取规则 1–4 或 8 能算出指数发作时长（给了具体数值或日期），且手术时间相对入院的时间清楚（不需要规则 9 的下限）。\nNo：只有含糊说法（规则 5：'数天''病史很长'）或没有时间。\nUnclear：有数值但含义不清（如间歇加重型，不清楚这次发作从何时算起，规则 6）；或只能按规则 9 作为下限。"),
    ("Ascertainment of outcome / 结局的确定",
     "协议：ascertainment of outcome (whether splenic viability is stated, with its basis)",
     "Yes：写明脾脏有无活力，并写明依据（术中所见和/或组织学）。\nNo：没有写脾脏有无活力，只能从术式推断。\nUnclear：写了有无活力但没有依据。"),
    ("Alternative explanations / 其他解释",
     "协议：alternative explanations",
     "Yes：报告描述了本身就可能导致脾梗死或延误的因素（扭转程度、伴随的胃扭转或膈肌异常、既往手术、血栓、脾大小等），并考虑或排除了它们。\nNo：存在明显的其他解释而没有讨论。\nUnclear：没有讨论，也看不出有明显的其他因素。\n（这一项最难统一，校准时重点讨论。）"),
    ("Adequacy of follow-up / 随访是否足够",
     "协议：adequacy of follow-up",
     "Yes：报告了随访时间和末次随访时的脾脏状态，时间长到足以发现保脾后的复发或再扭转（具体下限协议没规定，校准时由评审员定一个并记在这里）。\nNo：没有随访，或明显太短。\nUnclear：只写了'随访良好'之类。\nNot applicable：对主要结局不适用（主要结局在手术时判断），例如切除了无活力的脾脏、没有保脾；这一项不进入 adequate_primary。"),
    ("Sufficiency of reporting / 报告是否充分",
     "协议：sufficiency of reporting",
     "Yes：每名患者的关键项（年龄、性别、表现、时间线、影像、手术、脾脏活力）都有，提取时不需要大量填 NR。\nNo：多数患者的几个关键项都是 NR。\nUnclear：介于两者之间。"),
    ("adequate_primary（衍生）",
     "协议：a report will be judged adequate for the primary question when the exposure, outcome and reporting items are all satisfied",
     "暴露、结局、报告三项都是 Yes → Yes；其他 → No；任一项空白 → 留空。Unclear 算'不满足'（Claude 的读法，待确认）。敏感性分析 5 用它。"),
]


def build_header(ws, cols):
    for ci, (code, label, hint, kind, width, sec) in enumerate(cols, start=1):
        cells = [ws.cell(row=1, column=ci, value=code), ws.cell(row=2, column=ci, value=label),
                 ws.cell(row=3, column=ci, value=hint)]
        for c in cells:
            c.border = BORDER
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for c in cells[:2]:
            c.fill = PatternFill("solid", fgColor=FILL[sec])
            c.font = HDR_FONT
        cells[2].font = Font(italic=True, size=9, color="404040")
        cells[2].fill = PatternFill("solid", fgColor="F2F2F2")
        ws.column_dimensions[get_column_letter(ci)].width = width
    for r, h in zip((1, 2, 3), (22, 48, 90)):
        ws.row_dimensions[r].height = h


def add_validations(ws, cols, n_rows):
    letters = {code: get_column_letter(i) for i, (code, *_r) in enumerate(cols, start=1)}
    last = FIRST + n_rows - 1
    for ci, (code, label, hint, kind, width, sec) in enumerate(cols, start=1):
        L = get_column_letter(ci)
        rng = f"{L}{FIRST}:{L}{last}"
        if isinstance(kind, tuple) and kind[0] == "list":
            dv = DataValidation(type="list", formula1=LIST_RANGE[kind[1]], allow_blank=True, showErrorMessage=True,
                                errorStyle="warning", errorTitle="Not in list",
                                error="值不在列表里。请从下拉列表选择；确需例外请在 Notes 里说明。")
            ws.add_data_validation(dv)
            dv.add(rng)
        elif kind == "date":
            for r in range(FIRST, last + 1):
                ws.cell(row=r, column=ci).number_format = "yyyy-mm-dd"
        elif isinstance(kind, tuple) and kind[0] == "formula":
            for r in range(FIRST, last + 1):
                ws.cell(row=r, column=ci, value=FORMULAS[kind[1]](r, letters)).fill = GREY
    ws.freeze_panes = ws.cell(row=3, column=3)


def put_readme(ws):
    ws.column_dimensions["A"].width = 150
    for r, (kind, text) in enumerate(README_LINES, start=1):
        c = ws.cell(row=r, column=1, value=("• " + text) if kind == "b" else text)
        c.alignment = WRAP
        if kind == "title":
            c.font = Font(bold=True, size=14)
        elif kind == "h":
            c.font = Font(bold=True, size=12, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="2F75B5")
        ws.row_dimensions[r].height = 24 if kind == "title" else 18 * max(1, -(-len(text) // 70))


def put_rules(ws):
    for col, w in (("A", 28), ("B", 52), ("C", 96)):
        ws.column_dimensions[col].width = w
    ws.cell(row=1, column=1, value="Items and decision rules / 评估项目和判断规则（规则是 Claude 的操作化，待确认）").font = Font(bold=True, size=12)
    for j, h in enumerate(["Item", "Protocol wording (Section 15)", "Suggested decision rule"], start=1):
        c = ws.cell(row=2, column=j, value=h)
        c.font, c.border = HDR_FONT, BORDER
        c.fill = PatternFill("solid", fgColor=FILL["item"])
    for r, (a, b, c_) in enumerate(RULES, start=3):
        for j, v in enumerate((a, b, c_), start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.alignment, c.border = WRAP, BORDER
        ws.row_dimensions[r].height = 15 * max(3, c_.count("\n") + 1 + len(c_) // 95)


def put_lists(wb):
    ws = wb.create_sheet("Lists")
    for ci, (name, vals) in enumerate(LISTS.items(), start=1):
        L = get_column_letter(ci)
        ws.cell(row=1, column=ci, value=name).font = Font(bold=True)
        for ri, v in enumerate(vals, start=2):
            ws.cell(row=ri, column=ci, value=v)
        LIST_RANGE[name] = f"Lists!${L}$2:${L}${len(vals) + 1}"
        ws.column_dimensions[L].width = 18
    return ws


def main():
    wb = Workbook()
    ws_readme = wb.active
    ws_readme.title = "README"
    ws_rob = wb.create_sheet("RoB")
    ws_rules = wb.create_sheet("Rules")
    put_lists(wb)
    put_readme(ws_readme)
    put_rules(ws_rules)
    build_header(ws_rob, ROB_COLS)
    add_validations(ws_rob, ROB_COLS, N_ROWS)
    for ws, col in ((ws_readme, "7F7F7F"), (ws_rob, "2F75B5"), (ws_rules, "7F7F7F")):
        ws.sheet_properties.tabColor = col
    wb["Lists"].sheet_properties.tabColor = "BBBBBB"
    wb.properties.title = f"Risk of bias workbook template {VERSION}"
    wb.properties.creator = "Wandering spleen SR team"
    wb.save(OUT)
    print("saved", OUT, "columns", len(ROB_COLS))


if __name__ == "__main__":
    main()
