"""Analyse the third reviewer's blind judgements (see tools/make_blind_adjudication_sheet.py).

Usage (from the repo root):
  python3 -I tools/blind_adjudication_analysis.py FILLED.xlsx OUT_DIR [FOLDER]
  FOLDER defaults to search/screening/formal/; after the re-screen use the folder written by tools/compare_rescreen.py
FILLED.xlsx is the sheet returned by the third reviewer (sheet '判断', columns blind_id, 决定, 理由, 备注). The key is read from
search/screening/formal/blind_adjudication_key.csv and the full results from search/screening/formal/screening_results_all.csv.

Definitions: 'advance' = Include or Maybe; 'exclude' = Exclude. In every conflict exactly one reviewer's status equals the third
reviewer's, so each conflict counts as one error for the other reviewer.
Writes in OUT_DIR: blind_adjudication_summary.txt, blind_adjudication_decisions.csv (key + third-reviewer decision),
fulltext_list_after_adjudication.csv, and, if a pre-specified trigger fires, expansion_ws_ids.csv.
"""
import csv, math, sys, collections
from openpyxl import load_workbook

filled, outdir = sys.argv[1], sys.argv[2]
F = sys.argv[3].rstrip('/') + '/' if len(sys.argv) > 3 else 'search/screening/formal/'   # folder with the key and screening_results_all.csv
ADV = ('Include', 'Maybe')


def wilson(k, n, z=1.96):
    if n == 0:
        return (float('nan'), float('nan'))
    p = k / n; d = 1 + z * z / n; c = p + z * z / (2 * n); h = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n))
    return ((c - h) / d, (c + h) / d)


def cp_upper(k, n, alpha=0.05):
    """Clopper-Pearson one-sided 95% upper bound for k events in n (bisection on the binomial CDF)."""
    lo, hi = k / n, 1.0
    for _ in range(60):
        mid = (lo + hi) / 2
        cdf = sum(math.comb(n, i) * mid ** i * (1 - mid) ** (n - i) for i in range(k + 1))
        if cdf > alpha:
            lo = mid
        else:
            hi = mid
    return hi


def sign_test_p(a, b):
    n = a + b
    if n == 0:
        return 1.0
    k = min(a, b)
    return min(1.0, 2 * sum(math.comb(n, i) for i in range(k + 1)) / 2 ** n)


ws = load_workbook(filled, data_only=True)['判断']
hdr = [c.value for c in ws[1]]
col = {h: i for i, h in enumerate(hdr)}
jy = {}
for row in ws.iter_rows(min_row=2, values_only=True):
    bid = row[col['blind_id']]
    if bid:
        jy[bid] = dict(decision=(row[col['决定']] or '').strip(), reason=(row[col['理由']] or '').strip(), note=(row[col['备注']] or '') if '备注' in col else '')
key = list(csv.DictReader(open(F + 'blind_adjudication_key.csv', encoding='utf-8-sig')))
res = {r['ws_id']: r for r in csv.DictReader(open(F + 'screening_results_all.csv', encoding='utf-8-sig'))}
out = []
problems = []
for k in key:
    j = jy.get(k['blind_id'])
    if j is None or j['decision'] not in ('Include', 'Maybe', 'Exclude'):
        problems.append(f"{k['blind_id']}: no valid decision")
        continue
    if j['decision'] == 'Exclude' and not j['reason']:
        problems.append(f"{k['blind_id']}: Exclude without a reason (guide section 3)")
    k = dict(k); k['JY_decision'] = j['decision']; k['JY_reason'] = j['reason']; k['JY_note'] = j['note']
    k['JY_adv'] = j['decision'] in ADV
    out.append(k)
done = {k['blind_id'] for k in out}
L = []
L.append(f'Third-reviewer judgements: {len(out)} of {len(key)} sheet records usable; problems: {len(problems)}')
L += ['  ' + p for p in problems[:30]]
if len(out) < len(key):
    L.append('INCOMPLETE: results below are provisional.')

# ---- conflicts
conf = [k for k in out if k['stratum'] == 'conflict']
c = collections.Counter()
for k in conf:
    r1a, r2a = k['R1_decision'] in ADV, k['R2_decision'] in ADV
    c[('R1 excl / R2 adv' if (not r1a and r2a) else 'R2 excl / R1 adv', 'JY advance' if k['JY_adv'] else 'JY exclude')] += 1
n_conf = len(conf)
r1_right = c[('R1 excl / R2 adv', 'JY exclude')] + c[('R2 excl / R1 adv', 'JY advance')]
r2_right = c[('R1 excl / R2 adv', 'JY advance')] + c[('R2 excl / R1 adv', 'JY exclude')]
L.append(f'\nConflicts judged: {n_conf}')
for kk, v in sorted(c.items()):
    L.append(f'  {kk[0]:18s} -> {kk[1]:11s} {v}')
r1w = wilson(r1_right, n_conf); r2w = wilson(r2_right, n_conf)
L.append(f'  Third reviewer agrees with R1 (Jun Shu) in {r1_right} ({r1_right / max(n_conf, 1):.1%}; 95% CI {r1w[0]:.1%}-{r1w[1]:.1%}), with R2 (Fei Peng) in {r2_right} ({r2_right / max(n_conf, 1):.1%}; CI {r2w[0]:.1%}-{r2w[1]:.1%}); sign test p {"< 0.001" if sign_test_p(r1_right, r2_right) < 0.001 else "= " + format(sign_test_p(r1_right, r2_right), ".3f")}')
fe1 = c[('R1 excl / R2 adv', 'JY advance')]; fe2 = c[('R2 excl / R1 adv', 'JY advance')]
L.append(f'  False exclusions (an eligible-looking record excluded): R1 {fe1}, R2 {fe2}.  Over-inclusions: R1 {c[("R2 excl / R1 adv", "JY exclude")]}, R2 {c[("R1 excl / R2 adv", "JY exclude")]}.')
mm = collections.Counter(k['JY_decision'] for k in conf); L.append(f'  Third-reviewer decisions on conflicts: {dict(mm)}')

# ---- QC
trig = []
tg = [k for k in out if k['stratum'] in ('targeted_flag', 'targeted_noabstract')]
rq = [k for k in out if k['stratum'] == 'random_qc']
for name, grp in (('targeted (word-flagged + no-abstract)', tg), ('  of which word-flagged', [k for k in tg if k['stratum'] == 'targeted_flag']), ('  of which no abstract', [k for k in tg if k['stratum'] == 'targeted_noabstract']), ('random QC sample', rq)):
    a = sum(1 for k in grp if k['JY_adv']); lo, hi = wilson(a, len(grp))
    L.append(f'\nBoth-excluded records, {name}: {len(grp)} judged, third reviewer advances {a} ({a / max(len(grp), 1):.1%}; Wilson 95% CI {lo:.1%}-{hi:.1%})')
a_t = sum(1 for k in tg if k['JY_adv']); a_r = sum(1 for k in rq if k['JY_adv'])
if rq:
    L.append(f'  Random QC: one-sided 95% upper bound for the share of eligible-looking records among the unflagged both-excluded records: {cp_upper(a_r, len(rq)):.1%}')
t1 = tg and a_t / len(tg) > 0.10
t2 = a_r >= 2
L.append(f'\nPre-specified trigger 1 (targeted advanced > 10%): {"FIRED" if t1 else "not fired"} ({a_t}/{len(tg)}).  Trigger 2 (>=2 of the random QC advanced): {"FIRED" if t2 else "not fired"} ({a_r}/{len(rq)}).')
if t1 or t2:
    already = {k['ws_id'] for k in key}
    exp = [r for r in res.values() if r['R1_JunShu_decision'] == 'Exclude' and r['R2_FeiPeng_decision'] == 'Exclude'
           and ('E2' in r['R1_JunShu_reason'] + r['R2_FeiPeng_reason'] or 'E3' in r['R1_JunShu_reason'] + r['R2_FeiPeng_reason']) and r['ws_id'] not in already]
    with open(outdir + '/expansion_ws_ids.csv', 'w', encoding='utf-8-sig', newline='') as f:
        w = csv.writer(f); w.writerow(['ws_id']); [w.writerow([r['ws_id']]) for r in exp]
    L.append(f'  -> {len(exp)} further both-excluded records with E2 or E3 are written to expansion_ws_ids.csv for a second blind sheet.')

# ---- final list
adv_final = {}
for r in res.values():
    if r['R1_JunShu_decision'] in ADV and r['R2_FeiPeng_decision'] in ADV:
        adv_final[r['ws_id']] = 'both advanced'
for k in out:
    if k['JY_adv']:
        adv_final.setdefault(k['ws_id'], {'conflict': 'conflict, third reviewer advanced', 'targeted_flag': 'both excluded, third reviewer advanced', 'targeted_noabstract': 'both excluded, third reviewer advanced', 'random_qc': 'both excluded, third reviewer advanced (QC)'}[k['stratum']])
with open(outdir + '/fulltext_list_after_adjudication.csv', 'w', encoding='utf-8-sig', newline='') as f:
    w = csv.writer(f); w.writerow(['ws_id', 'year', 'title', 'language', 'source'])
    for wsid, src in sorted(adv_final.items()):
        r = res[wsid]; w.writerow([wsid, r['year'], r['title'], r['language'], src])
src_c = collections.Counter(adv_final.values())
L.append(f'\nRecords advancing to full text so far: {len(adv_final)} ({dict(src_c)}); excluded at title/abstract so far: {len(res) - len(adv_final)} of {len(res)}')
if problems or len(out) < len(key):
    L.append('(provisional: incomplete or flawed rows above)')
with open(outdir + '/blind_adjudication_decisions.csv', 'w', encoding='utf-8-sig', newline='') as f:
    w = csv.writer(f); w.writerow(['blind_id', 'ws_id', 'stratum', 'R1_decision', 'R1_reason', 'R2_decision', 'R2_reason', 'JY_decision', 'JY_reason', 'JY_note'])
    for k in out:
        w.writerow([k[x] for x in ('blind_id', 'ws_id', 'stratum', 'R1_decision', 'R1_reason', 'R2_decision', 'R2_reason', 'JY_decision', 'JY_reason', 'JY_note')])
open(outdir + '/blind_adjudication_summary.txt', 'w', encoding='utf-8').write('\n'.join(L) + '\n')
print('\n'.join(L))
