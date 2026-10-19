#!/usr/bin/env python3
"""Check and compare the two extraction files (Reports and Patients sheets of the extraction workbook).

Protocol v1.5 Section 14 / Appendix C: two reviewers (E1 = Jun Shu, E2 = Hongqiang Bian) extract independently after calibration on
10 reports; agreement on the primary exposure and the primary outcome is reported as Cohen's kappa; disagreements are settled by
discussion and, failing agreement, by Jun Yang; Ji Wang (who does not extract) checks the files. Nobody looks at the other's file until
both have finished the calibration set or the whole extraction.

Two commands
  check    one file: format and completeness check, pace check from the EX_time_log sheet. Safe at a midpoint: it shows nothing of
           the other extractor's file.
  compare  two files: coverage of reports and patients, kappa for the primary exposure (four duration categories plus NR) and the
           primary outcome, agreement for every coded field, and the lists of disagreements for the discussion and the adjudicator.
           Run it only when both extractors have finished the set being compared (or with --aggregate-only: counts only, no WS ids,
           no output files).

Usage (from the repo root; needs openpyxl):
  python3 -I tools/compare_extraction.py check extraction_E1.xlsx [--extractor E1] [--max-per-hour 8]
  python3 -I tools/compare_extraction.py compare extraction_E1.xlsx extraction_E2.xlsx --out DIR [--aggregate-only]

Claude's readings (the protocol does not say; for the extractors and the adjudicator to confirm)
  * Derived columns (category, primary outcome, period) are NOT read from the files: a workbook that was never saved by Excel has no
    stored values for formulas. The script recomputes them from the raw columns with the same rules as the Rules sheet (D2, D3).
  * Primary exposure: category from idx_dur_days (<=2, 3-7, 8-30, >30 days; boundary values in the lower class); if the days are NR, the
    imprecise-duration category if one was given, else NR. A blank cell means 'not extracted yet': such patients are left out of the
    kappa and counted separately. Days are also compared: equal within a tolerance (default 0.5 day or 10%, whichever is larger).
  * Primary outcome: Non-viable / Viable / other (Review, NR, NR - removed, Not applicable), derived from viability and procedure. The
    raw viability field is compared as well.
  * Patients are matched by patient_id (WS0123_P1). If the two extractors numbered the patients of a report differently the patients
    show as 'only in one file' and the report is listed with both counts; that is settled in the discussion, not by the script.
  * Kappa below 30 matched patients is printed with a warning (the CI is very wide); the 10-report calibration has too few patients for
    a meaningful kappa in most cases, so read the disagreement lists.
Outputs contain no personal data; they do contain both extractors' record-level values, so give them to the adjudicator only after
both have finished.
"""
import argparse
import collections
import csv
import datetime
import importlib.util
import math
import os
import re
import runpy
import statistics
import sys

sys.dont_write_bytecode = True
HERE = os.path.dirname(os.path.abspath(__file__))
_spec = importlib.util.spec_from_file_location("cft", os.path.join(HERE, "compare_fulltext_logs.py"))
cft = importlib.util.module_from_spec(_spec)
sys.modules["cft"] = cft
_spec.loader.exec_module(cft)
txt, to_date, to_time, read_sheet, kappa, fmt_k = cft.txt, cft.to_date, cft.to_time, cft.read_sheet, cft.kappa, cft.fmt_k

from openpyxl import load_workbook


def load_builder():
    """Field definitions and allowed values from the workbook builder (single source of truth)."""
    saved = sys.argv
    sys.argv = ["build_extraction_workbook.py", os.devnull]
    try:
        return runpy.run_path(os.path.join(HERE, "build_extraction_workbook.py"), run_name="builder_import")
    finally:
        sys.argv = saved


B = load_builder()
LISTS = B["LISTS"]
PATIENT_COLS = B["PATIENT_COLS"]
REPORT_COLS = B["REPORT_COLS"]
PATIENT_CODES = [c[0] for c in PATIENT_COLS]
REPORT_CODES = [c[0] for c in REPORT_COLS]
KIND = {c[0]: c[3] for c in PATIENT_COLS}
CATS = LISTS["dur_category"]  # <=2 d, 3-7 d, 8-30 d, >30 d
DERIVED = {"own_patient", "idx_dur_category", "thrombocytosis", "primary_outcome"}
NOT_COMPARED = DERIVED | {"patient_id", "report_id", "extractor", "date_extracted", "notes", "idx_dur_rules", "src_age_sex",
                          "src_exposure", "src_outcome", "src_other", "age_verbatim", "idx_dur_verbatim", "total_hist_verbatim"}
# tolerance for numeric fields: (absolute, relative); a pair agrees if |a-b| <= max(abs, rel * max(a, b))
TOL_DEFAULT = (0.5, 0.10)
TOL = {"age_years": (0.06, 0.0), "platelets": (0.0, 0.02), "torsion_degrees": (15.0, 0.0), "spleen_size_cm": (0.5, 0.0),
       "followup_months": (1.0, 0.10), "n_earlier_episodes": (0.0, 0.0)}
DAY_FIELDS = ("idx_dur_days", "idx_dur_alt_days", "total_hist_days", "waiting_interval_days")
YES_NO_NR_NOTE = "NR"


# ----------------------------------------------------------------------------- reading
def num(v):
    """float for a number, 'NR' for NR, None for blank, the text for anything else."""
    s = txt(v)
    if s == "":
        return None
    if s.upper() == "NR":
        return "NR"
    try:
        return float(s.replace(",", "."))
    except ValueError:
        return s


def load(path):
    wb = load_workbook(path, data_only=True)
    for n in ("Reports", "Patients"):
        if n not in wb.sheetnames:
            raise SystemExit("%s has no %s sheet - use the template workbook." % (path, n))
    reports = read_sheet(wb, "Reports", REPORT_CODES)
    patients = read_sheet(wb, "Patients", PATIENT_CODES)
    return wb, reports, patients


def category(p):
    """Duration category of a patient dict (same rule as the workbook formula, from the raw columns); None = not extracted."""
    d = num(p["idx_dur_days"])
    if isinstance(d, float):
        return CATS[0] if d <= 2 else CATS[1] if d <= 7 else CATS[2] if d <= 30 else CATS[3]
    imp = txt(p["idx_dur_cat_imprecise"])
    if imp:
        return imp
    if d == "NR":
        return "NR"
    return None


def outcome(p):
    """Primary outcome of a patient dict (Rules D3). None = not extracted."""
    v = txt(p["viability"])
    pr = txt(p["procedure"])
    removal = pr in ("Total splenectomy", "Partial splenectomy +/- splenopexy")
    if v == "":
        return None
    if v == "Non-viable":
        return "Non-viable"
    if v == "Viable":
        return "Viable"
    if v == "Partially non-viable":
        return "Non-viable" if removal else "Review"
    if v == "NR":
        return "NR - removed, check reason" if removal else "NR"
    return "Not applicable"


def outcome3(o):
    return None if o is None else (o if o in ("Non-viable", "Viable") else "Other")


# ----------------------------------------------------------------------------- check one file
class Issues(cft.Issues):
    pass


def pace(wb, patients, reports, max_per_hour, issues, lines):
    if "EX_time_log" not in wb.sheetnames:
        lines.append("EX_time_log sheet: missing (use the current template, v2026-10-18a or later).")
        return
    rows = read_sheet(wb, "EX_time_log", ["date", "extractor", "start", "end", "n_reports", "n_patients", "note"])
    sessions = []
    for r in rows:
        n = r["_row"]
        d, s, e = to_date(r["date"]), to_time(r["start"]), to_time(r["end"])
        k, kp = txt(r["n_reports"]), txt(r["n_patients"])
        if d is None or s is None or e is None or not re.match(r"^\d+$", k) or (kp and not re.match(r"^\d+$", kp)):
            issues.add("WARN", n, "-", "EX_time_log row: date, start, end (hh:mm), n_reports or n_patients missing or unreadable", kind="time-unreadable")
            continue
        mins = (datetime.datetime.combine(d, e) - datetime.datetime.combine(d, s)).total_seconds() / 60.0
        if mins <= 0:
            issues.add("WARN", n, "-", "EX_time_log row: end is not after start", kind="time-order")
            continue
        sessions.append({"row": n, "date": d, "mins": mins, "n": int(k), "np": int(kp) if kp else 0})
    if not sessions:
        lines.append("EX_time_log: no usable sessions.")
        return
    tot_min = sum(s["mins"] for s in sessions)
    tot_n = sum(s["n"] for s in sessions)
    tot_np = sum(s["np"] for s in sessions)
    lines.append("EX_time_log: %d sessions, %.1f h, %d reports worked on, %d patients entered%s" % (
        len(sessions), tot_min / 60.0, tot_n, tot_np, "; %.0f min per report" % (tot_min / tot_n) if tot_n else ""))
    fast = [s for s in sessions if s["n"] > 0 and s["n"] / (s["mins"] / 60.0) > max_per_hour]
    if fast:
        lines.append("  to ask about: %d session(s) above %.0f reports per hour (rows %s)" % (
            len(fast), max_per_hour, ", ".join(str(s["row"]) for s in fast[:15])))
        for s in fast:
            issues.add("WARN", s["row"], "-", "EX_time_log: %d reports in %.0f min (%.0f per hour, above %.0f)" % (
                s["n"], s["mins"], s["n"] / (s["mins"] / 60.0), max_per_hour))
    n_rep = len({txt(r["report_id"]) for r in reports if txt(r["report_id"])})
    lines.append("  reports in the Reports sheet: %d; patient rows: %d%s" % (
        n_rep, len(patients), "" if (n_rep == tot_n and len(patients) == tot_np) else "  (time log adds up to %d reports and %d patients)" % (tot_n, tot_np)))
    if n_rep != tot_n:
        issues.add("WARN", 0, "-", "Reports sheet has %d reports but EX_time_log adds up to %d" % (n_rep, tot_n))
    if tot_np and len(patients) != tot_np:
        issues.add("WARN", 0, "-", "Patients sheet has %d rows but EX_time_log adds up to %d patients" % (len(patients), tot_np))
    days = collections.Counter(to_date(p["date_extracted"]) for p in patients if to_date(p["date_extracted"]))
    tdays = collections.defaultdict(int)
    for s in sessions:
        tdays[s["date"]] += s["np"]
    for d in sorted(days):
        if days[d] > tdays.get(d, 0):
            issues.add("WARN", 0, "-", "%s: %d patient rows dated that day, %d in EX_time_log" % (d, days[d], tdays.get(d, 0)))


def run_check(args):
    wb, reports, patients = load(args.file)
    issues = Issues()
    lines = ["File: %s (%d report rows, %d patient rows)" % (os.path.basename(args.file), len(reports), len(patients))]
    rep_ids = collections.Counter(txt(r["report_id"]) for r in reports)
    pat_ids = collections.Counter(txt(p["patient_id"]) for p in patients)
    for r in reports:
        n, rid = r["_row"], txt(r["report_id"])
        if not rid:
            issues.add("ERROR", n, "-", "Reports row without report_id")
        elif rep_ids[rid] > 1:
            issues.add("ERROR", n, rid, "report_id appears %d times in Reports" % rep_ids[rid], kind="dup-report")
        elif rid != "OWN" and not cft.WS_RE.match(rid):
            issues.add("WARN", n, rid, "report_id is not of the form WS0123 or OWN")
        for col in ("n_children_individual", "n_children_aggregate"):
            v = txt(r[col])
            if v != "" and not re.match(r"^\d+$", v):
                issues.add("ERROR", n, rid, "%s is not a whole number (%r)" % (col, v))
        for col, lst in (("language", "lang"), ("design", "design"), ("consecutive", "consecutive"), ("extractor", "extractor")):
            v = txt(r[col])
            if v and v not in LISTS[lst]:
                issues.add("ERROR", n, rid, "%s %r is not in the list" % (col, v))
    per_report = collections.Counter(txt(p["report_id"]) for p in patients)
    n_missing = collections.Counter()
    for p in patients:
        n, pid, rid = p["_row"], txt(p["patient_id"]), txt(p["report_id"])
        if not pid:
            issues.add("ERROR", n, "-", "Patients row without patient_id")
            continue
        if pat_ids[pid] > 1:
            issues.add("ERROR", n, pid, "patient_id appears %d times" % pat_ids[pid], kind="dup-patient")
        if not re.match(r"^(WS\d{4}|OWN)_P\d+$", pid):
            issues.add("WARN", n, pid, "patient_id is not of the form WS0123_P1 or OWN_P1")
        if rid not in rep_ids:
            issues.add("ERROR", n, pid, "report_id %r is not in the Reports sheet" % rid)
        elif pid and not pid.startswith(rid + "_P"):
            issues.add("WARN", n, pid, "patient_id does not start with its report_id %s" % rid)
        ex = txt(p["extractor"])
        if args.extractor and ex and ex != args.extractor:
            issues.add("WARN", n, pid, "extractor column does not say %s" % args.extractor, kind="extractor-mismatch")
        elif not ex:
            issues.add("WARN", n, pid, "extractor column empty", kind="extractor-empty")
        d = to_date(p["date_extracted"])
        if d is None:
            issues.add("WARN", n, pid, "date_extracted missing or not a date", kind="date-missing")
        elif d > datetime.date.today():
            issues.add("WARN", n, pid, "date_extracted is in the future", kind="date-future")
        # list values
        for col, kind in KIND.items():
            if isinstance(kind, tuple) and kind[0] == "list" and col in p:
                v = txt(p[col])
                if v and kind[1] in LISTS and v not in LISTS[kind[1]]:
                    issues.add("ERROR", n, pid, "%s %r is not in the list" % (col, v))
        # numbers
        for col in ("age_years", "idx_dur_days", "idx_dur_alt_days", "total_hist_days", "waiting_interval_days", "platelets",
                    "torsion_degrees", "spleen_size_cm", "followup_months", "n_earlier_episodes"):
            v = num(p[col])
            if isinstance(v, str) and v != "NR":
                issues.add("ERROR", n, pid, "%s must be a number or NR (found %r)" % (col, v))
            elif isinstance(v, float) and v < 0:
                issues.add("ERROR", n, pid, "%s is negative" % col)
        age = num(p["age_years"])
        if isinstance(age, float) and age >= 18:
            issues.add("ERROR", n, pid, "age_years %.2f is 18 or more: not eligible, should not be in this sheet" % age)
        # completeness of the primary variables
        d_ = num(p["idx_dur_days"])
        if d_ is None:
            n_missing["idx_dur_days"] += 1
        if txt(p["viability"]) == "":
            n_missing["viability"] += 1
        if txt(p["procedure"]) == "":
            n_missing["procedure"] += 1
        if isinstance(d_, float) and not txt(p["src_exposure"]):
            issues.add("WARN", n, pid, "duration given but src_exposure (source) empty", kind="no-src-exp")
        if txt(p["viability"]) not in ("", "NR", "Not applicable") and not txt(p["src_outcome"]):
            issues.add("WARN", n, pid, "viability given but src_outcome (source) empty", kind="no-src-out")
        if isinstance(d_, float) and not txt(p["idx_dur_verbatim"]):
            issues.add("WARN", n, pid, "duration in days without the verbatim wording", kind="no-verbatim")
        if txt(p["viability"]) in ("Viable", "Partially non-viable", "Non-viable") and not txt(p["viab_basis"]):
            issues.add("WARN", n, pid, "viability given without its basis (viab_basis)", kind="no-basis")
    for r in reports:
        rid = txt(r["report_id"])
        k = txt(r["n_children_individual"])
        if rid and k.isdigit() and per_report.get(rid, 0) != int(k):
            issues.add("WARN", r["_row"], rid, "n_children_individual is %s but %d patient row(s) in Patients" % (k, per_report.get(rid, 0)), kind="count-mismatch")
    n_done = len(patients)
    lines.append("Patients with the primary variables still blank (= not extracted): idx_dur_days %d, viability %d, procedure %d of %d" % (
        n_missing["idx_dur_days"], n_missing["viability"], n_missing["procedure"], n_done))
    cat = collections.Counter(category(p) or "blank" for p in patients)
    out = collections.Counter(outcome(p) or "blank" for p in patients)
    lines.append("Duration categories (recomputed): " + ", ".join("%s %d" % (k, cat[k]) for k in CATS + ["NR", "blank"] if cat.get(k)))
    lines.append("Primary outcome (recomputed): " + ", ".join("%s %d" % (k, v) for k, v in sorted(out.items())))
    pace(wb, patients, reports, args.max_per_hour, issues, lines)
    lines.append("Format check: %d ERROR(S), %d WARNING(S)" % (issues.count("ERROR"), issues.count("WARN")))
    lines.extend(issues.lines())
    print("\n".join(lines))


# ----------------------------------------------------------------------------- compare two files
def agree_value(field, a, b, tol_abs=None, tol_rel=None):
    """(state, text_a, text_b): state in 'agree', 'differ', 'blank-in-one', 'both-blank'."""
    ka, kb = KIND.get(field), None
    if isinstance(ka, tuple) or ka == "text":
        sa, sb = txt(a).lower(), txt(b).lower()
        if field == "same_patient_as":
            sa, sb = txt(a).upper(), txt(b).upper()
        if sa == "" and sb == "":
            return "both-blank"
        if sa == "" or sb == "":
            return "blank-in-one"
        return "agree" if sa == sb else "differ"
    na, nb = num(a), num(b)
    if na is None and nb is None:
        return "both-blank"
    if na is None or nb is None:
        return "blank-in-one"
    if isinstance(na, float) and isinstance(nb, float):
        ta, tr = TOL.get(field, TOL_DEFAULT)
        if field in DAY_FIELDS:
            ta = tol_abs if tol_abs is not None else TOL_DEFAULT[0]
            tr = tol_rel if tol_rel is not None else TOL_DEFAULT[1]
        return "agree" if abs(na - nb) <= max(ta, tr * max(abs(na), abs(nb))) else "differ"
    return "agree" if str(na).lower() == str(nb).lower() else "differ"


def run_compare(args):
    lines = ["EXTRACTION: COMPARISON OF THE TWO FILES",
             "E1 file: %s | E2 file: %s | made %s" % (os.path.basename(args.file1), os.path.basename(args.file2), datetime.date.today())]
    if not args.aggregate_only:
        lines.append("Contains both extractors' record-level values: give it to the adjudicator only after both have finished.")
    wb1, rep1, pat1 = load(args.file1)
    wb2, rep2, pat2 = load(args.file2)
    P1 = {txt(p["patient_id"]): p for p in pat1 if txt(p["patient_id"])}
    P2 = {txt(p["patient_id"]): p for p in pat2 if txt(p["patient_id"])}
    R1 = {txt(r["report_id"]): r for r in rep1 if txt(r["report_id"])}
    R2 = {txt(r["report_id"]): r for r in rep2 if txt(r["report_id"])}
    both = sorted(set(P1) & set(P2))
    only1, only2 = sorted(set(P1) - set(P2)), sorted(set(P2) - set(P1))
    lines.append("")
    lines.append("COVERAGE")
    lines.append("  Reports: %d in E1, %d in E2, %d in both" % (len(R1), len(R2), len(set(R1) & set(R2))))
    lines.append("  Patients: %d in E1, %d in E2, %d matched by patient_id, %d only in E1, %d only in E2" % (
        len(P1), len(P2), len(both), len(only1), len(only2)))
    cnt1, cnt2 = collections.Counter(p["report_id"] for p in pat1), collections.Counter(p["report_id"] for p in pat2)
    rep_diff = []
    for rid in sorted(set(R1) | set(R2)):
        a = txt(R1.get(rid, {}).get("n_children_individual")) if rid in R1 else None
        b = txt(R2.get(rid, {}).get("n_children_individual")) if rid in R2 else None
        if (a, b, cnt1.get(rid, 0), cnt2.get(rid, 0)) and (a != b or cnt1.get(rid, 0) != cnt2.get(rid, 0)):
            rep_diff.append((rid, a, b, cnt1.get(rid, 0), cnt2.get(rid, 0)))
    lines.append("  Reports where the child count or the number of patient rows differs: %d" % len(rep_diff))

    # ---- primary exposure and outcome
    ex_pairs, ex_blank, out_pairs, out_blank, via_pairs, via_blank = [], 0, [], 0, [], 0
    primary_rows = []
    day_state = collections.Counter()
    for pid in both:
        a, b = P1[pid], P2[pid]
        ca, cb = category(a), category(b)
        if ca is None or cb is None:
            ex_blank += 1
        else:
            ex_pairs.append((ca, cb))
        oa, ob = outcome(a), outcome(b)
        if oa is None or ob is None:
            out_blank += 1
        else:
            out_pairs.append((outcome3(oa), outcome3(ob)))
        va, vb = txt(a["viability"]), txt(b["viability"])
        if va == "" or vb == "":
            via_blank += 1
        else:
            via_pairs.append((va, vb))
        st = agree_value("idx_dur_days", a["idx_dur_days"], b["idx_dur_days"], args.tol_abs, args.tol_rel)
        day_state[st] += 1
        reasons = []
        if ca is not None and cb is not None and ca != cb:
            reasons.append("duration category %s vs %s" % (ca, cb))
        if st == "differ":
            reasons.append("days differ beyond tolerance")
        if oa is not None and ob is not None and outcome3(oa) != outcome3(ob):
            reasons.append("primary outcome %s vs %s" % (oa, ob))
        if va and vb and va != vb:
            reasons.append("viability %s vs %s" % (va, vb))
        if (ca is None) != (cb is None) or (oa is None) != (ob is None):
            reasons.append("primary variable blank in one file")
        if reasons:
            primary_rows.append((pid, a, b, "; ".join(reasons)))
    lines.append("")
    lines.append("PRIMARY EXPOSURE (index-episode duration, categories recomputed from the days)")
    lines.append("  matched patients %d; compared %d; left out because blank in one or both files %d" % (len(both), len(ex_pairs), ex_blank))
    ex_cats = CATS + ["NR"]
    lines.append("  " + fmt_k(kappa(ex_pairs, ex_cats)))
    lines.append("  contingency (rows E1, columns E2; %s):" % " / ".join(ex_cats))
    for c in ex_cats:
        lines.append("    %-6s %s" % (c, "  ".join("%4d" % sum(1 for a, b in ex_pairs if a == c and b == d) for d in ex_cats)))
    nd = day_state["agree"] + day_state["differ"]
    lines.append("  days (idx_dur_days): agree within tolerance (%.1f day or %.0f%%) in %d of %d where both gave a value or NR; blank in one file %d" % (
        args.tol_abs, 100 * args.tol_rel, day_state["agree"], nd, day_state["blank-in-one"]))
    lines.append("")
    lines.append("PRIMARY OUTCOME (derived from viability and procedure: Non-viable / Viable / other)")
    lines.append("  compared %d; left out because blank in one or both files %d" % (len(out_pairs), out_blank))
    lines.append("  " + fmt_k(kappa(out_pairs, ["Non-viable", "Viable", "Other"])))
    lines.append("  raw viability field (Viable / Partially non-viable / Non-viable / NR / Not applicable): " + fmt_k(kappa(via_pairs, LISTS["viability"])))
    if 0 < len(ex_pairs) < 30 or 0 < len(out_pairs) < 30:
        lines.append("  Note: fewer than 30 matched patients; kappa is very imprecise, read the disagreement list.")

    # ---- every coded field
    fields = [c for c in PATIENT_CODES if c not in NOT_COMPARED]
    by_field = []
    diff_rows = []
    for f in fields:
        n_cmp = n_ag = n_blank = 0
        for pid in both:
            a, b = P1[pid], P2[pid]
            st = agree_value(f, a[f], b[f], args.tol_abs, args.tol_rel)
            if st == "both-blank":
                continue
            if st == "blank-in-one":
                n_blank += 1
            n_cmp += 1
            if st == "agree":
                n_ag += 1
            else:
                diff_rows.append({"patient_id": pid, "report_id": txt(a["report_id"]), "field": f, "kind": st,
                                  "E1": txt(a[f]), "E2": txt(b[f])})
        by_field.append((f, n_cmp, n_ag, n_blank))
    tot_c = sum(x[1] for x in by_field)
    tot_a = sum(x[2] for x in by_field)
    lines.append("")
    lines.append("ALL CODED FIELDS (%d fields; free text, sources and derived columns not compared)" % len(fields))
    lines.append("  field values compared (cells filled in at least one file): %d; agree %d (%.1f%%); blank in one file %d; differ %d" % (
        tot_c, tot_a, 100.0 * tot_a / tot_c if tot_c else 0, sum(x[3] for x in by_field), tot_c - tot_a - sum(x[3] for x in by_field)))
    worst = sorted((x for x in by_field if x[1] >= 3 and x[2] < x[1]), key=lambda x: x[2] / x[1])[:8]
    if worst:
        lines.append("  fields with the lowest agreement (at least 3 values compared): " + ", ".join(
            "%s %d/%d" % (x[0], x[2], x[1]) for x in worst) if not args.aggregate_only else
            "  fields with differences: %d of %d" % (sum(1 for x in by_field if x[1] and x[2] < x[1]), len(fields)))
    lines.append("")
    lines.append("PACE (from EX_time_log)")
    for name, wb, pats, reps in (("E1", wb1, pat1, rep1), ("E2", wb2, pat2, rep2)):
        sub, iss = [], Issues()
        pace(wb, pats, reps, args.max_per_hour, iss, sub)
        lines.append(name + ": " + ("\n    ".join(sub) if sub else "no data"))
    text = "\n".join(lines) + "\n"
    print(text)
    if args.aggregate_only:
        print("(--aggregate-only: no files written)")
        return
    os.makedirs(args.out, exist_ok=True)
    with open(os.path.join(args.out, "ex_compare_summary.txt"), "w", encoding="utf-8") as f:
        f.write(text)
        if only1 or only2:
            f.write("\nPATIENTS ONLY IN ONE FILE\n  E1 only: %s\n  E2 only: %s\n" % (", ".join(only1) or "-", ", ".join(only2) or "-"))
    with open(os.path.join(args.out, "ex_patient_matching.csv"), "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["type", "report_id", "patient_id", "E1", "E2"])
        for pid in only1:
            w.writerow(["patient only in E1", txt(P1[pid]["report_id"]), pid, "present", ""])
        for pid in only2:
            w.writerow(["patient only in E2", txt(P2[pid]["report_id"]), pid, "", "present"])
        for rid, a, b, n1, n2 in rep_diff:
            w.writerow(["child count / patient rows differ", rid, "", "n_children_individual=%s; patient rows=%d" % (a if a is not None else "(no row)", n1),
                        "n_children_individual=%s; patient rows=%d" % (b if b is not None else "(no row)", n2)])
    with open(os.path.join(args.out, "ex_primary_disagreements.csv"), "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["patient_id", "report_id", "what", "E1_idx_dur_verbatim", "E1_idx_dur_days", "E1_viability", "E1_procedure", "E1_src_exposure", "E1_src_outcome",
                    "E2_idx_dur_verbatim", "E2_idx_dur_days", "E2_viability", "E2_procedure", "E2_src_exposure", "E2_src_outcome",
                    "discussion_agreed_idx_dur_days", "discussion_agreed_viability", "discussion_note", "adjudicator_idx_dur_days", "adjudicator_viability",
                    "adjudicator_note", "final_idx_dur_days", "final_viability"])
        for pid, a, b, what in primary_rows:
            w.writerow([pid, txt(a["report_id"]), what]
                       + [txt(a[c]) for c in ("idx_dur_verbatim", "idx_dur_days", "viability", "procedure", "src_exposure", "src_outcome")]
                       + [txt(b[c]) for c in ("idx_dur_verbatim", "idx_dur_days", "viability", "procedure", "src_exposure", "src_outcome")] + [""] * 7)
    with open(os.path.join(args.out, "ex_field_differences.csv"), "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["patient_id", "report_id", "field", "kind", "E1", "E2", "agreed_value", "adjudicator_value", "note"])
        for d in diff_rows:
            w.writerow([d["patient_id"], d["report_id"], d["field"], d["kind"], d["E1"], d["E2"], "", "", ""])
    with open(os.path.join(args.out, "ex_agreement_by_field.csv"), "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["field", "values_compared", "agree", "blank_in_one_file", "differ"])
        for fld, n_cmp, n_ag, n_bl in by_field:
            w.writerow([fld, n_cmp, n_ag, n_bl, n_cmp - n_ag - n_bl])
    print("Wrote ex_compare_summary.txt, ex_primary_disagreements.csv (%d), ex_field_differences.csv (%d), ex_patient_matching.csv, ex_agreement_by_field.csv in %s" % (
        len(primary_rows), len(diff_rows), args.out))


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    c = sub.add_parser("check", help="format, completeness and pace check of one extraction file")
    c.add_argument("file")
    c.add_argument("--extractor", choices=["E1", "E2"], help="expected value of the extractor column")
    c.add_argument("--max-per-hour", type=float, default=8.0, help="sessions above this many reports per hour are flagged (default 8)")
    c.set_defaults(func=run_check)
    m = sub.add_parser("compare", help="compare the two extractors' files")
    m.add_argument("file1", help="E1 (Jun Shu)")
    m.add_argument("file2", help="E2 (Hongqiang Bian)")
    m.add_argument("--out", default="search/extraction/compare", help="output folder")
    m.add_argument("--tol-abs", type=float, default=0.5, help="days: absolute tolerance (default 0.5)")
    m.add_argument("--tol-rel", type=float, default=0.10, help="days: relative tolerance (default 0.10)")
    m.add_argument("--max-per-hour", type=float, default=8.0)
    m.add_argument("--aggregate-only", action="store_true", help="print counts only; no ws ids, no output files")
    m.set_defaults(func=run_compare)
    args = ap.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
