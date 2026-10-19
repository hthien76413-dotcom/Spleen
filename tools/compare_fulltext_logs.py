#!/usr/bin/env python3
"""Check and compare the full-text screening logs (FT_log sheets of the extraction workbook).

Guide: "Full-text screening guide (draft)" v2026-10-17a, sections 3, 9, 10, 11 and 13 (decisions D3, D4, D5, D7).
Each reviewer keeps their own copy of the workbook (FT_R1.xlsx, FT_R2.xlsx); nobody looks at the other's file until both finish.

Two commands
  check    one file: format check (decision values, reason names, required fields, duplicate rows, records that should not
           have been read). Safe to run on one reviewer's file at the midpoint
           (after the first 50 reports) because it shows nothing about the other reviewer.
  compare  two files: agreement and kappa, list of conflicts for discussion/adjudication, differences in details,
           PRISMA-style counts per reviewer. Run it only when both reviewers have finished (or with --aggregate-only,
           which prints counts only: no WS ids, no per-record output files).

Usage (from the repo root; needs openpyxl):
  python3 -I tools/compare_fulltext_logs.py check FT_R1.xlsx [--reviewer R1] [--list PATH]
  python3 -I tools/compare_fulltext_logs.py compare FT_R1.xlsx FT_R2.xlsx --out DIR [--list PATH] [--aggregate-only]

Rules applied (from the guide)
  * Decisions: Include / Exclude / Pending IPD. Reasons: the names in the FT_log drop-down list; Exclude needs a reason,
    Pending IPD takes 'A aggregate data only', Include has none.
  * Any difference in the decision is a conflict to be settled (section 10: every decision decides whether a report enters
    extraction). Claude's reading: Include vs Pending IPD is a conflict too, although the guide names only Include vs Exclude
    and Pending IPD vs Exclude explicitly.
  * Both Exclude with different reasons is NOT a conflict (D3): the proposed reason is the first applicable one in the order
    L, D, E5, E4, E1, E2, E3, A, N, E6 (Y is never entered in FT_log; it is counted from the retrieval list).
    Exception (author's decision 18 Oct 2026: reason D needs the other screener to try as well): if exactly one of the two
    reasons is D, the pair is a conflict ('D not confirmed by both'); one screener got the full text, so it is settled in the
    discussion after both have finished. Both D = confirmed. The decisions still count as agreeing for kappa.
  * Both Include but different ws_basis / n_children / duplicate_of: listed as detail differences, settled when the
    extraction consensus is made.
  * No pace check: the FT_time_log sheet was dropped on 19 Oct 2026 (author's decision). A file that still has the sheet
    (the two calibration files do) is read for the FT_log sheet only; the time log is ignored.
Nothing in the outputs contains personal data; they do contain both reviewers' record-level decisions, so give them to the
adjudicator only after both have finished.
"""
import argparse
import collections
import csv
import datetime
import math
import os
import re
import statistics
import sys

from openpyxl import load_workbook

DECISIONS = ["Include", "Exclude", "Pending IPD"]
REASONS = [
    "L other language", "D no full text", "E1 non-human", "E2 not wandering spleen", "E3 adults only",
    "E4 no original patient data", "E5 not a formal article", "E6 other",
    "A aggregate data only", "N age not confirmed under 18",
]
ORDER = ["L", "D", "E5", "E4", "E1", "E2", "E3", "A", "N", "E6"]  # guide section 5 / D3
REASON_BY_CODE = {r.split()[0]: r for r in REASONS}
DEFAULT_LIST = "search/screening/rescreen/fulltext_retrieval_list.csv"
WS_RE = re.compile(r"^WS\d{4}$")
PREFILLED = ("ws_id", "citation", "language", "reviewer")  # columns that a pre-filled workbook already carries
FT_COLS = ["ws_id", "citation", "language", "reviewer", "date", "fulltext_obtained", "route", "routes_tried", "decision",
           "reason", "n_children_individual", "n_children_aggregate", "ws_basis", "duplicate_of", "note"]


# ----------------------------------------------------------------------------- reading
def txt(v):
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        v = int(v)
    return str(v).strip()


def to_date(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    s = txt(v)
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d"):
        try:
            return datetime.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def to_time(v):
    if isinstance(v, datetime.datetime):
        return v.time()
    if isinstance(v, datetime.time):
        return v
    s = txt(v).replace("：", ":")
    m = re.match(r"^(\d{1,2})[:.](\d{2})(?::\d{2})?$", s)
    if m and int(m.group(1)) < 24 and int(m.group(2)) < 60:
        return datetime.time(int(m.group(1)), int(m.group(2)))
    return None


def read_sheet(wb, name, codes, first_row=4):
    """Rows of a sheet as dicts keyed by the field codes in row 1; skips completely empty rows."""
    ws = wb[name]
    head = {txt(c.value): i for i, c in enumerate(ws[1]) if txt(c.value)}
    missing = [c for c in codes if c not in head]
    if missing:
        raise SystemExit("Sheet %s: the header row has no column(s) %s - use the template workbook." % (name, ", ".join(missing)))
    rows = []
    for r_idx, row in enumerate(ws.iter_rows(min_row=first_row, values_only=True), start=first_row):
        rec = {c: row[head[c]] if head[c] < len(row) else None for c in codes}
        if all(txt(v) == "" for v in rec.values()):
            continue
        rec["_row"] = r_idx
        rows.append(rec)
    return rows


def read_retrieval_list(path):
    if not path or not os.path.exists(path):
        return None
    out = {}
    with open(path, encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f):
            out[r["ws_id"]] = r
    return out


def canon_decision(s):
    s = txt(s)
    for d in DECISIONS:
        if s.lower() == d.lower():
            return d
    return None


def reason_code(s):
    """Code of a reason (first token, e.g. 'E1'); '' if blank. A wrongly typed reason keeps its raw text as the code."""
    s = txt(s)
    if not s:
        return ""
    first = s.split()[0].upper()
    return first if first in REASON_BY_CODE else s.lower()


# ----------------------------------------------------------------------------- check one file
class Issues:
    """Problems found in a file. Problems that can occur on many rows (kind=...) are printed as one line with the row numbers."""

    def __init__(self):
        self.items = []  # (severity, row, ws_id, message, kind)
        self.pending = 0  # pre-filled rows (WS id, citation, language, reviewer only) that have not been handled yet

    def add(self, sev, row, ws, msg, kind=None):
        self.items.append((sev, row, ws, msg, kind))

    def count(self, sev):
        return sum(1 for i in self.items if i[0] == sev)

    def lines(self):
        out, groups = [], collections.OrderedDict()
        for sev, row, ws, msg, kind in sorted(self.items, key=lambda i: (i[0] != "ERROR", i[1])):
            if kind:
                groups.setdefault((sev, kind, msg), []).append(row)
            else:
                out.append("  %s row %s %s: %s" % (sev, row or "-", ws, msg))
        for (sev, kind, msg), rows in groups.items():
            shown = ", ".join(str(r) for r in rows[:12]) + (", ..." if len(rows) > 12 else "")
            out.append("  %s %s (%d row%s: %s)" % (sev, msg, len(rows), "" if len(rows) == 1 else "s", shown))
        return out


def load_log(path, expected_reviewer, retrieval, issues):
    wb = load_workbook(path, data_only=True)
    if "FT_log" not in wb.sheetnames:
        raise SystemExit("%s has no FT_log sheet." % path)
    rows = read_sheet(wb, "FT_log", FT_COLS)
    by_id = collections.OrderedDict()
    for r in rows:
        ws = txt(r["ws_id"]).upper()
        r["ws_id"] = ws
        r["decision_c"] = canon_decision(r["decision"])
        r["reason_c"] = reason_code(r["reason"])
        n = r["_row"]
        if not WS_RE.match(ws):
            issues.add("ERROR", n, ws or "-", "ws_id missing or not of the form WS0123")
            continue
        if all(txt(r[c]) == "" for c in FT_COLS if c not in PREFILLED):
            issues.pending += 1  # a pre-filled row that the reviewer has not touched yet: not a record, not an error
            continue
        if ws in by_id:
            prev = by_id[ws]
            same = (prev["decision_c"], prev["reason_c"]) == (r["decision_c"], r["reason_c"])
            issues.add("ERROR", n, ws, "duplicate row (first at row %d, %s); the later row is used" % (
                prev["_row"], "same decision and reason" if same else "decision or reason differs"))
        by_id[ws] = r
        # retrieval list
        if retrieval is not None:
            info = retrieval.get(ws)
            if info is None:
                issues.add("ERROR", n, ws, "not on the retrieval list")
            else:
                cls = info["retrieval_class"]
                if cls.startswith("Y"):
                    issues.add("ERROR", n, ws, "published before 2000 (class Y): not to be read or logged")
                elif cls.startswith("L") and r["decision_c"] != "Exclude":
                    issues.add("WARN", n, ws, "listed as other language (class L) but logged as %s" % (txt(r["decision"]) or "blank"))
        # decision and reason
        dec = r["decision_c"]
        reason = txt(r["reason"])
        if dec is None:
            issues.add("ERROR", n, ws, "decision missing or not one of Include / Exclude / Pending IPD (%r)" % txt(r["decision"]))
        else:
            if dec == "Exclude":
                if not reason:
                    issues.add("ERROR", n, ws, "Exclude without a reason")
            elif dec == "Pending IPD":
                if reason and r["reason_c"] != "A":
                    issues.add("WARN", n, ws, "Pending IPD should carry reason 'A aggregate data only' (found %r)" % reason)
            elif dec == "Include" and reason:
                issues.add("WARN", n, ws, "Include with a reason (%r); leave the reason blank" % reason)
        if reason and reason not in REASONS:
            issues.add("ERROR", n, ws, "reason %r is not exactly a name in the drop-down list" % reason)
        # consistency with full text, details
        obtained = txt(r["fulltext_obtained"]).lower()
        if dec in ("Include", "Pending IPD") and obtained == "no":
            issues.add("ERROR", n, ws, "%s but full text obtained = No" % dec)
        if r["reason_c"] == "D":
            if obtained == "yes":
                issues.add("WARN", n, ws, "reason D (no full text) but full text obtained = Yes")
            if not txt(r["routes_tried"]):
                issues.add("WARN", n, ws, "reason D without routes_tried")
        if r["reason_c"] == "E6" and not txt(r["note"]):
            issues.add("WARN", n, ws, "E6 without a note")
        if dec == "Include":
            if not txt(r["ws_basis"]):
                issues.add("WARN", n, ws, "Include without ws_basis", kind="no-ws-basis")
            if txt(r["n_children_individual"]) == "":
                issues.add("WARN", n, ws, "Include without n_children_individual", kind="no-n-ind")
            if txt(r["n_children_aggregate"]) == "":
                issues.add("WARN", n, ws, "Include without n_children_aggregate (write 0 if there are none)", kind="no-n-agg-inc")
        if dec == "Pending IPD" and txt(r["n_children_aggregate"]) == "":
            issues.add("WARN", n, ws, "Pending IPD without n_children_aggregate", kind="no-n-agg")
        for col in ("n_children_individual", "n_children_aggregate"):
            v = txt(r[col])
            if v != "" and not re.match(r"^\d+$", v):
                issues.add("ERROR", n, ws, "%s is not a whole number (%r)" % (col, v))
        dup = txt(r["duplicate_of"]).upper()
        if dup and (not WS_RE.match(dup) or dup == ws):
            issues.add("WARN", n, ws, "duplicate_of %r is not another WS id" % txt(r["duplicate_of"]))
        # reviewer and date
        rv = txt(r["reviewer"])
        if expected_reviewer and rv and rv != expected_reviewer:
            issues.add("WARN", n, ws, "reviewer column does not say %s" % expected_reviewer, kind="reviewer-mismatch")
        elif not rv:
            issues.add("WARN", n, ws, "reviewer column empty", kind="reviewer-empty")
        d = to_date(r["date"])
        r["date_c"] = d
        if d is None:
            issues.add("WARN", n, ws, "date missing or not a date", kind="date-missing")
        elif d > datetime.date.today():
            issues.add("WARN", n, ws, "date is in the future", kind="date-future")
    return by_id, wb


def summarise_log(by_id, retrieval, lines):
    dec = collections.Counter(r["decision_c"] or "invalid" for r in by_id.values())
    lines.append("Decisions: " + ", ".join("%s %d" % (k, dec[k]) for k in DECISIONS + (["invalid"] if dec.get("invalid") else [])))
    reasons = collections.Counter(r["reason_c"] for r in by_id.values() if r["decision_c"] == "Exclude")
    if reasons:
        lines.append("Exclude by reason: " + ", ".join("%s %d" % (k, reasons[k]) for k in ORDER if k in reasons)
                     + ("; blank or invalid " + str(sum(v for k, v in reasons.items() if k not in ORDER)) if any(k not in ORDER for k in reasons) else ""))
    if retrieval is not None:
        to_read = [w for w, i in retrieval.items() if i["retrieval_class"] == "retrieve"]
        done = [w for w in to_read if w in by_id]
        lines.append("Retrieval list: %d reports to retrieve; %d handled in this file, %d not yet." % (
            len(to_read), len(done), len(to_read) - len(done)))


def run_check(args):
    retrieval = read_retrieval_list(args.list)
    lines = []
    if retrieval is None:
        lines.append("Retrieval list not found (%s): list checks skipped." % args.list)
    issues = Issues()
    by_id, wb = load_log(args.file, args.reviewer, retrieval, issues)
    lines.insert(0, "File: %s (%d records in FT_log)" % (os.path.basename(args.file), len(by_id)))
    summarise_log(by_id, retrieval, lines)
    if issues.pending:
        lines.append("Pre-filled rows not handled yet: %d (not counted as records)" % issues.pending)
    lines.append("Format check: %d ERROR(S), %d WARNING(S)" % (issues.count("ERROR"), issues.count("WARN")))
    lines.extend(issues.lines())
    print("\n".join(lines))


# ----------------------------------------------------------------------------- compare two files
def kappa(pairs, cats):
    n = len(pairs)
    if n == 0:
        return None
    po = sum(1 for a, b in pairs if a == b) / n
    pa = {c: sum(1 for a, _ in pairs if a == c) / n for c in cats}
    pb = {c: sum(1 for _, b in pairs if b == c) / n for c in cats}
    pe = sum(pa[c] * pb[c] for c in cats)
    if abs(1 - pe) < 1e-12:
        return {"po": po, "kappa": None, "lo": None, "hi": None}
    k = (po - pe) / (1 - pe)
    se = math.sqrt(po * (1 - po) / (n * (1 - pe) ** 2))  # rough standard error, as used for the calibration
    return {"po": po, "kappa": k, "lo": max(-1.0, k - 1.96 * se), "hi": min(1.0, k + 1.96 * se)}  # rough interval, kept inside [-1, 1]


def fmt_k(k):
    if k is None:
        return "n/a"
    if k["kappa"] is None:
        return "agreement %.1f%%; kappa not defined (no variation)" % (100 * k["po"])
    return "agreement %.1f%%; kappa %.2f (rough 95%% CI %.2f to %.2f)" % (100 * k["po"], k["kappa"], k["lo"], k["hi"])


def proposed_reason(c1, c2):
    cands = [c for c in (c1, c2) if c in ORDER]
    if not cands:
        return ""
    return REASON_BY_CODE[min(cands, key=ORDER.index)]


def run_compare(args):
    retrieval = read_retrieval_list(args.list)
    lines = ["FULL-TEXT SCREENING: COMPARISON OF THE TWO FT_log FILES",
             "R1 file: %s | R2 file: %s | made %s" % (os.path.basename(args.file1), os.path.basename(args.file2), datetime.date.today())]
    if not args.aggregate_only:
        lines.append("Contains both reviewers' record-level decisions: give it to the adjudicator only after both have finished.")
    if retrieval is None:
        lines.append("Retrieval list not found (%s): list checks skipped, no titles in the outputs." % args.list)
    iss1, iss2 = Issues(), Issues()
    r1, wb1 = load_log(args.file1, "R1", retrieval, iss1)
    r2, wb2 = load_log(args.file2, "R2", retrieval, iss2)
    for name, iss in (("R1", iss1), ("R2", iss2)):
        lines.append("Format check %s: %d error(s), %d warning(s)%s%s" % (
            name, iss.count("ERROR"), iss.count("WARN"), "  -> run the 'check' command on this file" if iss.items else "",
            "; %d pre-filled rows not handled yet" % iss.pending if iss.pending else ""))
    ids = sorted(set(r1) | set(r2))
    both = [w for w in ids if w in r1 and w in r2]
    only1 = [w for w in ids if w in r1 and w not in r2]
    only2 = [w for w in ids if w in r2 and w not in r1]
    lines += ["", "Records: %d in R1, %d in R2, %d in both, %d only in R1, %d only in R2" % (len(r1), len(r2), len(both), len(only1), len(only2))]
    if retrieval is not None:
        to_read = [w for w, i in retrieval.items() if i["retrieval_class"] == "retrieve"]
        neither = [w for w in to_read if w not in r1 and w not in r2]
        lines.append("Retrieval list: %d to retrieve; handled by both %d, by one only %d, by neither %d" % (
            len(to_read), sum(1 for w in to_read if w in r1 and w in r2),
            sum(1 for w in to_read if (w in r1) != (w in r2)), len(neither)))

    rows, pairs, detail = [], [], []
    agg_blank_as_zero = 0
    cat_count = collections.Counter()
    reason_pairs = collections.Counter()
    for w in both:
        a, b = r1[w], r2[w]
        d1, d2 = a["decision_c"], b["decision_c"]
        if d1 is None or d2 is None:
            cat = "invalid decision in one file"
        elif d1 == d2 == "Exclude" and (a["reason_c"] == "D") != (b["reason_c"] == "D"):
            cat = "conflict: D not confirmed by both"
            pairs.append((d1, d2))
        elif d1 == d2:
            cat = "agree " + d1
            pairs.append((d1, d2))
            if d1 == "Exclude":
                reason_pairs[(a["reason_c"], b["reason_c"])] += 1
        else:
            cat = "conflict: " + " vs ".join(sorted((d1, d2), key=DECISIONS.index))
            pairs.append((d1, d2))
        cat_count[cat] += 1
        rec = {"ws_id": w, "category": cat, "R1_decision": d1 or txt(a["decision"]), "R1_reason": txt(a["reason"]),
               "R2_decision": d2 or txt(b["decision"]), "R2_reason": txt(b["reason"]),
               "proposed_reason_if_both_exclude": proposed_reason(a["reason_c"], b["reason_c"]) if cat == "agree Exclude" and a["reason_c"] != b["reason_c"] else ""}
        rows.append((rec, a, b, cat))
        if cat in ("agree Include", "agree Pending IPD"):
            diffs = []
            for col in (("ws_basis", "n_children_individual", "n_children_aggregate", "duplicate_of") if cat == "agree Include"
                        else ("n_children_aggregate",)):
                v1, v2 = txt(a[col]).upper(), txt(b[col]).upper()
                if cat == "agree Include" and col == "n_children_aggregate":
                    # for an included report a blank count means 'none' (check asks for 0); blank and 0 are the same here
                    if v1 == "" or v2 == "":
                        agg_blank_as_zero += 1
                    v1, v2 = v1 or "0", v2 or "0"
                if v1 != v2:
                    diffs.append((col, txt(a[col]), txt(b[col])))
            for col, v1, v2 in diffs:
                detail.append({"ws_id": w, "type": "both " + cat.split()[1] + ", different " + col, "R1": v1, "R2": v2})
        if cat == "agree Exclude" and a["reason_c"] != b["reason_c"]:
            detail.append({"ws_id": w, "type": "both Exclude, different reason (proposed: first applicable in the D3 order)",
                           "R1": txt(a["reason"]), "R2": txt(b["reason"])})

    lines.append("")
    lines.append("Categories (records in both files): " + ", ".join("%s %d" % (k, v) for k, v in sorted(cat_count.items())))
    n_conf = sum(v for k, v in cat_count.items() if k.startswith("conflict"))
    lines.append("Conflicts to settle: %d of %d (%.1f%%)" % (n_conf, len(both), 100.0 * n_conf / len(both) if both else 0))
    lines.append("3 categories (Include / Exclude / Pending IPD): " + fmt_k(kappa(pairs, DECISIONS)))
    bin_pairs = [("Exclude" if a == "Exclude" else "Continue", "Exclude" if b == "Exclude" else "Continue") for a, b in pairs]
    lines.append("2 categories (Exclude vs Include or Pending IPD): " + fmt_k(kappa(bin_pairs, ["Exclude", "Continue"])))
    if len(pairs) < 50:
        lines.append("Note: fewer than 50 records compared; the guide (D5) says to report raw agreement only for the 20-report calibration, "
                     "kappa is not meaningful at this size.")
    lines.append("Contingency (rows R1, columns R2; Include / Exclude / Pending IPD):")
    for d in DECISIONS:
        lines.append("  %-12s %s" % (d, "  ".join("%4d" % sum(1 for a, b in pairs if a == d and b == e) for e in DECISIONS)))
    n_ex = sum(reason_pairs.values())
    if n_ex:
        same = sum(v for (x, y), v in reason_pairs.items() if x == y)
        lines.append("Both Exclude: %d; same reason %d (%.1f%%), different reason %d" % (n_ex, same, 100.0 * same / n_ex, n_ex - same))
        diff_pairs = sorted(((x, y), v) for (x, y), v in reason_pairs.items() if x != y)
        if diff_pairs and not args.aggregate_only:
            lines.append("  different reasons (R1 / R2): " + ", ".join("%s/%s %d" % (x, y, v) for (x, y), v in diff_pairs))
        elif diff_pairs:
            lines.append("  pairs of different reasons: %d kinds" % len(diff_pairs))
    lines.append("Detail differences (both Include / both Pending IPD / both Exclude): %d" % len(detail))
    if agg_blank_as_zero:
        lines.append("  n_children_aggregate: a blank counted as 0 on %d both-Include record(s), so blank vs 0 is not listed as a difference" % agg_blank_as_zero)
    for name, by in (("R1", r1), ("R2", r2)):
        dec = collections.Counter(r["decision_c"] or "invalid" for r in by.values())
        rs = collections.Counter(r["reason_c"] for r in by.values() if r["decision_c"] == "Exclude")
        lines.append("%s: %s | Exclude reasons: %s" % (
            name, ", ".join("%s %d" % (d, dec[d]) for d in DECISIONS),
            ", ".join("%s %d" % (c, rs[c]) for c in ORDER if c in rs) or "none"))
    lines.append("")

    text = "\n".join(lines) + "\n"
    print(text)
    if args.aggregate_only:
        print("(--aggregate-only: no files written)")
        return
    os.makedirs(args.out, exist_ok=True)
    meta = {}
    if retrieval is not None:
        meta = {w: retrieval[w] for w in ids if w in retrieval}

    def info(w, col):
        return meta.get(w, {}).get(col, "")

    with open(os.path.join(args.out, "ft_compare_summary.txt"), "w", encoding="utf-8") as f:
        f.write(text)
        f.write("\nFORMAT ISSUES (R1)\n" + "\n".join(iss1.lines()) + "\n")
        f.write("\nFORMAT ISSUES (R2)\n" + "\n".join(iss2.lines()) + "\n")
        if only1 or only2:
            f.write("\nONLY IN ONE FILE\n  R1 only: %s\n  R2 only: %s\n" % (", ".join(only1) or "-", ", ".join(only2) or "-"))
    base = ["ws_id", "year", "language", "title", "category"]
    cols = ["decision", "reason", "ws_basis", "n_children_individual", "n_children_aggregate", "duplicate_of", "note"]
    with open(os.path.join(args.out, "ft_compare_all.csv"), "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(base + ["R1_" + c for c in cols] + ["R2_" + c for c in cols] + ["proposed_reason_if_both_exclude"])
        for rec, a, b, cat in rows:
            wid = rec["ws_id"]
            w.writerow([wid, info(wid, "year"), info(wid, "language_group"), info(wid, "title"), cat]
                       + [txt(a[c]) for c in cols] + [txt(b[c]) for c in cols] + [rec["proposed_reason_if_both_exclude"]])
        for wid in only1 + only2:
            src, by = ("R1", r1) if wid in r1 else ("R2", r2)
            x = by[wid]
            vals = [txt(x[c]) for c in cols]
            empty = [""] * len(cols)
            w.writerow([wid, info(wid, "year"), info(wid, "language_group"), info(wid, "title"), "only in " + src]
                       + (vals + empty if src == "R1" else empty + vals) + [""])
    conflicts = [(rec, a, b) for rec, a, b, cat in rows if cat.startswith("conflict")]
    with open(os.path.join(args.out, "ft_conflicts_for_discussion.csv"), "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(base + ["R1_" + c for c in cols] + ["R2_" + c for c in cols]
                   + ["discussion_agreed_decision", "discussion_agreed_reason", "discussion_note",
                      "adjudicator_decision", "adjudicator_reason", "adjudicator_note", "final_decision", "final_reason"])
        for rec, a, b in conflicts:
            wid = rec["ws_id"]
            w.writerow([wid, info(wid, "year"), info(wid, "language_group"), info(wid, "title"), rec["category"]]
                       + [txt(a[c]) for c in cols] + [txt(b[c]) for c in cols] + [""] * 8)
    with open(os.path.join(args.out, "ft_detail_differences.csv"), "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["ws_id", "year", "title", "type", "R1", "R2"])
        for d in detail:
            w.writerow([d["ws_id"], info(d["ws_id"], "year"), info(d["ws_id"], "title"), d["type"], d["R1"], d["R2"]])
    print("Wrote ft_compare_summary.txt, ft_compare_all.csv, ft_conflicts_for_discussion.csv (%d), ft_detail_differences.csv (%d) in %s" % (
        len(conflicts), len(detail), args.out))


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    c = sub.add_parser("check", help="format check of one FT_log file")
    c.add_argument("file")
    c.add_argument("--reviewer", choices=["R1", "R2"], help="expected value of the reviewer column")
    c.add_argument("--list", default=DEFAULT_LIST, help="retrieval list CSV (default %s)" % DEFAULT_LIST)
    c.set_defaults(func=run_check)
    m = sub.add_parser("compare", help="compare the two reviewers' FT_log files")
    m.add_argument("file1", help="R1 (Jun Shu)")
    m.add_argument("file2", help="R2 (Hongqiang Bian)")
    m.add_argument("--out", default="search/screening/fulltext/compare", help="output folder")
    m.add_argument("--list", default=DEFAULT_LIST)
    m.add_argument("--aggregate-only", action="store_true", help="print counts only; no ws ids, no output files")
    m.set_defaults(func=run_compare)
    args = ap.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
