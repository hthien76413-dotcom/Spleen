#!/usr/bin/env python3
"""Check and compare the two risk-of-bias workbooks (tools/make_rob_workbooks.py).

  check   FILE                      one file: values, unfilled rows, duplicate ids.
  compare FILE_R1 FILE_R2 --out DIR both files, run only when both reviewers have finished: agreement per item (raw and kappa),
                                    the disagreement list for the discussion (rob_disagreements.csv), and a draft rob.csv
                                    (report_id, adequate) for the analysis, written ONLY for reports where the two agree on
                                    the three items that define 'adequate'; the others stay blank until settled.
The derived `adequate` is recomputed here (not read from the file): Yes if exposure, outcome and reporting are all Yes; No if all
three are filled and any is not Yes; blank otherwise.
Usage: python3 -I tools/compare_rob.py check rob_R1.xlsx | compare rob_R1.xlsx rob_R2.xlsx --out DIR      (needs openpyxl)
"""
import argparse, collections, csv, os, sys
sys.dont_write_bytecode = True
from openpyxl import load_workbook
ITEMS = ["selection", "exposure", "outcome", "alternative", "followup", "reporting"]
OK = {"Yes", "No", "Unclear"}


def load(path):
    ws = load_workbook(path, data_only=True)["RoB"]
    h = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    rows, seen, dup = {}, set(), []
    for r in range(2, ws.max_row + 1):
        rid = str(ws.cell(r, h["report_id"]).value or "").strip()
        if not rid: continue
        if rid in seen: dup.append(rid)
        seen.add(rid)
        rows[rid] = {k: str(ws.cell(r, h[k]).value or "").strip() for k in ITEMS + ["note", "title"]}
    return rows, dup


def adequate(r):
    v = [r["exposure"], r["outcome"], r["reporting"]]
    if any(x == "" for x in v): return ""
    return "Yes" if all(x == "Yes" for x in v) else "No"


def kappa(a, b):
    n = len(a)
    if n == 0: return float("nan")
    po = sum(x == y for x, y in zip(a, b)) / n
    ca, cb = collections.Counter(a), collections.Counter(b)
    pe = sum(ca[k] * cb[k] for k in set(ca) | set(cb)) / n ** 2
    return float("nan") if pe == 1 else (po - pe) / (1 - pe)


def check(path):
    rows, dup = load(path)
    bad = [(rid, k, r[k]) for rid, r in rows.items() for k in ITEMS if r[k] not in OK | {""}]
    unf = [rid for rid, r in rows.items() if any(r[k] == "" for k in ITEMS)]
    print("File: %s, %d reports; fully filled %d; with a blank item %d" % (os.path.basename(path), len(rows), len(rows) - len(unf), len(unf)))
    print("Adequate (recomputed): Yes %d, No %d, not yet %d" % tuple(sum(adequate(r) == v for r in rows.values()) for v in ("Yes", "No", "")))
    print("Format check: %d ERROR(S), %d WARNING(S)" % (len(bad) + len(dup), 0))
    for b in bad: print("  ERROR", b)
    for d in dup: print("  ERROR duplicate report_id", d)


def compare(p1, p2, out):
    a, _ = load(p1); b, _ = load(p2); ids = sorted(set(a) & set(b))
    print("Reports in both: %d (only R1 %d, only R2 %d)" % (len(ids), len(set(a) - set(b)), len(set(b) - set(a))))
    os.makedirs(out, exist_ok=True)
    dis = []
    for k in ITEMS:
        pairs = [(a[i][k], b[i][k]) for i in ids if a[i][k] and b[i][k]]
        agree = sum(x == y for x, y in pairs)
        print("%-12s n=%d agree %.1f%% kappa %.2f" % (k, len(pairs), 100 * agree / len(pairs) if pairs else float("nan"), kappa([x for x, _ in pairs], [y for _, y in pairs])))
    for i in ids:
        diff = [k for k in ITEMS if a[i][k] != b[i][k]]
        if diff: dis.append([i, a[i]["title"], ";".join(diff)] + [a[i][k] for k in ITEMS] + [b[i][k] for k in ITEMS] + [a[i]["note"], b[i]["note"], "", "", ""])
    with open(os.path.join(out, "rob_disagreements.csv"), "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f); w.writerow(["report_id", "title", "items_differ"] + ["R1_" + k for k in ITEMS] + ["R2_" + k for k in ITEMS] + ["R1_note", "R2_note", "discussion", "adjudicator", "final"])
        w.writerows(dis)
    with open(os.path.join(out, "rob_draft.csv"), "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f); w.writerow(["report_id", "adequate"]); n = 0
        for i in ids:
            if all(a[i][k] == b[i][k] for k in ("exposure", "outcome", "reporting")) and adequate(a[i]):
                w.writerow([i, adequate(a[i])]); n += 1
    print("Reports with a disagreement on any item: %d; draft rob.csv has %d reports (agreed on the three defining items)" % (len(dis), n))


if __name__ == "__main__":
    ap = argparse.ArgumentParser(); sub = ap.add_subparsers(dest="cmd", required=True)
    c = sub.add_parser("check"); c.add_argument("file")
    d = sub.add_parser("compare"); d.add_argument("f1"); d.add_argument("f2"); d.add_argument("--out", required=True)
    x = ap.parse_args()
    check(x.file) if x.cmd == "check" else compare(x.f1, x.f2, x.out)
