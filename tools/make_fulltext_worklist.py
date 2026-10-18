#!/usr/bin/env python3
"""Retrieval worklist for the full-text stage (reports to retrieve that have an abstract in the database record).

Inputs (all in the repository):
  search/screening/rescreen/fulltext_retrieval_list.csv   retrieval_class 'retrieve' and has_abstract 'yes'
  search/dedup/records_deduplicated.ris                   citation details and links
  search/screening/fulltext/calibration_20_list.csv       only to mark the 20 calibration reports (never the key file)
  search/screening/fulltext/pmc_lookup_2026-10-18.csv     PMID -> PMCID pairs from the PubMed ID converter

The worklist carries NO screening decision, no stratum flag and no 'source' column (the source would reveal how a
record reached the full-text stage, e.g. the E4 safety net). Retrieving a PDF creates no decision.

Usage: python3 -I tools/make_fulltext_worklist.py OUT.xlsx [--csv OUT.csv] [--abstract yes|no|any]
Requires openpyxl.
"""
import argparse
import csv
import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
RETRIEVAL = ROOT / "search/screening/rescreen/fulltext_retrieval_list.csv"
RIS = ROOT / "search/dedup/records_deduplicated.ris"
CAL = ROOT / "search/screening/fulltext/calibration_20_list.csv"
PMC = ROOT / "search/screening/fulltext/pmc_lookup_2026-10-18.csv"

ROUTES = [
    "Publisher or open access", "PubMed Central", "Library subscription", "Interlibrary loan or document delivery",
    "CNKI or Wanfang", "Contact author", "Other", "Not obtained",
]
GROUPS = {
    1: "1 PMC record exists",
    2: "2 English, DOI or PubMed link",
    3: "3 Chinese",
    4: "4 English, no DOI and no PubMed link",
}


def read_ris(path):
    recs, cur = [], None
    for line in open(path, encoding="utf-8-sig"):
        m = re.match(r"^([A-Z][A-Z0-9])  - ?(.*)$", line.rstrip("\n"))
        if not m:
            continue
        tag, val = m.groups()
        if tag == "TY":
            cur = {}
            recs.append(cur)
        cur.setdefault(tag, []).append(val)
    return {r["ID"][0]: r for r in recs}


def first_author(rec):
    au = rec.get("AU", [""])[0]
    if not au:
        return ""
    return au.split(",")[0].strip() if "," in au else au.strip()


def pages(rec):
    vol, iss = rec.get("VL", [""])[0], rec.get("IS", [""])[0]
    sp, ep = rec.get("SP", [""])[0], rec.get("EP", [""])[0]
    out = vol + (f"({iss})" if iss else "")
    if sp:
        out += (":" if out else "") + sp + (f"-{ep}" if ep else "")
    return out


def urls(rec):
    allv = rec.get("UR", []) + rec.get("N1", [])
    text = " ".join(allv)
    found = re.findall(r"https?://[^\s;]+", text)
    pubmed = next((u for u in found if "pubmed.ncbi.nlm.nih.gov" in u), "")
    cnki = next((u for u in found if "cnki.net" in u), "")
    wf = next((u for u in found if "wanfangdata.com.cn" in u), "")
    scopus = next((u for u in found if "scopus.com" in u), "")
    m = re.search(r"PubMed (\d+)", text)
    if not pubmed and m:
        pubmed = f"https://pubmed.ncbi.nlm.nih.gov/{m.group(1)}/"
    pmid = re.search(r"pubmed\.ncbi\.nlm\.nih\.gov/(\d+)", pubmed)
    return pubmed, (pmid.group(1) if pmid else ""), cnki, wf, scopus


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("out")
    ap.add_argument("--csv")
    ap.add_argument("--abstract", default="yes", choices=["yes", "no", "any"])
    args = ap.parse_args()

    recs = read_ris(RIS)
    cal = {r["ws_id"] for r in csv.DictReader(open(CAL, encoding="utf-8-sig"))}
    pmc = {}
    if PMC.exists():
        for r in csv.DictReader(open(PMC, encoding="utf-8-sig")):
            pmc[r["pmid"]] = r["pmcid"]
    rows = []
    for r in csv.DictReader(open(RETRIEVAL, encoding="utf-8-sig")):
        if r["retrieval_class"] != "retrieve":
            continue
        if args.abstract != "any" and r["has_abstract"] != args.abstract:
            continue
        rows.append(r)

    items = []
    for r in rows:
        w = r["ws_id"]
        rec = recs[w]
        pubmed, pmid, cnki, wf, scopus = urls(rec)
        doi = rec.get("DO", [""])[0]
        pmcid = pmc.get(pmid, "") if pmid else ""
        if pmcid:
            grp = 1
        elif r["language_group"] == "Chinese":
            grp = 3
        elif doi or pubmed:
            grp = 2
        else:
            grp = 4
        items.append({
            "ws_id": w, "year": r["year"], "language": r["language_group"], "record_type": r["record_type"],
            "first_author": first_author(rec), "journal": rec.get("T2", [""])[0], "vol_pages": pages(rec),
            "title": r["title"], "doi": doi, "doi_url": f"https://doi.org/{doi}" if doi else "",
            "pubmed_url": pubmed, "pmcid": pmcid,
            "pmc_url": f"https://pmc.ncbi.nlm.nih.gov/articles/{pmcid}/" if pmcid else "",
            "cnki_url": cnki, "wanfang_url": wf, "scopus_url": scopus,
            "group": grp, "calibration": "Y" if w in cal else "",
            "target_file": f"{w}.pdf",
        })
    items.sort(key=lambda d: (d["group"], -int(d["year"] or 0), d["ws_id"]))
    for i, d in enumerate(items, 1):
        d["seq"] = i

    COLS = [
        ("seq", "No.", 6), ("ws_id", "WS id", 9), ("year", "Year", 7), ("language", "Language", 10),
        ("record_type", "Type", 7), ("first_author", "First author", 16), ("journal", "Journal", 28),
        ("vol_pages", "Vol(issue):pages", 16), ("title", "Title", 60), ("doi", "DOI", 28), ("doi_url", "DOI link", 14),
        ("pubmed_url", "PubMed link", 14), ("pmcid", "PMC id", 13), ("pmc_url", "PMC link", 14),
        ("cnki_url", "CNKI link", 12), ("wanfang_url", "Wanfang link", 12), ("scopus_url", "Scopus link", 12),
        ("group", "Route group", 24), ("calibration", "Calibration report", 12), ("target_file", "Save PDF as", 13),
    ]
    TRACK = [
        ("retrieved_by", "Retrieved by", 12), ("retrieved_date", "Date (yyyy-mm-dd)", 14),
        ("obtained", "Full text obtained", 12), ("route", "Route", 26), ("routes_tried", "Routes tried (if not obtained)", 34),
        ("retry_by", "Second try by", 12), ("retry_date", "Second try date", 14), ("retry_obtained", "Second try obtained", 12),
        ("retry_routes_tried", "Second try: routes tried", 34), ("note", "Note", 30),
    ]

    wb = Workbook()
    ws = wb.active
    ws.title = "Worklist"
    head_fill = PatternFill("solid", fgColor="D9E1F2")
    track_fill = PatternFill("solid", fgColor="FFF2CC")
    thin = Side(style="thin", color="BBBBBB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    allcols = COLS + TRACK
    for c, (_, label, width) in enumerate(allcols, 1):
        cell = ws.cell(row=1, column=c, value=label)
        cell.font = Font(bold=True)
        cell.fill = head_fill if c <= len(COLS) else track_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = border
        ws.column_dimensions[get_column_letter(c)].width = width
    for r, d in enumerate(items, 2):
        for c, (key, _, _) in enumerate(COLS, 1):
            v = d[key]
            if key == "group":
                v = GROUPS[v]
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=(key == "title"))
            if key.endswith("_url") and v:
                cell.hyperlink = v
                cell.value = "link"
                cell.font = Font(color="0563C1", underline="single")
        for c in range(len(COLS) + 1, len(allcols) + 1):
            ws.cell(row=r, column=c).border = border
    n = len(items)
    last = n + 1
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(allcols))}{last}"

    idx = {k: i for i, (k, _, _) in enumerate(allcols, 1)}
    dv_yes = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    dv_route = DataValidation(type="list", formula1='"' + ",".join(ROUTES) + '"', allow_blank=True)
    dv_by = DataValidation(type="list", formula1='"舒俊,卞红强,彭飞,其他"', allow_blank=True)
    for dv in (dv_yes, dv_route, dv_by):
        ws.add_data_validation(dv)
    for key in ("obtained", "retry_obtained"):
        L = get_column_letter(idx[key])
        dv_yes.add(f"{L}2:{L}{last}")
    L = get_column_letter(idx["route"])
    dv_route.add(f"{L}2:{L}{last}")
    for key in ("retrieved_by", "retry_by"):
        L = get_column_letter(idx[key])
        dv_by.add(f"{L}2:{L}{last}")

    # summary sheet with live counts
    sm = wb.create_sheet("Summary")
    g = get_column_letter(idx["group"])
    ob = get_column_letter(idx["obtained"])
    rob = get_column_letter(idx["retry_obtained"])
    lines = [("Reports on the worklist", f"=COUNTA(Worklist!B2:B{last})")]
    for k, label in GROUPS.items():
        lines.append((f"Route group {label}", f'=COUNTIF(Worklist!{g}2:{g}{last},"{label}")'))
    lines += [
        ("Full text obtained (first try)", f'=COUNTIF(Worklist!{ob}2:{ob}{last},"Yes")'),
        ("Not obtained at first try", f'=COUNTIF(Worklist!{ob}2:{ob}{last},"No")'),
        ("Second try obtained", f'=COUNTIF(Worklist!{rob}2:{rob}{last},"Yes")'),
        ("Second try not obtained (candidates for reason D if both logged all routes)",
         f'=COUNTIF(Worklist!{rob}2:{rob}{last},"No")'),
        ("Not yet tried", f'=COUNTA(Worklist!B2:B{last})-COUNTA(Worklist!{ob}2:{ob}{last})'),
    ]
    sm.column_dimensions["A"].width = 70
    sm.column_dimensions["B"].width = 12
    for i, (a, b) in enumerate(lines, 1):
        sm.cell(row=i, column=1, value=a)
        sm.cell(row=i, column=2, value=b)

    rd = wb.create_sheet("README", 0)
    text = [
        "全文取回清单 / Full-text retrieval worklist (draft, generated 2026-10-18)",
        "",
        f"内容：{n} 篇待取全文的报告（retrieval class = retrieve，数据库记录里有摘要）。不含 Y（2000 年以前，不取）和 L（其他语言，不取）。",
        "这份清单只用于取全文，不含任何筛选判断，也不含来源列或分层标记。取回 PDF 不产生任何决定。",
        "",
        "怎么用（Claude 的建议，尚未经三人确认）：",
        "1. PDF 按 WS 编号命名（列 'Save PDF as'），保存在自己的电脑或网盘，不放进仓库。",
        "2. 取全文途径按全文筛选指南 4.2 的顺序：出版社或开放获取 / PMC → 单位图书馆订阅 → 图书馆文献传递 → CNKI、万方（中文）→ 联系通讯作者。只用合法途径。",
        "3. 每篇在黄色列记：谁取的、日期、是否取到、最终途径。没取到：写 'Routes tried'，试过什么、什么时候试的。",
        "4. 第一次没取到的，由另一位筛选员再试一遍（第二次尝试列）。只有两人都试过所有途径、都没取到，才在 FT_log 里记 D。",
        "5. 20 篇校准报告（列 'Calibration report' = Y）：两位筛选员各自独立取，不共用这里的记录，也不互相交换 PDF 之外的任何信息。",
        "6. 'Route group' 只是查找的先后提示：1 = PubMed 的 ID 转换工具显示在 PMC 有记录（2026-10-18 查询；链接要点开确认是同一篇，可能有延迟开放或只是作者稿）；2 = 英文，有 DOI 或 PubMed 链接；3 = 中文；4 = 英文且无 DOI、无 PubMed 链接。",
        "7. 链接来自数据库记录，没有逐条点开核实；个别链接可能失效或指向摘要页。",
        "8. Summary 表里的计数是公式；本文件从未在 Excel 或 WPS 中打开过，下拉列表和公式显示如有问题请告诉舒俊。",
        "",
        "分组计数（生成时）：" + "；".join(f"{GROUPS[k]}: {c}" for k, c in sorted(Counter(d["group"] for d in items).items())),
        f"其中校准报告 {sum(1 for d in items if d['calibration'])} 篇。",
    ]
    rd.column_dimensions["A"].width = 150
    for i, t in enumerate(text, 1):
        c = rd.cell(row=i, column=1, value=t)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 1:
            c.font = Font(bold=True, size=13)

    wb.save(args.out)
    if args.csv:
        with open(args.csv, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow([lab for _, lab, _ in COLS])
            for d in items:
                w.writerow([GROUPS[d[k]] if k == "group" else d[k] for k, _, _ in COLS])
    print(f"{n} reports; groups: {dict(sorted(Counter(d['group'] for d in items).items()))}; "
          f"calibration: {sum(1 for d in items if d['calibration'])}; pmc: {sum(1 for d in items if d['pmcid'])}")


if __name__ == "__main__":
    sys.exit(main())
