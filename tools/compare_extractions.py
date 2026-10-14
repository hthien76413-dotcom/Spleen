#!/usr/bin/env python3
"""Compare two independent extractions made with the extraction workbook (protocol v1.4, Section 14).

Usage:
    python3 -I tools/compare_extractions.py extraction_E1.xlsx extraction_E2.xlsx -o comparison.xlsx
    python3 -I tools/compare_extractions.py E1.xlsx E2.xlsx -o comparison.xlsx --sheets Patients Reports
    python3 -I tools/compare_extractions.py E1.xlsx E2.xlsx -o comparison.xlsx --tol-days 1 0.10
    python3 -I tools/compare_extractions.py rob_R1.xlsx rob_R2.xlsx -o rob_comparison.xlsx --sheets RoB

What it does
  * Matches rows by patient_id (Patients sheet) or report_id (Reports sheet, and RoB sheet of the risk-of-bias
    workbook built by tools/build_rob_workbook.py; for RoB, kappa is reported for adequate_primary and for each item).
  * Compares every categorical and numeric field. Free-text fields, source columns, notes, extractor and date
    are not compared (listed in the Summary sheet).
  * Blank = "not yet extracted" (protocol README); NR = "not reported". A field that is blank for one extractor and
    filled for the other goes to the sheet "Incomplete", not "Disagreements", and is left out of the agreement figures.
  * Reports Cohen's kappa (unweighted) for the registered primary exposure (index-episode duration category:
    <=2 d, 3-7 d, 8-30 d, >30 d, plus missing) and for the primary outcome (derived primary_outcome, the raw
    viability field, and the registered binary Non-viable vs other). Days are also reported as exact agreement and as
    agreement within a tolerance (default: 1 day or 10 %, whichever is larger; the protocol does not say how a
    continuous exposure is to be compared, so both are shown and the tolerance is an assumption to be confirmed).
  * The grey (derived) columns of the workbook are formulas that carry no cached value until Excel or WPS has saved the
    file, so they are recomputed here with the same rules (workbook README sheet, derived columns D1-D5).

Input files hold extracted data from published reports and, for OWN_P1-OWN_P5, non-identifiable data. Keep the filled
workbooks and the comparison output outside the git repository; the repository only holds the template.

Requires openpyxl. Python 3.8+.
"""
import argparse
import math
import sys
from collections import Counter, defaultdict
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FIRST_DATA_ROW = 4  # rows 1-3 are the header block of the template; row 1 holds the field codes

# ----------------------------------------------------------------------------- field classification
# Every column code of the template must appear in exactly one group below (checked in the test script against
# tools/build_extraction_workbook.py, so that a change to the template cannot silently go uncompared).
PATIENT_FIELDS = {
    "key": ["patient_id"],
    "identity": ["report_id"],  # a different report_id for the same patient_id is flagged
    "derived": ["own_patient", "idx_dur_category", "thrombocytosis", "primary_outcome"],
    "categorical": [
        "sex", "cond_gv", "cond_goo", "cond_panc_tail", "cond_bowel_obs", "cond_dh", "cond_kidney", "cond_prune",
        "cond_prev_surg", "aetiology", "ws_basis", "presentation", "idx_dur_intermittent", "idx_dur_lower_bound",
        "idx_dur_cat_imprecise", "missed_dx", "known_dx_deferred", "deferred_reason", "torsion_awaiting",
        "img_us", "img_doppler", "img_ct", "img_mri", "whirl_sign", "perfusion", "approach", "torsion_at_op",
        "viability", "viab_basis", "procedure", "fixation", "reason_viable_removed", "recurrence",
        "further_surgery", "status_last_fu", "same_patient_as",
    ],
    "numeric": [
        "age_years", "idx_dur_days", "idx_dur_alt_days", "total_hist_days", "n_earlier_episodes",
        "waiting_interval_days", "platelets", "torsion_degrees", "spleen_size_cm", "followup_months",
    ],
    "free": [
        "extractor", "date_extracted", "age_verbatim", "prev_surg_type", "cond_syndrome", "cond_other",
        "aetiology_cause", "idx_dur_verbatim", "idx_dur_rules", "total_hist_verbatim", "missed_dx_alt",
        "missed_dx_contact", "procedure_other", "fixation_other", "recurrence_details", "post_splenectomy_compl",
        "src_age_sex", "src_exposure", "src_outcome", "src_other", "notes",
    ],
}
REPORT_FIELDS = {
    "key": ["report_id"],
    "identity": [],
    "derived": ["pub_period"],
    "categorical": ["language", "design", "consecutive"],
    "numeric": ["year", "n_children_individual", "n_children_aggregate"],
    "free": [
        "first_author", "journal", "country", "institution", "study_period", "linked_duplicates",
        "extractor", "date_extracted", "notes",
    ],
}
ROB_FIELDS = {  # risk-of-bias workbook (tools/build_rob_workbook.py); one row per report, one file per reviewer
    "key": ["report_id"],
    "identity": [],
    "derived": ["adequate_primary"],
    "categorical": ["q_selection", "q_exposure", "q_outcome", "q_alternative", "q_followup", "q_reporting"],
    "numeric": [],
    "free": [
        "reviewer", "date_assessed", "support_selection", "support_exposure", "support_outcome", "support_other",
        "notes",
    ],
}
SHEETS = {"Patients": PATIENT_FIELDS, "Reports": REPORT_FIELDS, "RoB": ROB_FIELDS}

# absolute and relative tolerance for "agreement within tolerance"; (0, 0) = exact. Day fields can be overridden.
DAY_FIELDS = ["idx_dur_days", "idx_dur_alt_days", "total_hist_days", "waiting_interval_days"]
NUM_TOL = {
    "age_years": (0.02, 0.0),   # two decimals (months / 12)
    "followup_months": (0.5, 0.0),
    "spleen_size_cm": (0.1, 0.0),
}

DUR_CATS = ["<=2 d", "3-7 d", "8-30 d", ">30 d"]
VIABILITY = ["Viable", "Partially non-viable", "Non-viable", "NR", "Not applicable"]


# ----------------------------------------------------------------------------- cell helpers
def clean(v):
    """Normalise a cell: None/blank -> None; numbers stay numbers; numeric-looking text becomes a number."""
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip()
    if s == "":
        return None
    if s.startswith("="):  # an unevaluated formula; never data
        return None
    try:
        return float(s)
    except ValueError:
        return s


def is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool) and not math.isnan(v)


def canon(v):
    """Case-insensitive, whitespace-collapsed key for categorical comparison; NR is always 'NR'."""
    if v is None:
        return None
    if is_num(v):
        return str(v)
    s = " ".join(str(v).split())
    return "NR" if s.lower() == "nr" else s.lower()


# ----------------------------------------------------------------------------- derived columns (workbook D1-D5)
def derive_category(days, imprecise_cat):
    """Workbook column idx_dur_category: the numeric value wins; otherwise the imprecise-duration category (rule 5)."""
    if is_num(days):
        if days <= 2:
            return DUR_CATS[0]
        if days <= 7:
            return DUR_CATS[1]
        if days <= 30:
            return DUR_CATS[2]
        return DUR_CATS[3]
    if imprecise_cat:
        return imprecise_cat
    if days is not None and canon(days) == "NR":
        return "NR"
    return None


def derive_outcome(viability, procedure):
    """Workbook column primary_outcome (D3)."""
    if not viability:
        return None
    v = canon(viability)
    removal = canon(procedure) in ("total splenectomy", "partial splenectomy +/- splenopexy")
    if v == "non-viable":
        return "Non-viable"
    if v == "viable":
        return "Viable"
    if v == "partially non-viable":
        return "Non-viable" if removal else "Review"
    if v == "NR":  # canon() returns NR in capitals
        return "NR - removed, check reason" if removal else "NR"
    return "Not applicable"


def derive_adequate(exposure, outcome, reporting):
    """Risk-of-bias workbook column adequate_primary: Yes only if exposure, outcome and reporting items are all Yes."""
    items = [canon(x) for x in (exposure, outcome, reporting)]
    if any(i is None for i in items):
        return None
    return "Yes" if all(i == "yes" for i in items) else "No"


def derive_thrombocytosis(platelets):
    if is_num(platelets):
        return "Yes" if platelets > 450 else "No"
    return "NR" if platelets is not None and canon(platelets) == "NR" else None


# ----------------------------------------------------------------------------- statistics
def cohen_kappa(pairs):
    """Unweighted Cohen's kappa for a list of (a, b) category pairs.

    Returns dict(n, po, pe, kappa, se, lo, hi). kappa is None when it is undefined (pe == 1, i.e. no variation).
    The standard error is the large-sample formula sqrt(po(1-po)/(n(1-pe)^2)), approximate, and unreliable for small n;
    the interval is cut to [-1, 1]. With the 10-report calibration the interval will be very wide: report it as such.
    """
    n = len(pairs)
    if n == 0:
        return dict(n=0, po=None, pe=None, kappa=None, se=None, lo=None, hi=None)
    cats = sorted({a for a, _ in pairs} | {b for _, b in pairs}, key=str)
    ca, cb = Counter(a for a, _ in pairs), Counter(b for _, b in pairs)
    po = sum(1 for a, b in pairs if a == b) / n
    pe = sum((ca[c] / n) * (cb[c] / n) for c in cats)
    if abs(1 - pe) < 1e-12:
        return dict(n=n, po=po, pe=pe, kappa=None, se=None, lo=None, hi=None)
    k = (po - pe) / (1 - pe)
    se = math.sqrt(po * (1 - po) / (n * (1 - pe) ** 2))
    return dict(n=n, po=po, pe=pe, kappa=k, se=se, lo=max(-1.0, k - 1.96 * se), hi=min(1.0, k + 1.96 * se))


def confusion(pairs, order=None):
    cats = sorted({a for a, _ in pairs} | {b for _, b in pairs}, key=str)
    if order:
        cats = [c for c in order if c in cats] + [c for c in cats if c not in order]
    m = Counter(pairs)
    return cats, [[m[(r, c)] for c in cats] for r in cats]


# ----------------------------------------------------------------------------- reading
def read_sheet(path, sheet, fields):
    wb = load_workbook(path, data_only=False)
    if sheet not in wb.sheetnames:
        raise SystemExit(f"{path}: no sheet named {sheet!r} (found: {', '.join(wb.sheetnames)})")
    ws = wb[sheet]
    codes = {}
    for ci in range(1, ws.max_column + 1):
        code = ws.cell(row=1, column=ci).value
        if code:
            codes[str(code).strip()] = ci
    key = fields["key"][0]
    if key not in codes:
        raise SystemExit(f"{path}, sheet {sheet}: no column with field code {key!r} in row 1 (is this the template?)")
    rows, problems = {}, []
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        kv = clean(ws.cell(row=r, column=codes[key]).value)
        row = {c: clean(ws.cell(row=r, column=ci).value) for c, ci in codes.items()}
        if kv is None:
            if any(v is not None for c, v in row.items() if c not in fields["derived"]):
                problems.append((f"row {r}", f"row has data but no {key}; skipped"))
            continue
        kv = str(kv)
        if kv in rows:
            problems.append((kv, f"{key} appears twice in {path} (row {r}); the first row is used"))
            continue
        rows[kv] = row
    return rows, codes, problems


# ----------------------------------------------------------------------------- comparison
def numbers_agree(a, b, absol, rel):
    if is_num(a) and is_num(b):
        return abs(a - b) <= max(absol, rel * max(abs(a), abs(b)))
    return canon(a) == canon(b)  # NR vs NR; NR vs number is a disagreement


def compare(rows1, rows2, fields, tol):
    disagreements, incomplete = [], []
    stats = defaultdict(lambda: dict(n=0, agree=0, exact=0))
    matched = sorted(set(rows1) & set(rows2))
    for pid in matched:
        a, b = rows1[pid], rows2[pid]
        for f in fields["identity"]:
            if canon(a.get(f)) != canon(b.get(f)):
                disagreements.append((pid, f, a.get(f), b.get(f), "identity"))
        for kind in ("categorical", "numeric"):
            for f in fields[kind]:
                if f not in a and f not in b:
                    continue
                va, vb = a.get(f), b.get(f)
                if va is None and vb is None:
                    continue
                if va is None or vb is None:
                    incomplete.append((pid, f, va, vb))
                    continue
                s = stats[f]
                s["n"] += 1
                if kind == "categorical":
                    ok = canon(va) == canon(vb)
                    s["exact"] += ok
                else:
                    absol, rel = tol.get(f, (0.0, 0.0))
                    ok = numbers_agree(va, vb, absol, rel)
                    s["exact"] += (va == vb) if (is_num(va) and is_num(vb)) else (canon(va) == canon(vb))
                s["agree"] += ok
                if not ok:
                    disagreements.append((pid, f, va, vb, kind))
    only1 = sorted(set(rows1) - set(rows2))
    only2 = sorted(set(rows2) - set(rows1))
    return matched, disagreements, incomplete, stats, only1, only2


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("file1", help="extractor 1 workbook (E1)")
    ap.add_argument("file2", help="extractor 2 workbook (E2)")
    ap.add_argument("-o", "--out", required=True, help="comparison report (.xlsx) to write")
    ap.add_argument("--sheets", nargs="+", default=["Patients"], choices=sorted(SHEETS))
    ap.add_argument("--tol-days", nargs=2, type=float, metavar=("ABS", "REL"), default=[1.0, 0.10],
                    help="tolerance for the duration fields: |a-b| <= max(ABS, REL*max(a,b)); default 1 0.10")
    args = ap.parse_args(argv)

    tol = dict(NUM_TOL)
    for f in DAY_FIELDS:
        tol[f] = tuple(args.tol_days)

    out = Workbook()
    out.remove(out.active)
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="2F75B5")

    def sheet(name, header, rows, widths=None):
        ws = out.create_sheet(name[:31])
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = hdr_font, hdr_fill
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for r in rows:
            ws.append(list(r))
        for i, w in enumerate(widths or [], start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        return ws

    summary_lines = []
    log = lambda s="": (summary_lines.append(s), print(s))  # noqa: E731

    log("提取比较 / extraction comparison")
    log(f"  E1: {args.file1}")
    log(f"  E2: {args.file2}")
    log(f"  日期容差 / tolerance for day fields: |a-b| <= max({args.tol_days[0]:g} d, {args.tol_days[1]:.0%} of the larger value)")
    log("  空白 = 尚未提取，不计入一致率；NR = 报告里没有，与数值不一致算分歧。")
    all_dis, all_inc, all_unm, field_rows, kappa_rows, conf_blocks = [], [], [], [], [], []

    for sh in args.sheets:
        fields = SHEETS[sh]
        r1, codes1, p1 = read_sheet(args.file1, sh, fields)
        r2, codes2, p2 = read_sheet(args.file2, sh, fields)
        for k, msg in p1:
            all_unm.append((sh, "E1", k, msg))
        for k, msg in p2:
            all_unm.append((sh, "E2", k, msg))
        matched, dis, inc, stats, only1, only2 = compare(r1, r2, fields, tol)
        for k in only1:
            all_unm.append((sh, "E1 only", k, "present in E1, missing in E2"))
        for k in only2:
            all_unm.append((sh, "E2 only", k, "present in E2, missing in E1"))
        log("")
        log(f"[{sh}] E1 {len(r1)} 行, E2 {len(r2)} 行, 匹配 {len(matched)}, 只在 E1 {len(only1)}, 只在 E2 {len(only2)}")
        tot_n = sum(s["n"] for s in stats.values())
        tot_ok = sum(s["agree"] for s in stats.values())
        if tot_n:
            log(f"  可比较的项目 {tot_n}，一致 {tot_ok} ({tot_ok / tot_n:.1%})，分歧 {len(dis)}，一方空白 {len(inc)}")
        for pid, f, va, vb, kind in dis:
            all_dis.append((sh, pid, f, va, vb, kind, ""))
        for pid, f, va, vb in inc:
            all_inc.append((sh, pid, f, va, vb))
        for f in sorted(stats):
            s = stats[f]
            field_rows.append((sh, f, s["n"], s["agree"], s["agree"] / s["n"], s["exact"], s["exact"] / s["n"]))

        if sh in ("Patients", "RoB"):
            def derived_pairs(fn):
                out_pairs = []
                for pid in matched:
                    a, b = fn(r1[pid]), fn(r2[pid])
                    if a is not None and b is not None:
                        out_pairs.append((canon(a), canon(b)))
                return out_pairs

            if sh == "Patients":
                def cat(row):
                    return derive_category(row.get("idx_dur_days"), row.get("idx_dur_cat_imprecise"))

                def outcome(row):
                    return derive_outcome(row.get("viability"), row.get("procedure"))

                def binary(row):
                    o = outcome(row)
                    return None if o is None else ("Non-viable" if o == "Non-viable" else "Other")

                specs = [
                    ("主要暴露：指数发作时长类别（四类+缺失）/ primary exposure category", cat, DUR_CATS + ["NR"]),
                    ("主要结局：primary_outcome（推导列，全部类别）", outcome, ["Non-viable", "Viable", "Review"]),
                    ("主要结局：Non-viable vs 其他（协议第 12 节的二分类）", binary, ["Non-viable", "Other"]),
                    ("脾脏存活原始字段 viability", lambda r: r.get("viability"), VIABILITY),
                ]
            else:
                def adequate(row):
                    return derive_adequate(row.get("q_exposure"), row.get("q_outcome"), row.get("q_reporting"))

                specs = [("偏倚风险：adequate_primary（推导列）/ adequate for the primary question", adequate, ["Yes", "No"])]
                for item in ROB_FIELDS["categorical"]:
                    specs.append((f"偏倚风险：{item}", (lambda r, item=item: r.get(item)), ["Yes", "No", "Unclear", "Not applicable"]))
            log("")
            log("  Cohen's kappa（未加权；SE 为大样本近似，n 小时不可靠）")
            for label, fn, order in specs:
                pairs = derived_pairs(fn)
                kp = cohen_kappa(pairs)
                kappa_rows.append((sh, label, kp["n"],
                                   kp["po"], kp["kappa"], kp["lo"], kp["hi"],
                                   "undefined: no variation in the ratings" if (kp["n"] and kp["kappa"] is None) else ""))
                if kp["n"] == 0:
                    log(f"    {label}: 没有两人都填了的患者")
                elif kp["kappa"] is None:
                    log(f"    {label}: n={kp['n']}, 一致 {kp['po']:.1%}, kappa 无法计算（两人评分没有变化）")
                else:
                    log(f"    {label}: n={kp['n']}, 一致 {kp['po']:.1%}, kappa {kp['kappa']:.2f} "
                        f"(95% CI {kp['lo']:.2f} to {kp['hi']:.2f})")
                cats, mat = confusion(pairs, order)
                conf_blocks.append((label, cats, mat))
            # days: exact vs within tolerance, among patients where both gave a number
            for f in (("idx_dur_days", "total_hist_days") if sh == "Patients" else ()):
                s = stats.get(f)
                if s and s["n"]:
                    log(f"  {f}: 两人都填 {s['n']} 例，精确一致 {s['exact']}，容差内一致 {s['agree']}")

    log("")
    log(f"未比较的自由文本/来源列 / not compared: {', '.join(sorted({f for sh in args.sheets for f in SHEETS[sh]['free']}))}")

    sheet("Summary", ["Summary"], [(s,) for s in summary_lines], [120])
    sheet("Disagreements", ["sheet", "id", "field", "E1", "E2", "kind", "consensus (fill in)"], all_dis,
          [10, 16, 26, 28, 28, 12, 28])
    sheet("Incomplete", ["sheet", "id", "field", "E1", "E2"], all_inc, [10, 16, 26, 28, 28])
    sheet("Unmatched", ["sheet", "where", "id", "note"], all_unm, [10, 12, 20, 70])
    sheet("By_field", ["sheet", "field", "n both filled", "agree (with tolerance)", "% agree",
                       "exact agree", "% exact"], field_rows, [10, 26, 12, 14, 10, 12, 10])
    sheet("Kappa", ["sheet", "measure", "n", "observed agreement", "kappa", "95% CI low", "95% CI high", "note"],
          kappa_rows, [10, 60, 8, 14, 10, 12, 12, 40])
    ws = out.create_sheet("Confusion")
    for label, cats, mat in conf_blocks:
        ws.append([label + "  (rows = E1, columns = E2)"])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        ws.append([""] + cats)
        for c, row in zip(cats, mat):
            ws.append([c] + row)
        ws.append([])
    ws.column_dimensions["A"].width = 40
    out.save(args.out)
    print(f"\n已写出 / written: {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
