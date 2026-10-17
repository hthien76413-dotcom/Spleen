"""Calibration practice sheet (50 records) for a new second reviewer (17 Oct 2026).

The 50 records are the calibration set screened by Jun Shu and Fei Peng on 5 Oct 2026 (search/screening/calibration_set_50.ris).
The new reviewer screens them in an Excel/WPS sheet, outside Rayyan, blind to the earlier decisions; afterwards
tools/calibration_compare_new_reviewer.py compares her decisions with the two earlier ones for the discussion of the guide.

Usage (repo root): python3 -I tools/make_calibration_sheet_new_reviewer.py
Writes search/screening/calibration_new_reviewer/calibration_sheet_50_new_reviewer.xlsx  (for the new reviewer)
       search/screening/calibration_new_reviewer/calibration_key.csv                    (NOT for the new reviewer)
Needs openpyxl (pip install openpyxl).
"""
import csv, random, re
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

OUT = 'search/screening/calibration_new_reviewer/'
SEED = 20261017
recs = {}
for blk in open('search/screening/calibration_set_50.ris', encoding='utf-8-sig').read().split('\nER  -'):
    m = re.search(r'^ID  - (.*)$', blk, re.M)
    if not m:
        continue
    ti = ' '.join(re.findall(r'^TI  - (.*)$', blk, re.M))
    ab = ' '.join(re.findall(r'^AB  - (.*)$', blk, re.M)).strip()
    py = (re.search(r'^PY  - (\d{4})', blk, re.M) or [None, ''])[1]
    la = (re.search(r'^LA  - (.*)$', blk, re.M) or [None, ''])[1].strip()[:3]
    recs[m.group(1)] = (ti, ab, py, la)
assert len(recs) == 50, len(recs)
ids = sorted(recs)
random.Random(SEED).shuffle(ids)
with open(OUT + 'calibration_key.csv', 'w', encoding='utf-8-sig', newline='') as f:
    w = csv.writer(f); w.writerow(['blind_id', 'ws_id'])
    for i, w_id in enumerate(ids, 1):
        w.writerow([f'C{i:02d}', w_id])

wb = Workbook()
ws0 = wb.active; ws0.title = '说明'
lines = [
    ('t', '第二位筛选员：50 条校准练习表（儿童游走脾系统综述，PROSPERO CRD420261520329）'),
    ('p', '这是正式筛选前的一次性练习：50 条记录，已随机排序。表里没有任何人的决定。请独立完成，做完之前不要向舒俊或别人询问，也不要去看别人的决定。做完后把整个文件发给舒俊；之后会列出和第一轮校准结果不一致的记录，大家一起讨论指南哪里没写清楚。这 50 条在正式筛选里照样要筛一遍。'),
    ('h', '一、怎么判断（完整规则见《筛选指南（标题和摘要）》；下面是摘要）'),
    ('b', '只依据标题和摘要判断，不要上网查全文或作者信息，也不要用任何 AI 工具、翻译工具或自动化工具辅助决定（协议写明不使用 AI）。每条自己读完再选。'),
    ('b', '三种决定：Include（明确符合）、Maybe（拿不准，进入全文再看）、Exclude（明确不符合）。Include 和 Maybe 都进入全文。拿不准就选 Maybe。'),
    ('b', '纳入条件：游走（异位、可移动）脾，伴或不伴脾蒂扭转，先天性或获得性都符合；患者小于 18 岁；病例报告、病例系列、队列研究，或描述了原始患者的信件和综述。摘要里看不出年龄或是否有个体数据时，不要因此排除，选 Maybe。'),
    ('b', 'Exclude 必须同时在"理由"列选一个固定理由；说不出适用哪一条，就不要 Exclude，改选 Maybe。理由只有这六个：E1 non-human（动物、兽医、实验）；E2 not wandering spleen（副脾含游走副脾、脾种植、正常位置脾脏扭转、脾异位妊娠、移植脾、保脾手术、与脾无关的文章等）；E3 adults only（全部患者 ≥18 岁）；E4 no original patient data（无新病例的综述、社论、评论、只评论别人文章的信件）；E5 not a formal article（只有摘要的会议报告、专利、报纸、产品说明）；E6 other（在备注写一句具体说明）。一条记录符合两条理由时任选一条。'),
    ('b', '没有摘要的记录（摘要栏写"（无摘要）"）：只选 Maybe 或 Exclude，不选 Include；只有标题明确符合某条理由时才 Exclude，其他选 Maybe。'),
    ('h', '二、怎么填'),
    ('b', '在"判断"表里，每条记录填"决定"（下拉选择）和需要时的"理由"（下拉选择）；备注可以留空。请按表里的顺序做，不要漏行，不要改列名、不要删行。'),
    ('b', '在"时间日志"表里记录开始和结束的时间。'),
]
ws0.column_dimensions['A'].width = 130
for i, (k, t) in enumerate(lines, 1):
    c = ws0.cell(row=i, column=1, value=('• ' + t) if k == 'b' else t)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    if k == 't': c.font = Font(bold=True, size=14)
    elif k == 'h': c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='2F75B5')
ws1 = wb.create_sheet('判断')
for j, (h, w_) in enumerate(zip(['blind_id', '年份', '语种', '标题', '摘要', '决定', '理由', '备注'], [9, 7, 7, 46, 100, 11, 26, 30]), 1):
    c = ws1.cell(row=1, column=j, value=h); c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='2F75B5')
    ws1.column_dimensions[chr(64 + j)].width = w_
for i, w_id in enumerate(ids, 2):
    ti, ab, py, la = recs[w_id]
    for j, v in enumerate([f'C{i - 1:02d}', py, la, ti, ab or '（无摘要）', None, None, None], 1):
        c = ws1.cell(row=i, column=j, value=v); c.alignment = Alignment(wrap_text=True, vertical='top')
ws1.freeze_panes = 'D2'
dv1 = DataValidation(type='list', formula1='"Include,Maybe,Exclude"', allow_blank=True)
dv2 = DataValidation(type='list', formula1='"E1 non-human,E2 not wandering spleen,E3 adults only,E4 no original patient data,E5 not a formal article,E6 other"', allow_blank=True)
ws1.add_data_validation(dv1); ws1.add_data_validation(dv2); dv1.add('F2:F51'); dv2.add('G2:G51')
ws2 = wb.create_sheet('时间日志')
for j, (h, w_) in enumerate(zip(['日期', '开始时间', '结束时间', '备注'], [12, 11, 11, 40]), 1):
    c = ws2.cell(row=1, column=j, value=h); c.font = Font(bold=True); ws2.column_dimensions[chr(64 + j)].width = w_
wb.save(OUT + 'calibration_sheet_50_new_reviewer.xlsx')
print('written', OUT)
