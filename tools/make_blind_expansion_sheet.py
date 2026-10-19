"""Second blind sheet for the third reviewer after pre-specified trigger 1 fired (17 Oct 2026).

Why: in the first blind sheet the third reviewer advanced 42 of the 192 targeted both-excluded records (21.9%), above the
pre-specified 10% (tools/blind_adjudication_analysis.py wrote expansion_ws_ids.csv). The pre-specified consequence is that he also
judges ALL remaining both-excluded records with reason E2 or E3 from either reviewer.

Usage (from the repo root): python3 -I tools/make_blind_expansion_sheet.py [FOLDER]    (default search/screening/rescreen/)
Reads  FOLDER/expansion_ws_ids.csv, FOLDER/screening_results_all.csv, search/dedup/records_deduplicated.ris
Writes FOLDER/blind_expansion_sheet.xlsx   (for the third reviewer; no ids, no decisions, no reason for inclusion)
       FOLDER/blind_expansion_key.csv      (NOT for the third reviewer)
Needs openpyxl. Records are shuffled with a fixed seed; blind ids X0001...
"""
import csv, random, re, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

SEED = 20261017
F = sys.argv[1].rstrip('/') + '/' if len(sys.argv) > 1 else 'search/screening/rescreen/'

ids = [r['ws_id'] for r in csv.DictReader(open(F + 'expansion_ws_ids.csv', encoding='utf-8-sig'))]
res = {r['ws_id']: r for r in csv.DictReader(open(F + 'screening_results_all.csv', encoding='utf-8-sig'))}
title, abstract = {}, {}
for blk in open('search/dedup/records_deduplicated.ris', encoding='utf-8-sig').read().split('\nER  -'):
    m = re.search(r'^ID  - (.*)$', blk, re.M)
    if m:
        title[m.group(1)] = ' '.join(re.findall(r'^TI  - (.*)$', blk, re.M))
        abstract[m.group(1)] = ' '.join(re.findall(r'^AB  - (.*)$', blk, re.M)).strip()
rng = random.Random(SEED)
order = ids[:]
rng.shuffle(order)
n = len(order)

with open(F + 'blind_expansion_key.csv', 'w', encoding='utf-8-sig', newline='') as f:
    w = csv.writer(f)
    w.writerow(['blind_id', 'ws_id', 'R1_decision', 'R1_reason', 'R2_decision', 'R2_reason'])
    for i, k in enumerate(order, start=1):
        r = res[k]
        w.writerow([f'X{i:04d}', k, r['R1_JunShu_decision'], r['R1_JunShu_reason'], r['R2_FeiPeng_decision'], r['R2_FeiPeng_reason']])

wb = Workbook()
ws0 = wb.active
ws0.title = '说明'
lines = [
    ('t', '第三位评审：标题摘要盲评表，第二张（儿童游走脾系统综述，PROSPERO CRD420261520329）'),
    ('p', f'本表共 {n} 条记录，已随机排序。表里没有任何人的决定，也没有说明为什么某条记录在这张表里。请逐条独立判断，不要向舒俊或彭飞询问，也不要去看他们的决定。判断规则和第一张表完全相同。'),
    ('h', '一、怎么判断（完整规则见《筛选指南（标题和摘要）》；下面是摘要）'),
    ('b', '只依据标题和摘要判断，不要上网查全文或作者信息，也不要用任何 AI 工具或翻译工具辅助决定（协议写明不使用 AI）。'),
    ('b', '三种决定：Include（明确符合）、Maybe（拿不准，进入全文再看）、Exclude（明确不符合）。Include 和 Maybe 都进入全文。拿不准就选 Maybe。'),
    ('b', '纳入条件：游走（异位、可移动）脾，伴或不伴脾蒂扭转，先天性或获得性都符合；患者小于 18 岁；病例报告、病例系列、队列研究，或描述了原始患者的信件和综述。摘要里看不出年龄或是否有个体数据时，不要因此排除，选 Maybe。'),
    ('b', 'Exclude 必须同时在"理由"列选一个固定理由；说不出适用哪一条，就不要 Exclude，改选 Maybe。理由只有这六个：E1 non-human（动物、兽医、实验）；E2 not wandering spleen（副脾含游走副脾、脾种植、正常位置脾脏扭转、脾异位妊娠、移植脾、保脾手术等）；E3 adults only（全部患者 ≥18 岁）；E4 no original patient data（无新病例的综述、社论、评论）；E5 not a formal article（只有摘要的会议报告、专利、报纸）；E6 other（在备注写一句具体说明）。一条记录符合两条理由时任选一条。'),
    ('b', '没有摘要的记录（摘要栏写"（无摘要）"）：只选 Maybe 或 Exclude，不选 Include；只有标题明确符合某条理由时才 Exclude，其他选 Maybe。'),
    ('b', '"脾扭转""异位脾"没有说明脾脏位置是否正常、患者是否成人时，选 Maybe。'),
    ('h', '二、怎么填'),
    ('b', '在"判断"表里，每条记录填"决定"（下拉选择）和需要时的"理由"（下拉选择）；备注可写一句话，也可留空。请按表里的顺序做，不要漏行。'),
    ('b', '请在"时间日志"表里记录每次开始和结束的时间以及做到第几条（上一张表的文件里没有这张日志，这一张请务必保留并填写）；这是为了在方法部分说明评审用时。'),
    ('b', '做完后把整个文件发给舒俊（不要改列名、不要删行、不要删表）。'),
]
ws0.column_dimensions['A'].width = 130
for i, (k, t) in enumerate(lines, start=1):
    c = ws0.cell(row=i, column=1, value=('• ' + t) if k == 'b' else t)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    if k == 't':
        c.font = Font(bold=True, size=14)
    elif k == 'h':
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='2F75B5')

ws1 = wb.create_sheet('判断')
heads = ['blind_id', '年份', '语种', '标题', '摘要', '决定', '理由', '备注']
widths = [9, 7, 7, 46, 100, 11, 26, 30]
for j, (h, w_) in enumerate(zip(heads, widths), start=1):
    c = ws1.cell(row=1, column=j, value=h)
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = PatternFill('solid', fgColor='2F75B5')
    ws1.column_dimensions[chr(64 + j)].width = w_
for i, k in enumerate(order, start=2):
    vals = [f'X{i - 1:04d}', res[k]['year'] or '', res[k]['language'][:3], title[k], abstract[k] or '（无摘要）', None, None, None]
    for j, v in enumerate(vals, start=1):
        c = ws1.cell(row=i, column=j, value=v)
        c.alignment = Alignment(wrap_text=True, vertical='top')
ws1.freeze_panes = 'D2'
dv1 = DataValidation(type='list', formula1='"Include,Maybe,Exclude"', allow_blank=True)
dv2 = DataValidation(type='list', formula1='"E1 non-human,E2 not wandering spleen,E3 adults only,E4 no original patient data,E5 not a formal article,E6 other"', allow_blank=True)
ws1.add_data_validation(dv1)
ws1.add_data_validation(dv2)
dv1.add(f'F2:F{n + 1}')
dv2.add(f'G2:G{n + 1}')
ws2 = wb.create_sheet('时间日志')
for j, (h, w_) in enumerate(zip(['日期', '开始时间', '结束时间', '从第几条（blind_id）', '到第几条', '本次完成条数', '备注'], [12, 11, 11, 20, 12, 14, 30]), start=1):
    c = ws2.cell(row=1, column=j, value=h)
    c.font = Font(bold=True)
    ws2.column_dimensions[chr(64 + j)].width = w_
wb.properties.title = 'Blind expansion sheet'
wb.save(F + 'blind_expansion_sheet.xlsx')
print('expansion sheet rows', n)
