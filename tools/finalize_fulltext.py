#!/usr/bin/env python3
"""Final decisions of the full-text stage: from the two FT_log files and the filled conflict list to the lists and PRISMA counts.

Guide: "Full-text screening guide (draft)", sections 8, 10 and 11. Run it after
  1. both reviewers have finished and `compare_fulltext_logs.py compare` has written ft_conflicts_for_discussion.csv,
  2. the reviewers have discussed the conflicts and the adjudicator has judged the rest (the file was filled in),
  3. optionally, the IPD requests have an outcome (IPD_requests sheet of the extraction workbook).
It can also be run earlier: the lists are then written as PROVISIONAL (see the status line) so that a half-finished
state cannot be mistaken for the final one.

Usage (from the repo root; needs openpyxl; the helper functions come from tools/compare_fulltext_logs.py):
  python3 -I tools/finalize_fulltext.py FT_R1.xlsx FT_R2.xlsx --conflicts FILLED_ft_conflicts_for_discussion.csv \
      [--ipd EXTRACTION.xlsx] [--out DIR] [--list PATH]

Rules applied
  * Both reviewers gave the same decision -> that decision. Both Exclude with different reasons -> the first applicable reason in
    the D3 order (L, D, E5, E4, E1, E2, E3, A, N, E6), unless exactly one of them is D (that is a conflict, 4.3).
  * Conflicts (any difference in the decision, and 'D not confirmed by both') need an entry in the filled conflict list.
    Final decision = final_decision; if that cell is blank, adjudicator_decision; if that is blank too, discussion_agreed_decision
    (the 'resolved_by' column says which one was used; the first fallback is listed as a warning). A final Exclude needs a reason
    (final_reason, else adjudicator_reason, else discussion_agreed_reason, else the reason both/either reviewer gave).
  * A conflict that is in the current comparison but not in the filled list, or has no decision there, is UNRESOLVED and is not
    counted in any list. The same holds for records handled by one reviewer only.
  * Pending IPD (agreed or resolved): with --ipd, the IPD_requests sheet decides: final_status 'Included with IPD' -> Include,
    'Narrative only' -> Exclude (A); 'Not eligible after check' -> unresolved (the reason has to be settled by the reviewers);
    'Pending' or no row -> stays Pending IPD. Without --ipd every Pending IPD stays pending.
  * Records of retrieval class Y (published before 2000) are never counted as read; a row for one in an FT_log is reported and ignored.
Outputs (DIR, default search/screening/fulltext/final): fulltext_final_list.csv (or fulltext_provisional_list.csv), fulltext_included.csv,
fulltext_pending_ipd.csv, fulltext_unresolved.csv, fulltext_prisma_counts.txt. They contain titles and decisions, no personal data.
"""
import argparse
import collections
import csv
import importlib.util
import os
import sys

sys.dont_write_bytecode = True  # do not leave a __pycache__ folder in tools/
HERE = os.path.dirname(os.path.abspath(__file__))
_spec = importlib.util.spec_from_file_location("cft", os.path.join(HERE, "compare_fulltext_logs.py"))
cft = importlib.util.module_from_spec(_spec)
sys.modules["cft"] = cft
_spec.loader.exec_module(cft)

txt = cft.txt
DETAIL_COLS = ("ws_basis", "n_children_individual", "n_children_aggregate", "duplicate_of")


def read_conflicts(path, issues):
    """Filled conflict list -> {ws_id: row}. Values are checked against the allowed names."""
    out = {}
    if not path:
        return out
    with open(path, encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f):
            ws = txt(r.get("ws_id")).upper()
            if not cft.WS_RE.match(ws):
                issues.append("conflict list: row without a valid ws_id (%r)" % r.get("ws_id"))
                continue
            for col in ("final_decision", "adjudicator_decision", "discussion_agreed_decision"):
                v = txt(r.get(col))
                if v and cft.canon_decision(v) is None:
                    issues.append("%s: %s %r is not Include / Exclude / Pending IPD" % (ws, col, v))
            for col in ("final_reason", "adjudicator_reason", "discussion_agreed_reason"):
                v = txt(r.get(col))
                if v and v not in cft.REASONS:
                    issues.append("%s: %s %r is not exactly a name in the reason list" % (ws, col, v))
            out[ws] = r
    return out


def resolve_conflict(row):
    """(decision, reason, resolved_by, note) from a filled conflict row; decision None if nothing was entered."""
    chain = (("final_decision", "final_reason", "final column"),
             ("adjudicator_decision", "adjudicator_reason", "adjudicator"),
             ("discussion_agreed_decision", "discussion_agreed_reason", "discussion"))
    for dcol, rcol, name in chain:
        d = cft.canon_decision(row.get(dcol))
        if d:
            reason = txt(row.get(rcol))
            if not reason:
                for _, r2, _ in chain:  # a reason entered in another column of the same row
                    if txt(row.get(r2)):
                        reason = txt(row.get(r2))
                        break
            note = ""
            if name != "final column":
                note = "final_decision blank, taken from %s" % name
            return d, reason, name, note
    return None, "", "", ""


def read_ipd(path, issues):
    """IPD_requests sheet -> {ws_id: dict(status, requested, replied, received)}."""
    if not path:
        return None
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    if "IPD_requests" not in wb.sheetnames:
        raise SystemExit("%s has no IPD_requests sheet." % path)
    codes = ["ws_id", "email1_date", "email2_date", "reply_date", "data_received", "final_status"]
    rows = cft.read_sheet(wb, "IPD_requests", codes)
    out = {}
    for r in rows:
        ws = txt(r["ws_id"]).upper()
        if not cft.WS_RE.match(ws):
            issues.append("IPD_requests row %s: ws_id missing or not of the form WS0123" % r["_row"])
            continue
        out[ws] = {"status": txt(r["final_status"]).lower(), "requested": txt(r["email1_date"]) != "",
                   "replied": txt(r["reply_date"]) != "", "received": txt(r["data_received"]).lower() in ("yes", "partial")}
    return out


def pct(a, b):
    return "%.1f%%" % (100.0 * a / b) if b else "n/a"


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file1", help="R1 (Jun Shu) FT_log workbook")
    ap.add_argument("file2", help="R2 (Hongqiang Bian) FT_log workbook")
    ap.add_argument("--conflicts", help="filled ft_conflicts_for_discussion.csv")
    ap.add_argument("--ipd", help="extraction workbook whose IPD_requests sheet holds the outcomes of the requests")
    ap.add_argument("--out", default="search/screening/fulltext/final")
    ap.add_argument("--list", default=cft.DEFAULT_LIST)
    args = ap.parse_args()

    retrieval = cft.read_retrieval_list(args.list)
    if retrieval is None:
        raise SystemExit("Retrieval list not found (%s): needed for the counts." % args.list)
    iss1, iss2 = cft.Issues(), cft.Issues()
    r1, _ = cft.load_log(args.file1, "R1", retrieval, iss1)
    r2, _ = cft.load_log(args.file2, "R2", retrieval, iss2)
    notes = []  # warnings of this script
    for name, iss in (("R1", iss1), ("R2", iss2)):
        if iss.count("ERROR"):
            notes.append("%s file has %d format error(s): run `compare_fulltext_logs.py check` on it first." % (name, iss.count("ERROR")))
    ignored = [(n, w) for n, by in (("R1", r1), ("R2", r2)) for w in by if retrieval.get(w, {}).get("retrieval_class", "").startswith("Y")]
    for n, w in ignored:
        notes.append("%s: %s is class Y (published before 2000) and was logged; ignored." % (n, w))
    for by in (r1, r2):
        for w in [w for w in by if retrieval.get(w, {}).get("retrieval_class", "").startswith("Y")]:
            del by[w]

    conf_issues = []
    filled = read_conflicts(args.conflicts, conf_issues)
    notes.extend("conflict list: " + m for m in conf_issues)
    ipd_issues = []
    ipd = read_ipd(args.ipd, ipd_issues) if args.ipd else None
    notes.extend("IPD_requests: " + m for m in ipd_issues)

    ids = sorted(set(r1) | set(r2))
    final, unresolved = [], []
    used_conflict_ids = set()
    for w in ids:
        info = retrieval.get(w, {})
        base = {"ws_id": w, "year": info.get("year", ""), "title": info.get("title", ""), "language_group": info.get("language_group", ""),
                "source": info.get("source", "")}
        a, b = r1.get(w), r2.get(w)
        if a is None or b is None:
            unresolved.append(dict(base, problem="handled by %s only" % ("R1" if a else "R2")))
            continue
        d1, d2 = a["decision_c"], b["decision_c"]
        if d1 is None or d2 is None:
            unresolved.append(dict(base, problem="invalid or missing decision in %s" % ("R1" if d1 is None else "R2")))
            continue
        c1, c2 = a["reason_c"], b["reason_c"]
        dec = reason = resolved_by = note = ""
        if d1 == d2 and not (d1 == "Exclude" and (c1 == "D") != (c2 == "D")):
            dec = d1
            resolved_by = "agreement"
            if dec == "Exclude":
                reason = cft.proposed_reason(c1, c2) if c1 != c2 else txt(a["reason"])
                if c1 != c2:
                    note = "reasons differ (R1 %s, R2 %s): first applicable in the D3 order" % (c1 or "-", c2 or "-")
        else:
            row = filled.get(w)
            used_conflict_ids.add(w)
            kind = "D not confirmed by both" if d1 == d2 else "%s vs %s" % tuple(sorted((d1, d2), key=cft.DECISIONS.index))
            if row is None:
                unresolved.append(dict(base, problem="conflict (%s) not in the filled conflict list" % kind))
                continue
            dec, reason, resolved_by, note = resolve_conflict(row)
            if dec is None:
                unresolved.append(dict(base, problem="conflict (%s) has no decision in the filled list" % kind))
                continue
            if note:
                notes.append("%s: %s" % (w, note))
            if dec == "Exclude" and not reason:
                reason = cft.proposed_reason(c1, c2)
                if reason:
                    notes.append("%s: final Exclude without a reason in the list; the reviewers' reason %s was used" % (w, reason.split()[0]))
            if dec == "Exclude" and not reason:
                unresolved.append(dict(base, problem="resolved to Exclude but no reason anywhere"))
                continue
            if dec != "Exclude":
                reason = "A aggregate data only" if dec == "Pending IPD" else ""
        if dec == "Exclude" and not reason:
            unresolved.append(dict(base, problem="Exclude without a valid reason in the logs"))
            continue
        # Pending IPD -> outcome of the request
        ipd_note = ""
        if dec == "Pending IPD" and ipd is not None:
            st = ipd.get(w, {}).get("status", "")
            if st == "included with ipd":
                dec, reason, ipd_note = "Include", "", "IPD received"
            elif st.startswith("narrative only"):  # the drop-down value is 'Narrative only (no individual data)'
                dec, reason, ipd_note = "Exclude", "A aggregate data only", "IPD requested, not received"
            elif st == "not eligible after check":
                unresolved.append(dict(base, problem="IPD_requests says 'Not eligible after check': the reviewers must settle the reason"))
                continue
        # details of the two reviewers
        # Details (basis for wandering spleen, child counts, duplicate link). The final value is filled only when it is clear:
        # both reviewers gave the same decision as the final one and agree, or only one of them did and its values are taken.
        det = {c: "" for c in DETAIL_COLS}
        for c in DETAIL_COLS:
            for tag, x in (("R1_", a), ("R2_", b)):
                det[tag + c] = txt(x[c]).upper() if c == "duplicate_of" else txt(x[c])
        # IPD received: the child counts are settled when the data arrive; the count at the time of the request stays in the
        # IPD_requests sheet and is not used as a count of children with individual data.
        same_side = [(t, x) for t, x in (("R1_", a), ("R2_", b)) if x["decision_c"] == dec]
        details_state = ""
        if dec in ("Include", "Pending IPD") and not ipd_note:
            cols = DETAIL_COLS if dec == "Include" else ("n_children_aggregate",)
            if len(same_side) == 2:
                diff = [c for c in cols if det["R1_" + c].lower() != det["R2_" + c].lower()]
                details_state = "agree" if not diff else "differ: " + ", ".join(diff)
                for c in cols:
                    if c not in diff:
                        det[c] = det["R1_" + c]
            elif len(same_side) == 1:
                details_state = "one reviewer only (%s)" % same_side[0][0].rstrip("_")
                for c in cols:
                    det[c] = det[same_side[0][0] + c]
            else:
                details_state = "no reviewer gave this decision"
        elif ipd_note == "IPD received":
            details_state = "IPD received: counts to be taken from the data"
        final.append(dict(base, decision=dec, reason=reason, resolved_by=resolved_by, ipd_outcome=ipd_note,
                          R1_decision=d1, R1_reason=txt(a["reason"]), R2_decision=d2, R2_reason=txt(b["reason"]),
                          details_agree=details_state, note=note, **det))

    # conflicts in the filled list that are not conflicts now (stale file)
    stale = [w for w in filled if w not in used_conflict_ids]
    for w in stale:
        notes.append("%s is in the filled conflict list but is no conflict in the current logs (stale list, or a log was changed)" % w)

    # ------------------------------------------------------------------ counts
    to_read = [w for w, i in retrieval.items() if i["retrieval_class"] == "retrieve"]
    cls = collections.Counter(i["retrieval_class"][0] for i in retrieval.values())
    by_id = {f["ws_id"]: f for f in final}
    unres_ids = {u["ws_id"] for u in unresolved}
    not_handled = [w for w in to_read if w not in r1 and w not in r2]
    complete = not unresolved and not not_handled
    dec_n = collections.Counter(f["decision"] for f in final)
    ex_reason = collections.Counter(f["reason"].split()[0] for f in final if f["decision"] == "Exclude")
    resolved_by = collections.Counter(f["resolved_by"] for f in final)
    n_conf_total = sum(1 for w in used_conflict_ids)
    n_conf_ok = sum(1 for f in final if f["resolved_by"] not in ("agreement",))

    inc = [f for f in final if f["decision"] == "Include"]
    pend = [f for f in final if f["decision"] == "Pending IPD"]
    inc_ind = sum(int(f["n_children_individual"]) for f in inc if f["n_children_individual"].isdigit())
    inc_agg = sum(int(f["n_children_aggregate"]) for f in inc if f["n_children_aggregate"].isdigit() and not f["ipd_outcome"])
    inc_missing = sum(1 for f in inc if not f["n_children_individual"].isdigit() and not f["ipd_outcome"])
    pend_agg = sum(int(f["n_children_aggregate"]) for f in pend if f["n_children_aggregate"].isdigit())
    dups = sum(1 for f in inc if f["R1_duplicate_of"] or f["R2_duplicate_of"])
    n_d = ex_reason.get("D", 0)
    assessed = len(final) - n_d  # records for which a full text was obtained and read
    safety = [w for w, i in retrieval.items() if "E4 safety net" in i.get("source", "")]
    def safety_state(w):
        if w in by_id:
            return by_id[w]["decision"] + (" " + by_id[w]["reason"].split()[0] if by_id[w]["decision"] == "Exclude" else "")
        c = retrieval[w]["retrieval_class"]
        return "not read (class %s)" % c[0] if c[0] in "YL" else "not decided yet"
    safety_final = collections.Counter(safety_state(w) for w in safety)

    L = []
    status = "FINAL" if complete else "PROVISIONAL"
    L.append("FULL-TEXT STAGE: FINAL DECISIONS AND PRISMA COUNTS - %s" % status)
    if not complete:
        L.append("  PROVISIONAL because: %d unresolved record(s), %d record(s) on the retrieval list not handled by either reviewer. Do not quote these numbers." % (len(unresolved), len(not_handled)))
    L.append("R1 file %s | R2 file %s | conflicts %s | IPD %s" % (os.path.basename(args.file1), os.path.basename(args.file2),
                                                              os.path.basename(args.conflicts) if args.conflicts else "(none given)",
                                                              os.path.basename(args.ipd) if args.ipd else "(none given: every Pending IPD stays pending)"))
    L.append("")
    L.append("FROM THE RETRIEVAL LIST (search/screening/rescreen/fulltext_retrieval_list.csv)")
    L.append("  Records advanced from title/abstract screening (incl. the E4 safety net): %d" % len(retrieval))
    L.append("  Not sought, published before 2000 (class Y): %d" % cls.get("Y", 0))
    L.append("  Not sought, other language, published 2000 or later (class L): %d" % cls.get("L", 0))
    L.append("  Full text to be sought (class retrieve): %d" % len(to_read))
    L.append("")
    L.append("FULL-TEXT ASSESSMENT")
    L.append("  Handled by both reviewers: %d | one reviewer only: %d | neither: %d" % (
        sum(1 for w in to_read if w in r1 and w in r2), sum(1 for w in to_read if (w in r1) != (w in r2)), len(not_handled)))
    L.append("  Records with a final decision: %d (of %d handled by at least one)" % (len(final), len(ids)))
    L.append("  Full text not obtained, reason D (confirmed): %d" % n_d)
    L.append("  Full texts assessed (final decision, D excluded): %d" % assessed)
    L.append("  Excluded after full-text assessment: %d" % sum(v for k, v in ex_reason.items() if k != "D"))
    for k in cft.ORDER:
        if k != "D" and ex_reason.get(k):
            L.append("    %-3s %-32s %d" % (k, cft.REASON_BY_CODE[k][len(k) + 1:], ex_reason[k]))
    other = sum(v for k, v in ex_reason.items() if k not in cft.ORDER)
    if other:
        L.append("    (other or invalid reason text: %d)" % other)
    L.append("  Pending IPD (still waiting): %d records, %d children aggregate-only" % (len(pend), pend_agg))
    L.append("  Included: %d reports; children with individual data %d, aggregate-only %d%s" % (
        len(inc), inc_ind, inc_agg, "; %d included report(s) have no agreed child count yet" % inc_missing if inc_missing else ""))
    L.append("  Linked duplicate reports among the included (duplicate_of filled): %d" % dups)
    L.append("  Included via IPD received: %d report(s); their child counts come from the data, not from this stage" % sum(1 for f in inc if f["ipd_outcome"] == "IPD received"))
    L.append("  Included records whose R1 and R2 details differ (ws_basis / child counts / duplicate link): %d (settle at the extraction consensus)" % sum(1 for f in inc if f["details_agree"].startswith("differ")))
    L.append("  Included records whose details come from one reviewer only (the other did not include): %d" % sum(1 for f in inc if f["details_agree"].startswith("one reviewer")))
    L.append("")
    L.append("HOW CONFLICTS WERE SETTLED")
    L.append("  Conflicts in the current logs: %d | resolved: %d | unresolved: %d" % (n_conf_total, n_conf_ok, sum(1 for u in unresolved if u["problem"].startswith("conflict"))))
    for k in ("final column", "adjudicator", "discussion"):
        if resolved_by.get(k):
            L.append("    resolved by %s: %d" % (k, resolved_by[k]))
    L.append("  Both Exclude with different reasons settled by the D3 order: %d" % sum(1 for f in final if f["note"].startswith("reasons differ")))
    if safety:
        L.append("")
        L.append("E4 SAFETY-NET RECORDS (%d on the list; those under class Y or L are not read)" % len(safety))
        for k, v in sorted(safety_final.items()):
            L.append("    %s: %d" % (k, v))
    if ipd is not None:
        req = [w for w in ipd if ipd[w]["requested"]]
        L.append("")
        L.append("IPD REQUESTS (IPD_requests sheet)")
        L.append("  Requested: %d | replied: %d | individual data received: %d | final Include (IPD received): %d | final Exclude A (not received): %d" % (
            len(req), sum(1 for w in req if ipd[w]["replied"]), sum(1 for w in req if ipd[w]["received"]),
            sum(1 for f in final if f["ipd_outcome"] == "IPD received"), sum(1 for f in final if f["ipd_outcome"].startswith("IPD requested"))))
        not_pending_rows = [w for w in ipd if w not in by_id and w not in unres_ids]
        if not_pending_rows:
            L.append("  IPD_requests rows for records that have no final decision in the FT logs: %s" % ", ".join(sorted(not_pending_rows)[:12]))
    if unresolved:
        L.append("")
        L.append("UNRESOLVED (%d) - see fulltext_unresolved.csv" % len(unresolved))
        for k, v in collections.Counter(u["problem"].split(" (")[0] if u["problem"].startswith("conflict") else u["problem"].split(":")[0] for u in unresolved).items():
            L.append("    %s: %d" % (k, v))
    if notes:
        L.append("")
        L.append("NOTES AND WARNINGS (%d)" % len(notes))
        L.extend("  " + n for n in notes[:60])
        if len(notes) > 60:
            L.append("  ... %d more" % (len(notes) - 60))
    text = "\n".join(L) + "\n"
    print(text)

    os.makedirs(args.out, exist_ok=True)
    cols = ["ws_id", "year", "language_group", "title", "source", "decision", "reason", "resolved_by", "ipd_outcome", "R1_decision", "R1_reason",
            "R2_decision", "R2_reason", "details_agree", "ws_basis", "n_children_individual", "n_children_aggregate", "duplicate_of",
            "R1_ws_basis", "R2_ws_basis", "R1_n_children_individual", "R2_n_children_individual", "R1_n_children_aggregate",
            "R2_n_children_aggregate", "R1_duplicate_of", "R2_duplicate_of", "note"]

    def write(name, rows, columns):
        with open(os.path.join(args.out, name), "w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore")
            w.writeheader()
            for r in rows:
                w.writerow(r)

    for stale_name in ("fulltext_final_list.csv", "fulltext_provisional_list.csv"):
        p = os.path.join(args.out, stale_name)
        if os.path.exists(p):
            os.remove(p)  # only one of the two names may exist, so a provisional list is never mistaken for the final one
    write("fulltext_final_list.csv" if complete else "fulltext_provisional_list.csv", final, cols)
    write("fulltext_included.csv", inc, cols)
    write("fulltext_pending_ipd.csv", pend, cols)
    write("fulltext_unresolved.csv", unresolved, ["ws_id", "year", "language_group", "title", "source", "problem"])
    with open(os.path.join(args.out, "fulltext_prisma_counts.txt"), "w", encoding="utf-8") as f:
        f.write(text)
    print("Wrote the lists and fulltext_prisma_counts.txt in %s (%s)" % (args.out, status))


if __name__ == "__main__":
    main()
