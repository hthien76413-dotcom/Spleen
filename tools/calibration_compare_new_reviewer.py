"""Compare the new reviewer's 50-record calibration sheet with the first-round calibration (Jun Shu, Fei Peng).

Usage (repo root): python3 -I tools/calibration_compare_new_reviewer.py FILLED_SHEET.xlsx OUT_DIR
Writes OUT_DIR/calibration_new_reviewer_results.csv (one row per record, with titles) and OUT_DIR/calibration_new_reviewer_summary.txt.
The summary is for Claude, Jun Shu and the new reviewer AFTER she has finished the sheet (it shows the earlier decisions).
"""
import csv, sys
from collections import Counter
from openpyxl import load_workbook

sheet, out = sys.argv[1], sys.argv[2].rstrip('/') + '/'
key = {r['blind_id']: r['ws_id'] for r in csv.DictReader(open('search/screening/calibration_new_reviewer/calibration_key.csv', encoding='utf-8-sig'))}
old = {r['final_id']: r for r in csv.DictReader(open('search/screening/calibration_results_filled.csv', encoding='utf-8-sig'))}
ws = load_workbook(sheet, data_only=True)['判断']
new = {}
for row in ws.iter_rows(min_row=2, values_only=True):
    if row[0]:
        new[row[0]] = ((row[5] or '').strip(), (row[6] or '').strip(), row[3])
assert set(new) == set(key), 'blind ids do not match the key'


def kappa(a, b):
    n = len(a); po = sum(x == y for x, y in zip(a, b)) / n
    ca, cb = Counter(a), Counter(b)
    pe = sum(ca[k] * cb[k] for k in set(ca) | set(cb)) / n ** 2
    return po, (po - pe) / (1 - pe) if pe < 1 else float('nan')


adv = lambda d: 'advance' if d in ('Include', 'Maybe') else 'Exclude'
L, rows = [], []
for b in sorted(new):
    w = key[b]; d, rs, ti = new[b]
    rows.append([b, w, ti, d, rs, old[w]['R1_JunShu_decision'], old[w]['R1_reason'], old[w]['R2_FeiPeng_decision'], old[w]['R2_reason'], old[w]['has_abstract']])
with open(out + 'calibration_new_reviewer_results.csv', 'w', encoding='utf-8-sig', newline='') as f:
    wr = csv.writer(f)
    wr.writerow(['blind_id', 'ws_id', 'title', 'new_decision', 'new_reason', 'R1_JunShu_decision', 'R1_reason', 'R2_FeiPeng_decision', 'R2_reason', 'has_abstract'])
    wr.writerows(rows)
blank = [r[0] for r in rows if r[3] not in ('Include', 'Maybe', 'Exclude')]
L.append('records without a valid decision: %d %s' % (len(blank), blank))
L.append('Exclude without a reason: %d; Include on a record without abstract: %d' % (
    sum(1 for r in rows if r[3] == 'Exclude' and not r[4]), sum(1 for r in rows if r[3] == 'Include' and r[9] == 'no')))
ok = [r for r in rows if r[3] in ('Include', 'Maybe', 'Exclude')]
for name, idx in (('Jun Shu (first-round calibration)', 5), ('Fei Peng (first-round calibration)', 7)):
    a, b = [r[3] for r in ok], [r[idx] for r in ok]
    po, k = kappa(a, b); po2, k2 = kappa([adv(x) for x in a], [adv(x) for x in b])
    L.append('%s: 3-category agreement %.3f, kappa %.2f; advance-vs-exclude agreement %.3f, kappa %.2f' % (name, po, k, po2, k2))
L.append('new reviewer: Include %d / Maybe %d / Exclude %d' % tuple(sum(1 for r in ok if r[3] == d) for d in ('Include', 'Maybe', 'Exclude')))
L.append('')
L.append('records where the new reviewer differs in advance-vs-exclude from either earlier reviewer (for discussion):')
for r in ok:
    if adv(r[3]) != adv(r[5]) or adv(r[3]) != adv(r[7]):
        L.append('  %s %s | new %s %s | JS %s %s | FP %s %s | %s' % (r[0], r[1], r[3], r[4][:2], r[5], r[6][:2], r[7], r[8][:2], r[2][:70]))
open(out + 'calibration_new_reviewer_summary.txt', 'w', encoding='utf-8').write('\n'.join(L) + '\n')
print('\n'.join(L))
