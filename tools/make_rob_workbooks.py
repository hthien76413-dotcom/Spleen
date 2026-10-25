#!/usr/bin/env python3
"""Write the risk-of-bias workbooks (protocol v1.8 Section 15; tool of Murad et al. 2018 for case reports and case series).

One row per included report (Include rows of search/screening/fulltext/final/fulltext_included.csv; the own series OWN is not
assessed by its own authors). Two copies, R1 and R2, to be filled independently. Six items, each answered Yes (satisfied) / No /
Unclear; the two items that do not apply (challenge-rechallenge, dose-response) are omitted, as in the protocol. No summary score.
Derived column `adequate`: Yes when exposure, outcome and reporting are all Yes; No when all three are filled and any is not Yes;
blank until all three are filled (protocol: 'adequate for the primary question when the exposure, outcome and reporting items are
all satisfied'). The Unclear-counts-as-No reading is Claude's.
Item wording is Claude's paraphrase of Murad 2018 and of the protocol's Section 15 wording; check it against the paper.

Usage:  python3 -I tools/make_rob_workbooks.py [--included CSV] [--prefill CSV] [--out-dir DIR]     (needs openpyxl)
"""
import argparse, csv, os, sys
sys.dont_write_bytecode = True
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ap = argparse.ArgumentParser()
ap.add_argument("--included", default=os.path.join(ROOT, "search/screening/fulltext/final/fulltext_included.csv"))
ap.add_argument("--prefill", default=os.path.join(ROOT, "search/screening/fulltext/formal_prefill_540.csv"))
ap.add_argument("--out-dir", default=os.path.join(ROOT, "search/rob"))
A = ap.parse_args()
ITEMS = [
    ("selection", "Selection", "Do the patients represent the whole experience of the investigators (centre), or is the selection method stated, so that other similar patients are unlikely to have been left out?"),
    ("exposure", "Ascertainment of exposure", "Are the timing of symptom onset and of operation stated precisely enough to give the duration of the index episode?"),
    ("outcome", "Ascertainment of outcome", "Is splenic viability at operation stated, with its basis (operative findings or histology)?"),
    ("alternative", "Alternative explanations", "Were other explanations for the outcome considered or ruled out (for example, the spleen removed for a technical reason, or a pre-existing splenic disease)?"),
    ("followup", "Adequacy of follow-up", "Was follow-up long enough for the outcomes of interest (splenic status, recurrence) to occur, and is it reported?"),
    ("reporting", "Sufficiency of reporting", "Are the patients described in enough detail for others to use the information in their own practice?"),
]
LAB = {1: "R1", 2: "R2"}
NAMES = {1: "舒俊", 2: "卞红强"}
README = [
    ("title", "偏倚风险评价（Murad 2018 条目）"),
    ("warn", "两位评价员各填各的，做完之前不要看对方的文件，不要讨论具体的报告。这份文件只有报告编号，没有任何评价。"),
    ("h", "怎么填"),
    ("b", "每篇报告一行；每个条目选 Yes（满足）/ No / Unclear。条目文字见下方，是对 Murad 等 2018 和协议第 15 节的转述，请对照原文核对。"),
    ("b", "不算总分。adequate 列是公式：持续时间（exposure）、结局（outcome）、报告充分性（reporting）三项都是 Yes 才是 Yes；三项都填了而有一项不是 Yes 为 No；没填完为空。Unclear 按未满足算（这是 Claude 的读法，待确认）。"),
    ("b", "按这篇报告里用于主要问题的患者整体来评价；一篇里不同患者情况不同时，按多数或最关键的患者判断，并在 note 里写明。"),
    ("b", "挑战-再挑战、剂量-反应两项不适用，已省略（协议第 15 节）。本院 5 例不评价。"),
    ("h", "条目"),
] + [("b", "%s（%s）：%s" % (k, lab, t)) for k, lab, t in ITEMS] + [
    ("h", "做完以后"),
    ("b", "各自另存带日期的副本。两人都做完后运行 `python3 -I tools/compare_rob.py compare 文件R1.xlsx 文件R2.xlsx --out DIR`，得到分歧清单；讨论，讨论不了的交杨俊；统一后的结果写入 rob.csv（报告编号、adequate），供敏感性分析 5 使用。"),
    ("b", "评价员是谁协议第 15 节只写了'两位评价员'；这里按提取员（舒俊、卞红强）安排，这是提议，待确认。"),
    ("h", "没有验证的事"),
    ("b", "没有在 Excel 或 WPS 里打开过；下拉列表、公式和版面是否正常没有验证。"),
]

def main():
    inc = sorted(r["ws_id"] for r in csv.DictReader(open(A.included, encoding="utf-8-sig", newline="")) if r["decision"] == "Include")
    pre = {r["ws_id"]: r for r in csv.DictReader(open(A.prefill, encoding="utf-8-sig", newline=""))}
    miss = [i for i in inc if i not in pre]
    if miss: sys.exit("not in prefill list: %s" % miss)
    os.makedirs(A.out_dir, exist_ok=True)
    blue = PatternFill("solid", fgColor="2F75B5"); grey = PatternFill("solid", fgColor="E7E6E6")
    for k in (1, 2):
        wb = Workbook(); rd = wb.active; rd.title = "README"; rd.column_dimensions["A"].width = 150
        for r, (kind, t) in enumerate(README, 1):
            c = rd.cell(r, 1, ("• " + t) if kind == "b" else t); c.alignment = Alignment(wrap_text=True, vertical="top")
            if kind == "title": c.font = Font(bold=True, size=14)
            elif kind == "warn": c.font = Font(bold=True, color="C00000"); c.fill = PatternFill("solid", fgColor="FDE9D9")
            elif kind == "h": c.font = Font(bold=True, size=12, color="FFFFFF"); c.fill = blue
            rd.row_dimensions[r].height = max(18, 16 * (-(-len(t) // 70)))
        ws = wb.create_sheet("RoB")
        heads = ["report_id", "year", "language", "title"] + [i[0] for i in ITEMS] + ["adequate", "reviewer", "note"]
        for j, h in enumerate(heads, 1):
            c = ws.cell(1, j, h); c.font = Font(bold=True, color="FFFFFF"); c.fill = blue
        for i, ws_id in enumerate(inc, 2):
            p = pre[ws_id]
            for j, v in enumerate([ws_id, p["year"], p["language"], p["title"]], 1):
                c = ws.cell(i, j, v); c.fill = grey; c.alignment = Alignment(wrap_text=(j == 4), vertical="top")
            ws.cell(i, 11, '=IF(OR(F%d="",G%d="",J%d=""),"",IF(AND(F%d="Yes",G%d="Yes",J%d="Yes"),"Yes","No"))' % ((i,) * 6)).fill = grey
            ws.cell(i, 12, LAB[k]).fill = grey
        last = len(inc) + 1
        dv = DataValidation(type="list", formula1='"Yes,No,Unclear"', allow_blank=True); ws.add_data_validation(dv); dv.add("E2:J%d" % last)
        for col, w in zip("ABCDEFGHIJKLM", (10, 7, 10, 70, 11, 11, 11, 11, 11, 11, 10, 9, 50)): ws.column_dimensions[col].width = w
        ws.freeze_panes = "E2"; ws.auto_filter.ref = "A1:M%d" % last
        wb.properties.title = "Risk of bias %s (blank)" % LAB[k]; wb.properties.creator = "Wandering spleen SR team"
        out = os.path.join(A.out_dir, "rob_%s_blank.xlsx" % LAB[k]); wb.save(out)
        print("saved", out, NAMES[k], len(inc), "reports")

main()
