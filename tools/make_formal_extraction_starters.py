#!/usr/bin/env python3
"""Write the two starter workbooks for the FORMAL extraction (E1 = Jun Shu, E2 = Hongqiang Bian).

Population: the Include rows of search/screening/fulltext/final/fulltext_included.csv (all included reports, including the 10
calibration reports and the pilot papers, which are extracted again). Each starter is the normal extraction template with
FT_log and IPD_requests removed, report_id and extractor filled in Reports, a Report_list sheet (WS id, year, language, abstract
yes/no, title, DOI, link; no decision, no source, no strata) and a short README. Patients is empty. The five own children (OWN)
have their own starters (tools/make_own_series_draft.py --starter).

Usage:  python3 -I tools/make_formal_extraction_starters.py [--included CSV] [--prefill CSV] [--out-dir DIR]
Needs openpyxl; builds the template with tools/build_extraction_workbook.py first.
"""
import argparse, csv, os, subprocess, sys, tempfile
sys.dont_write_bytecode = True
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
ap.add_argument("--included", default=os.path.join(ROOT, "search/screening/fulltext/final/fulltext_included.csv"))
ap.add_argument("--prefill", default=os.path.join(ROOT, "search/screening/fulltext/formal_prefill_540.csv"))
ap.add_argument("--out-dir", default=os.path.join(ROOT, "search/extraction/formal"))
A = ap.parse_args()
FIRST = 4
NAMES = {"E1": "舒俊", "E2": "卞红强"}
README = [
    ("title", "正式提取（起始表）"),
    ("warn", "这份文件里只有纳入报告的编号和提取员代码，没有任何数据。请只根据全文填写；两人都交文件之前，不要打开对方的文件，不要讨论具体的报告，不要看预试验提取表。"),
    ("h", "要做什么"),
    ("b", "提取员：E1 = 舒俊，E2 = 卞红强。Reports 表每篇报告一行（已建好，extractor 已填）。Patients 表每名符合条件的儿童一行，patient_id 写成 WS0159_P1…，report_id 填 WS 编号，extractor 填 E1 或 E2。"),
    ("b", "Report_list 表是名单（编号、年份、语言、标题、DOI、链接）。包含 10 篇校准报告和预试验的报告，都要按校准后的规则重新提取。全文和提取表不放进仓库。逐例资料在线上附录里的，附录也要取到。"),
    ("b", "规则以《提取操作说明》v2026-10-25d 为准，尤其第 12 节 X13–X20：脾脏活力只采纳作者对术中脾状态的直接描述；未提及填 NR；只有至就诊或入院前时长时标时长下限；可能重复的患者只在报告层 notes 说明，不填 same_patient_as。协议 v1.7 同样规定了活力的记录方式。"),
    ("b", "空白 = 还没提取；NR = 报告里没有。不要猜值。拿不准的写在 notes。灰色列是公式，不要改；表头前三行和 Lists 表不要改。不写姓名、住院号、出生日期、具体日期。"),
    ("b", "和偏倚风险评价合并阅读：每篇全文只读一遍，同一遍里填这份提取表和 search/rob/ 里自己的偏倚风险表。两份表各自独立，由两人分别填；偏倚风险不是抄提取值。"),
    ("b", "遇到提取说明没有覆盖的情况，在 notes 里写明，并记下 WS 编号告诉杨俊，不要自己发明规则。"),
    ("h", "做完以后"),
    ("b", "保存前对照提取说明第 11 节的自查清单；运行 `python3 -I tools/compare_extraction.py check 你的文件.xlsx --extractor E1`（E2 同理）改掉格式错误。每次用完另存带日期的副本，不覆盖旧的。"),
    ("b", "把文件交杨俊保管；两人都交了以后，由杨俊决定谁运行 compare，分歧先由两位提取员讨论，讨论不了的交杨俊裁定。"),
    ("h", "没有验证的事"),
    ("b", "这份文件没有在 Excel 或 WPS 里打开过；下拉列表和版面实际是否正常没有验证。"),
]

def readme(ws):
    ws.delete_rows(1, ws.max_row + 5)
    ws.column_dimensions["A"].width = 150
    for r, (k, t) in enumerate(README, start=1):
        c = ws.cell(row=r, column=1, value=("• " + t) if k == "b" else t)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if k == "title": c.font = Font(bold=True, size=14)
        elif k == "warn":
            c.font = Font(bold=True, color="C00000"); c.fill = PatternFill("solid", fgColor="FDE9D9")
        elif k == "h":
            c.font = Font(bold=True, size=12, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="2F75B5")
        ws.row_dimensions[r].height = max(18, 16 * (-(-len(t) // 70)))

def main():
    with open(A.included, encoding="utf-8-sig", newline="") as fh:
        inc = [r for r in csv.DictReader(fh) if r["decision"] == "Include"]
    with open(A.prefill, encoding="utf-8-sig", newline="") as fh:
        pre = {r["ws_id"]: r for r in csv.DictReader(fh)}
    ids = sorted(r["ws_id"] for r in inc)
    missing = [i for i in ids if i not in pre]
    if missing: sys.exit("not in prefill list: %s" % missing)
    rows = [pre[i] for i in ids]
    tmp = os.path.join(tempfile.mkdtemp(prefix="fx_"), "t.xlsx")
    subprocess.run([sys.executable, "-I", os.path.join(HERE, "build_extraction_workbook.py"), tmp], check=True, stdout=subprocess.DEVNULL)
    os.makedirs(A.out_dir, exist_ok=True)
    for code in ("E1", "E2"):
        wb = load_workbook(tmp)
        for n in ("FT_log", "IPD_requests"):
            if n in wb.sheetnames: del wb[n]
        ws = wb["Reports"]
        cols = {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1) if ws.cell(row=1, column=c).value}
        for i, r in enumerate(rows):
            ws.cell(row=FIRST + i, column=cols["report_id"], value=r["ws_id"])
            ws.cell(row=FIRST + i, column=cols["extractor"], value=code)
        readme(wb["README"])
        cl = wb.create_sheet("Report_list", 1)
        for j, h in enumerate(["No.", "WS ID", "Year", "Language", "Abstract in database record", "Title", "DOI", "Link"], 1):
            c = cl.cell(row=1, column=j, value=h); c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="2F75B5")
        link = "url" if "url" in rows[0] else [k for k in rows[0] if "link" in k or "url" in k][0]
        for i, r in enumerate(rows, 2):
            vals = [i - 1, r["ws_id"], r["year"], r["language"], r["has_abstract"], r["title"], r["doi"], r[link]]
            for j, v in enumerate(vals, 1):
                cl.cell(row=i, column=j, value=v).alignment = Alignment(wrap_text=(j == 6), vertical="top")
        for col, w in zip("ABCDEFGH", (6, 10, 7, 10, 12, 70, 34, 60)): cl.column_dimensions[col].width = w
        cl.freeze_panes = "A2"
        ws.sheet_properties.tabColor = "C00000"; wb["Patients"].sheet_properties.tabColor = "C00000"
        wb.active = 0
        wb.properties.title = "Extraction %s (formal, blank)" % code
        wb.properties.creator = "Wandering spleen SR team"
        out = os.path.join(A.out_dir, "extraction_%s_formal_starter.xlsx" % code)
        wb.save(out); print("saved", out, NAMES[code], len(rows), "reports")

main()
