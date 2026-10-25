# SYNTHETIC consensus workbook (invented data; true log-odds 0.4 per doubling) to test analysis_crosscheck.py. Run from the repo root.
import sys, subprocess, random, math
from openpyxl import load_workbook
from openpyxl import Workbook
from pathlib import Path
wb=Workbook(); wb.remove(wb.active)
for name, cols in [('Reports',['report_id', 'year', 'language', 'design', 'extractor', 'linked_duplicates', 'consecutive']),('Patients',['aetiology', 'age_years', 'cond_dh', 'cond_goo', 'cond_gv', 'extractor', 'fixation', 'idx_dur_alt_days', 'idx_dur_days', 'idx_dur_lower_bound', 'known_dx_deferred', 'missed_dx', 'patient_id', 'platelets', 'procedure', 'report_id', 'same_patient_as', 'status_last_fu', 'torsion_at_op', 'torsion_awaiting', 'total_hist_days', 'viability', 'ws_basis'])]:
    ws=wb.create_sheet(name); ws.append(cols); ws.append(['']*len(cols)); ws.append(['']*len(cols))
random.seed(7)
R=wb['Reports']; P=wb['Patients']
rc={R.cell(1,c).value:c for c in range(1,R.max_column+1)}
pc={P.cell(1,c).value:c for c in range(1,P.max_column+1)}
rr=4; pr=4; TRUE=0.4  # log-odds per doubling
reports=[]
for i in range(60):
    rid="OWN" if i==0 else "WS%04d"%(100+i)
    year=2026 if i==0 else random.choice(range(2000,2026))
    reports.append(rid)
    vals=dict(report_id=rid,year=year,language="Chinese" if i%5==0 else "English",design="Case report" if i%3 else "Case series",extractor="Consensus")
    if rid=="WS0110": vals["linked_duplicates"]="WS0105"
    for k,v in vals.items(): R.cell(rr,rc[k],v)
    rr+=1
    n=5 if i==0 else (1 if i%3 else random.randint(3,6))
    for j in range(n):
        l2=random.uniform(-1,7); days=round(2**l2,1); age=round(random.uniform(0.2,15),2)
        pi=1/(1+math.exp(-(-1.5+TRUE*(l2-3)+0.02*age)))
        y=random.random()<pi
        v={"patient_id":"%s_P%d"%(rid,j+1),"report_id":rid,"extractor":"Consensus","age_years":age,"idx_dur_days":days,
           "viability":"Non-viable" if y else "Viable","procedure":"Total splenectomy" if y else "Splenopexy",
           "missed_dx":random.choice(["Yes","No","NR"]),"known_dx_deferred":random.choice(["Yes","No"]),"torsion_at_op":random.choice(["Yes","No"]),
           "idx_dur_lower_bound":"No","cond_gv":random.choice(["Yes","No"]),"cond_goo":"No","cond_dh":"No","aetiology":random.choice(["Congenital","Acquired"]),
           "platelets":random.randint(150,700),"ws_basis":"Torsion only, position not stated" if random.random()<.1 else "Imaging shows ectopic position or mobility"}
        if random.random()<.1: v["idx_dur_days"]="NR"
        if random.random()<.05: v["viability"]="NR"; v["procedure"]="Total splenectomy"
        if v["known_dx_deferred"]=="Yes": v["torsion_awaiting"]=random.choice(["Yes","No"])
        for k,x in v.items(): P.cell(pr,pc[k],x)
        pr+=1
wb.save(Path(__file__).parent / "synthetic_consensus.xlsx"); print(rr-4,"reports",pr-4,"patients")
