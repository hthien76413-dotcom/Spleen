from openpyxl import load_workbook
def txt(x): return '' if x is None else str(x).strip()
def load(path):
    wb=load_workbook(path,data_only=True)
    def rows(name):
        ws=wb[name]; codes=[c.value for c in ws[1]]
        return [dict(zip(codes,row)) for row in ws.iter_rows(min_row=4,values_only=True) if row[0] is not None]
    return wb, rows('Reports'), rows('Patients')
def outcome(p):
    v=txt(p.get('viability')); removal=txt(p.get('procedure')) in ['Total splenectomy','Partial splenectomy +/- splenopexy']
    if v=='Partially non-viable' and removal: return 'Non-viable'
    if v=='NR' and removal: return 'NR - removed, check reason'
    return v
