from pathlib import Path
from openpyxl import load_workbook
import random,math
p=Path(__file__).parent
w=load_workbook(p/'synthetic_consensus.xlsx'); s=w['Patients']; cols={c.value:i for i,c in enumerate(s[1])}; original=list(s.iter_rows(min_row=4,values_only=True)); s.delete_rows(4,s.max_row)
r=random.Random(20261025)
for repeat in range(6):
 for row in original:
  z=list(row); z[cols['patient_id']]=str(z[cols['patient_id']])+f'_S{repeat}'
  days=round(2**r.uniform(-1,7),3); age=float(z[cols['age_years']]); prob=1/(1+math.exp(-(-.5+.4*(math.log2(days)-3)+.02*age)))
  y=r.random()<prob
  z[cols['idx_dur_days']]=days; z[cols['viability']]='Non-viable' if y else 'Viable'; z[cols['procedure']]='Total splenectomy' if y else 'Splenopexy'
  s.append(z)
w.save(p/'synthetic_spline.xlsx')
print('Synthetic spline fixture: 768 patients; no real data')
